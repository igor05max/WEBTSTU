import io
import tempfile
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from django.test import SimpleTestCase
from docx import Document
from docx.shared import Pt, Twips
from lxml import etree

from apps.template_workspace.v2.classification.roles import RoleClassifierV2
from apps.template_workspace.v2.editor.safe_word_editor import SafeWordEditor
from apps.template_workspace.v2.inspector.document import DocumentInspector
from apps.template_workspace.v2.ooxml.namespaces import NS, qn
from apps.template_workspace.v2.planning.qwen_like import QwenLikePlanningEngine
from apps.template_workspace.v2.planning.qwen_provider import QwenPlanningProvider
from apps.template_workspace.v2.profile.template import TemplateProfileBuilder
from apps.template_workspace.v2.word_inputs import prepare_word_file


class CrossTemplateTests(SimpleTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def save(self, doc, name):
        path = self.root / (name + '.docx')
        doc.save(path)
        return path

    def article(self):
        doc = Document()
        doc.add_paragraph('A scientific title with reliable evidence')
        doc.add_paragraph('© Ivan I. Author')
        doc.add_paragraph('Abstract. Scientific content must survive formatting.')
        doc.add_paragraph('Keywords: science; formatting')
        doc.add_paragraph('1. Introduction')
        doc.add_paragraph('Body text containing original scientific findings and measurements.')
        return doc

    def template(self):
        doc = Document()
        doc.styles['Normal'].font.name = 'Calibri'
        doc.styles['Normal'].font.size = Pt(11)
        title = doc.add_paragraph('Title', style='Title')
        title.runs[0].font.size = Pt(13)
        title.paragraph_format.alignment = 0
        doc.add_paragraph('Firstname Lastname, Second Author')
        doc.add_paragraph('Abstract. An example summary.')
        doc.add_paragraph('Keywords: alpha; beta')
        doc.add_paragraph('1. Introduction')
        p = doc.add_paragraph('This is an example of the main scientific body paragraph with sufficient textual evidence.')
        p.paragraph_format.left_indent = Twips(720)
        p.paragraph_format.first_line_indent = Twips(0)
        return doc

    def render(self, source=None, template=None):
        source = self.save(source or self.article(), 'source')
        template = self.save(template or self.template(), 'template')
        output = self.root/'result.docx'
        result = SafeWordEditor().render(article_path=source, template_path=template, output_path=output)
        with ZipFile(output) as z:
            root = etree.fromstring(z.read('word/document.xml'))
        return output, root, result

    def test_short_template_title_and_exact_size_font_alignment(self):
        output, root, result = self.render()
        title = root.find('w:body/w:p', NS)
        self.assertEqual(title.find('w:pPr/w:jc', NS).get(qn('w:val')), 'left')
        self.assertEqual(title.find('w:r/w:rPr/w:sz', NS).get(qn('w:val')), '26')
        self.assertEqual(title.find('w:r/w:rPr/w:rFonts', NS).get(qn('w:ascii')), 'Calibri')
        self.assertTrue(result.metrics['native_integrity']['passed'])

    def test_one_column_template_is_not_forced_into_two_columns(self):
        _, root, _ = self.render()
        for cols in root.xpath('//w:sectPr/w:cols', namespaces=NS):
            self.assertEqual(cols.get(qn('w:num'), '1'), '1')
        body = next(p for p in root.findall('w:body/w:p', NS) if ''.join(p.xpath('.//w:t/text()', namespaces=NS)).startswith('Body text'))
        ind = body.find('w:pPr/w:ind', NS)
        self.assertEqual(ind.get(qn('w:left')), '720')
        self.assertEqual(ind.get(qn('w:firstLine')), '0')

    def test_nested_spacing_and_font_attributes_inherit_independently(self):
        doc = Document()
        doc.styles['Normal'].paragraph_format.line_spacing = 1.5
        doc.styles['Normal'].font.name = 'Arial'
        p = doc.add_paragraph('A paragraph using inherited line spacing')
        p.paragraph_format.space_after = Pt(4)
        fonts = etree.SubElement(p.runs[0]._r.get_or_add_rPr(), qn('w:rFonts'))
        fonts.set(qn('w:eastAsia'), 'SimSun')
        report = DocumentInspector(self.save(doc, 'cascade')).inspect()
        p = report.paragraphs[0]
        self.assertEqual(p.effective_formatting['paragraph']['spacing'][qn('w:line')], '360')
        self.assertEqual(p.runs[0].effective_formatting['fonts'][qn('w:ascii')], 'Arial')

    def test_authors_on_multiple_lines_do_not_become_title(self):
        doc = Document()
        doc.add_paragraph('Quantitative assessment of walking symmetry')
        doc.add_paragraph('Authors:')
        doc.add_paragraph('Антон Потлов [0000-0001-9376-3688]')
        doc.add_paragraph('Artem Obukhov [0000-0002-3450-5213],')
        doc.add_paragraph('Abstract. Real abstract text.')
        doc.add_paragraph('Keywords: gait, walking')
        doc.add_paragraph('1. Introduction')
        report = DocumentInspector(self.save(doc, 'authors')).inspect()
        roles = RoleClassifierV2().classify(report).block_roles
        self.assertEqual(roles[0]['role_hint'], 'title')
        self.assertEqual([roles[i]['role_hint'] for i in (2,3)], ['author','author'])
        self.assertEqual(roles[5]['role_hint'], 'keywords')
        _, root, _ = self.render(doc)
        self.assertNotIn('Authors:', ''.join(root.xpath('//w:t/text()', namespaces=NS)))

    def test_later_keywords_do_not_swallow_body_and_bibliography(self):
        doc = self.article()
        doc.add_paragraph('Перечень использованных источников')
        doc.add_paragraph('1. Author A. Research paper. 2025.')
        doc.add_paragraph('A. Author, B. Writer')
        doc.add_paragraph('A TRANSLATED SCIENTIFIC TITLE')
        doc.add_paragraph('Abstract. Translated summary of the same work.')
        doc.add_paragraph('Keywords: research; study')
        report = DocumentInspector(self.save(doc, 'tail')).inspect()
        roles = RoleClassifierV2().classify(report).block_roles
        self.assertEqual(roles[4]['role_hint'], 'heading_1')
        self.assertEqual(roles[6]['role_hint'], 'references_heading')
        self.assertEqual(roles[7]['role_hint'], 'reference_item')
        self.assertEqual(roles[-2]['role_hint'], 'abstract')

    def test_merging_front_matter_preserves_superscripts_and_math(self):
        doc = self.article()
        summary = doc.paragraphs[2]
        summary.text = 'Abstract.'
        extra = doc.add_paragraph('Result: 10')
        extra.add_run('-5').font.superscript = True
        extra.add_run(' cm')
        extra.add_run('2').font.superscript = True
        summary._p.addnext(extra._p)
        _, root, result = self.render(doc)
        superscripts = root.xpath('//w:vertAlign[@w:val="superscript"]', namespaces=NS)
        self.assertEqual(len(superscripts), 2)
        self.assertTrue(result.metrics['native_integrity']['passed'])

    def test_doi_in_template_bibliography_is_never_a_metadata_placeholder(self):
        doc = self.template()
        doc.add_paragraph('References')
        doc.add_paragraph('A. Writer. Reference title. DOI: 10.1234/unrelated')
        _, root, result = self.render(template=doc)
        self.assertEqual(result.metrics['template_placeholders_inserted'], 0)
        self.assertNotIn('unrelated', etree.tostring(root).decode())

    def test_dotx_changes_only_main_content_type(self):
        source = self.save(self.template(), 'dotx-source')
        template = self.root/'template.dotx'
        with ZipFile(source) as src, ZipFile(template, 'w', ZIP_DEFLATED) as dst:
            for item in src.infolist():
                payload = src.read(item.filename)
                if item.filename == '[Content_Types].xml':
                    payload = payload.replace(b'wordprocessingml.document.main+xml', b'wordprocessingml.template.main+xml')
                dst.writestr(item, payload)
        output = prepare_word_file(template, self.root/'converted.docx')
        with ZipFile(template) as src, ZipFile(output) as out:
            for name in src.namelist():
                if name != '[Content_Types].xml':
                    self.assertEqual(src.read(name), out.read(name))

    def test_single_column_qwen_cannot_enable_two_column_floating(self):
        snapshot = {'article': {'blocks': [], 'tables': []}, 'template': {
            'paragraphs': [], 'front_sequence': [], 'front_language_order': [], 'roles': {},
            'layout': {'default_body_column_count': 1}}}
        planner = QwenLikePlanningEngine()
        local = planner._local_plan(snapshot)
        result = planner._merge_provider(local, {'flow': {'allow_safe_prose_relocation': True, 'float_compact_tables_forward': True}}, snapshot)
        self.assertFalse(result.flow['allow_safe_prose_relocation'])
        self.assertFalse(result.flow['float_compact_tables_forward'])

    def test_qwen_prompt_size_is_bounded_for_long_articles(self):
        import json
        snapshot = {'article': {'blocks': [
            {'id': f'block_{i}', 'role':'heading_1', 'text':'Длинный научный текст ' * 100}
            for i in range(2000)], 'tables': []}, 'template': {'paragraphs': [], 'layout': {}}}
        compact = QwenPlanningProvider._compact_snapshot(snapshot)
        self.assertLessEqual(len(json.dumps(compact, ensure_ascii=False)), 22000)

    def test_qwen_prompt_is_bounded_for_huge_table_descriptions(self):
        import json
        snapshot = {'article': {'blocks': [], 'tables': [{'grid': '1'*60000}]*100}, 'template': {}}
        compact = QwenPlanningProvider._compact_snapshot(snapshot)
        self.assertLessEqual(len(json.dumps(compact, ensure_ascii=False)), 22000)

    def test_native_integrity_rejects_removed_scientific_text(self):
        from apps.template_workspace.v2.editor.integrity import native_integrity
        source = self.save(self.article(), 'integrity')
        with ZipFile(source) as package:
            root = etree.fromstring(package.read('word/document.xml'))
            text = root.xpath('//w:t', namespaces=NS)[-1]
            text.text = 'Missing findings'
            check = native_integrity(package, root, {})
        self.assertFalse(check['passed'])
        self.assertIn('scientific_text', check['losses'])

    def test_table_cells_cancel_inherited_manuscript_indents(self):
        source, template = self.article(), self.template()
        source.styles['Normal'].paragraph_format.first_line_indent = Twips(709)
        source.add_table(rows=2, cols=2).cell(0,0).text = '0.13'
        template.add_table(rows=2, cols=2).cell(0,0).text = 'Example'
        _, root, _ = self.render(source, template)
        for p in root.xpath('//w:tbl/w:tr/w:tc/w:p', namespaces=NS):
            ind = p.find('w:pPr/w:ind', NS)
            self.assertIsNotNone(ind)
            self.assertEqual(ind.get(qn('w:firstLine')), '0')

    def test_ole_display_equation_uses_template_text_area(self):
        source = self.article()
        p = source.add_paragraph()
        etree.SubElement(p.add_run()._r, qn('w:object'))
        _, root, result = self.render(source)
        p = root.xpath('//w:p[w:r/w:object]', namespaces=NS)[0]
        self.assertEqual(p.find('w:pPr/w:ind', NS).get(qn('w:left')), '720')
        self.assertTrue(result.metrics['native_integrity']['passed'])

    def test_story_import_keeps_colliding_source_media(self):
        from apps.template_workspace.v2.editor.template_evidence import StoryImporter
        template_data, source_data = io.BytesIO(), io.BytesIO()
        content_types = b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="png" ContentType="image/png"/></Types>'
        with ZipFile(template_data, 'w') as z:
            z.writestr('[Content_Types].xml', content_types)
            z.writestr('word/header1.xml', b'<header/>')
            z.writestr('word/media/image1.png', b'template logo')
            z.writestr('word/_rels/header1.xml.rels', b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="image" Target="media/image1.png"/></Relationships>')
        with ZipFile(source_data, 'w') as z:
            z.writestr('[Content_Types].xml', content_types)
            z.writestr('word/media/image1.png', b'article figure')
        with ZipFile(template_data) as template, ZipFile(source_data) as article:
            importer = StoryImporter(template, article)
            copied = importer.copy('word/header1.xml')
        self.assertEqual(copied, 'word/header1_v2.xml')
        self.assertNotIn('word/media/image1.png', importer.parts)
        self.assertEqual(importer.parts['word/media/image1_v2.png'], b'template logo')
        self.assertIn(b'media/image1_v2.png', importer.parts['word/_rels/header1_v2.xml.rels'])
