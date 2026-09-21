"""Compile measured JAMT rules into the native editor's formatting contract.

The empty OOXML carrier supplies section geometry and journal furniture only.
It contains no sample article, identifiers, dates or research text. Footer
authors, when available, come only from the current source manuscript.
All paragraph formatting comes directly from the versioned JSON style.
"""
import json
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Pt, Twips

from ..inspector.document import DocumentInspector
from ..models.template_profile import RoleFormatProfile, TemplateProfile
from ..ooxml.namespaces import qn
from ..profile.template import TemplateProfileBuilder


def load_style():
    return json.loads(Path(__file__).with_name("jamt.json").read_text(encoding="utf-8"))


def visual_review_rules():
    """Trusted style evidence from the shipped profile, never from a manuscript."""
    style = load_style()
    return {
        "id": style["id"], "version": style["version"], "font": style["font"],
        "body_columns": style["columns"],
        "full_width": ["bilingual front matter", "author biographies", "editorial dates", "license"],
        "wide_objects": "Large figures and tables may span both columns.",
        "front_language_order": style["front_language_order"],
        "typography": {role: {key: value for key, value in style["roles"][role].items()
                              if key in {"pt", "align", "bold", "italic", "first_indent"}}
                       for role in ("title", "author", "affiliation", "abstract", "body", "heading_1",
                                    "heading_2", "figure_caption", "table_caption", "table_body", "reference_item",
                                    "funding_text", "author_information", "author_bio", "editorial_metadata",
                                    "doi_metadata", "article_type", "rubric", "received_metadata")},
        "journal_header": "Published reference: black bold italic, alternating outer alignment and double rules. Preserve existing source issue metadata.",
        "quality_rules": style.get('quality_rules', []),
    }


def _rule(parent, name, **attributes):
    element = OxmlElement("w:" + name)
    for key, value in attributes.items():
        element.set(qn("w:" + key), str(value))
    parent.append(element)
    return element


def _furniture(paragraph, style, *, footer=False, author="", even=False):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1
    properties = paragraph._p.get_or_add_pPr()
    if not footer:
        width = style['page_twips']['w'] - style['margins_twips']['left'] - style['margins_twips']['right']
        _rule(properties, 'framePr', w=width, h=340, hAnchor='page', vAnchor='page',
              x=style['margins_twips']['left'], y=style['furniture']['header_frame_y_twips'],
              wrap='notBeside', hRule='atLeast')
    _rule(properties, "jc", val="left" if footer or even else "right")
    borders = _rule(properties, "pBdr")
    _rule(borders, "top" if footer else "bottom", val="thickThinSmallGap" if footer else "thinThickSmallGap", sz=24, space=1, color="000000")
    run = paragraph.add_run() if footer else paragraph.add_run(style["journal"])
    run.font.name = style["font"]
    run.font.size = Pt(9 if footer else 10)
    run.bold = True
    run.italic = not footer
    _rule(run._r.get_or_add_rPr(), "color", val="000000")
    if footer:
        _rule(run._r, "fldChar", fldCharType="begin")
        instruction = _rule(run._r, "instrText")
        instruction.text = " PAGE "
        _rule(run._r, "fldChar", fldCharType="separate")
        _rule(run._r, "t").text = "1"
        _rule(run._r, "fldChar", fldCharType="end")
        width = style['page_twips']['w'] - style['margins_twips']['left'] - style['margins_twips']['right']
        tabs = _rule(properties, 'tabs')
        _rule(tabs, 'tab', val='center', pos=width // 2)
        _rule(tabs, 'tab', val='right', pos=width)
        name_run = paragraph.add_run('\t' + author + ('' if even else '\t'))
        name_run.font.name = style['font']
        name_run.font.size = Pt(10)
        name_run.italic = True
        if not even:
            paragraph._p.remove(run._r);paragraph._p.append(run._r)


def prepare_jamt_style(directory, *, article_report=None, article_structure=None):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    style = load_style()
    document = Document()
    document.core_properties.title = "JAMT style geometry"
    document.core_properties.author = ""
    normal = document.styles["Normal"]
    normal.font.name = style["font"]
    normal.font.size = Pt(style['roles']['body']['pt'])
    # The stock DOCX theme otherwise wins over explicit font names when native
    # Header/Footer formatting is materialised into an unrelated manuscript.
    for fonts in document.styles.element.xpath('.//w:rFonts'):
        for key in list(fonts.attrib):
            if key.endswith('Theme'):
                del fonts.attrib[key]
        for script in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + script), style['font'])
    section = document.sections[0]
    section.page_width = Twips(style["page_twips"]["w"])
    section.page_height = Twips(style["page_twips"]["h"])
    for key in ("top", "bottom", "left", "right"):
        setattr(section, key + "_margin", Twips(style["margins_twips"][key]))
    section.header_distance = Twips(style["margins_twips"]["header"])
    section.footer_distance = Twips(style["margins_twips"]["footer"])
    document.settings.odd_and_even_pages_header_footer = True
    columns = section._sectPr.find(qn("w:cols"))
    columns.set(qn("w:num"), str(style["columns"]))
    columns.set(qn("w:space"), str(style["column_gap_twips"]))
    _furniture(section.header.paragraphs[0], style)
    _furniture(section.even_page_header.paragraphs[0], style, even=True)
    author = ''
    if article_report is not None and article_structure is not None:
        from ..editor.safe_word_editor import _article_author_shortline
        author = _article_author_shortline(article_structure, article_report)
    _furniture(section.footer.paragraphs[0], style, footer=True, author=author)
    _furniture(section.even_page_footer.paragraphs[0], style, footer=True, author=author, even=True)
    carrier = directory / "style-geometry.docx"
    document.save(carrier)
    report = DocumentInspector(carrier).inspect()
    roles = {}
    for role, rule in style["roles"].items():
        paragraph = {
            "alignment": rule["align"],
            "spacing": {qn("w:before"): str(round(rule.get("before", 0) * 20)),
                        qn("w:after"): str(round(rule.get("after", 0) * 20)),
                        qn("w:line"): str(round(240 * style['body_line_multiple'])) if role in {"body", "funding_text", "acknowledgements_text", "conflict_text"} else "240", qn("w:lineRule"): "auto"},
            "indentation": {qn("w:firstLine"): str(round(rule.get("first_indent", 0) * 20))},
        }
        if "right_tab" in rule:
            paragraph["tabs"] = [{qn("w:val"): "right", qn("w:pos"): str(rule["right_tab"])}]
        roles[role] = RoleFormatProfile(
            role=role, confidence=1, examples_count=len(style["references"]),
            typical_paragraph_formatting=paragraph,
            typical_run_formatting={"fonts": {qn("w:ascii"): style["font"], qn("w:hAnsi"): style["font"]},
                                    "size": str(round(rule["pt"] * 2)),
                                    "bold": rule.get("bold", False), "italic": rule.get("italic", False)},
            derived_from="jamt:" + style["version"],
        )
    layout = TemplateProfileBuilder()._layout_profile(report)
    layout.default_body_column_count = style["columns"]
    profile = TemplateProfile(
        source_path="style:jamt:" + style["version"], roles=roles, layout=layout,
        table_profiles={}, drawing_profiles={}, missing_roles=[],
        front_language_order=style["front_language_order"],
        front_role_sequence=style["front_role_sequence"],
        front_merge_roles=["author", "affiliation", "abstract"],
        style_id=style["id"], style_version=style["version"],
    )
    (directory / "style_source.json").write_text(json.dumps(style, ensure_ascii=False, indent=2), encoding="utf-8")
    return carrier, report, profile
