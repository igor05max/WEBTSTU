from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from zipfile import ZipFile

from django.test import SimpleTestCase
from docx import Document
from lxml import etree
import pymupdf

from .v2.editor.object_flow import normalize_vml_canvases
from .v2.editor.safe_word_editor import SafeWordEditor
from .v2.ooxml.namespaces import NS, qn
from .v2.readability import empty_body_pages, review_pdf
from .v2.styles.jamt import prepare_jamt_style


def canvas():
    return etree.fromstring(f'''<w:p xmlns:w="{NS['w']}" xmlns:v="{NS['v']}">
      <w:r><w:pict><v:group id="canvas" editas="canvas" coordorigin="686,7048"
        coordsize="9865" style="width:493.25pt;height:50pt">
        <v:shape id="frame" style="position:absolute;left:686;top:7048;width:9865;height:1000"/>
        <v:shape id="license" style="position:absolute;left:2394;top:7299;width:8141;height:581">
          <v:textbox><w:txbxContent><w:p><w:r><w:rPr><w:sz w:val="18"/></w:rPr>
            <w:t>Copyright: Original authors. This original license must remain visible.</w:t>
          </w:r></w:p></w:txbxContent></v:textbox>
        </v:shape>
      </v:group></w:pict></w:r>
    </w:p>''')


class VmlCanvasTests(SimpleTestCase):
    def test_editor_preserves_native_canvas_text_and_completes_its_geometry(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            document = Document()
            document.add_paragraph('Original research', style='Title')
            document.add_paragraph('1. Introduction', style='Heading 1')
            document.add_paragraph('The measured result is 42.5 MPa.')
            paragraph = canvas()
            document.element.body.insert(len(document.element.body) - 1, paragraph)
            source = root/'source.docx'; document.save(source)
            carrier, template, profile = prepare_jamt_style(root)
            result = SafeWordEditor().render(article_path=source, template_path=carrier,
                template_report=template, template_profile=profile, output_path=root/'result.docx')
            self.assertTrue(result.metrics['native_integrity']['passed'])
            self.assertEqual(result.metrics['vml_canvases_normalized'], 1)
            with ZipFile(root/'result.docx') as archive:
                output = etree.fromstring(archive.read('word/document.xml'))
            group = output.find('.//v:group', NS)
            self.assertEqual(group.get('coordsize'), '9865,1000')
            before_text = paragraph.xpath('.//w:txbxContent//w:t/text()', namespaces=NS)
            self.assertEqual(group.xpath('.//w:txbxContent//w:t/text()', namespaces=NS), before_text)
            self.assertEqual(group.xpath('.//w:sz/@w:val', namespaces=NS), ['18'])
            self.assertIn('mso-wrap-style:square', group.findall('v:shape', NS)[1].get('style'))
            self.assertEqual(normalize_vml_canvases(output), 0)

    def test_ambiguous_or_complete_geometry_is_left_untouched(self):
        for case in ('complete', 'no_frame', 'wrong_origin', 'wrong_aspect'):
            with self.subTest(case=case):
                paragraph = canvas(); group = paragraph.find('.//v:group', NS)
                frame = group.find('v:shape', NS)
                if case == 'complete':
                    group.set('coordsize', '9865,1000')
                elif case == 'no_frame':
                    group.remove(frame)
                elif case == 'wrong_origin':
                    group.set('coordorigin', '0,0')
                else:
                    group.set('style', 'width:493.25pt;height:100pt')
                before = etree.tostring(paragraph)
                self.assertEqual(normalize_vml_canvases(paragraph), 0)
                self.assertEqual(etree.tostring(paragraph), before)

    def test_explicit_no_wrap_and_auto_fit_remain_authoritative(self):
        for case in ('no_wrap', 'auto_fit'):
            with self.subTest(case=case):
                paragraph = canvas()
                shape = paragraph.xpath('.//v:shape[@id="license"]', namespaces=NS)[0]
                if case == 'no_wrap':
                    shape.set('style', shape.get('style') + ';mso-wrap-style:none')
                else:
                    shape.find('v:textbox', NS).set('style', 'mso-fit-shape-to-text:t')
                before = etree.tostring(shape)
                self.assertEqual(normalize_vml_canvases(paragraph), 1)
                self.assertEqual(etree.tostring(shape), before)


class EmptyPageReviewTests(SimpleTestCase):
    def test_empty_journal_page_is_flagged_even_when_ai_reports_no_issues(self):
        with TemporaryDirectory() as directory, pymupdf.open() as pdf:
            for index in range(4):
                page = pdf.new_page()
                page.insert_text((50, 35), 'Journal running header')
                page.insert_text((50, 810), str(index + 1))
            pdf[0].insert_text((50, 130), 'Original research is visible.')
            pdf[2].draw_line((100, 100), (100, 200))  # Vector-only art is content.
            pdf[3].insert_image(pymupdf.Rect(100, 100, 300, 400), stream=pdf[0].get_pixmap().tobytes('png'))
            self.assertEqual(empty_body_pages(pdf), [2])
            path = Path(directory)/'journal.pdf'; pdf.save(path)
            provider = SimpleNamespace(review=lambda *args, **kwargs: [])
            report = review_pdf(path, provider=provider)
            self.assertEqual(report['pages_checked'], [1, 2, 3, 4])
            self.assertEqual([(item['page'], item['kind']) for item in report['issues']], [(2, 'empty_page')])
            self.assertTrue(report['issues'][0]['advisory'])

    def test_short_document_at_page_edge_is_not_assumed_to_be_only_a_header(self):
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((40, 40), 'Short document')
            self.assertEqual(empty_body_pages(pdf), [])
