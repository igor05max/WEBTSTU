import json
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile
from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

from apps.template_workspace import test_v2_cross_template as fixtures
from apps.template_workspace.v2.editor.protected_blocks import (
    code_nodes, is_code, is_display_equation, visible_code_text, format_display_equation)
from apps.template_workspace.v2.editor.safe_word_editor import _rebalance_data_column_widths
from apps.template_workspace.v2.editor.template_evidence import NativeTemplateFormatting
from apps.template_workspace.v2.ooxml.namespaces import NS, qn
from apps.template_workspace.v2.readability import validate_issues, review_pdf, run_readability_review


class ReadabilityLayoutTests(SimpleTestCase):
    setUp = fixtures.CrossTemplateTests.setUp
    save = fixtures.CrossTemplateTests.save
    article = fixtures.CrossTemplateTests.article
    template = fixtures.CrossTemplateTests.template
    render = fixtures.CrossTemplateTests.render
    def test_json_soft_breaks_are_not_justified_and_content_is_exact(self):
        source = self.article()
        code = '{\n  "input": {\n    "name": "Название направления",\n    "scores": [1, 2]\n  }\n}'
        p = source.add_paragraph(code)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _, root, result = self.render(source)
        out = next(p for p in root.findall('w:body/w:p', NS) if '"input"' in visible_code_text(p))
        self.assertEqual(visible_code_text(out), code)
        self.assertEqual(out.find('w:pPr/w:jc', NS).get(qn('w:val')), 'left')
        self.assertEqual(out.find('w:pPr/w:ind', NS).get(qn('w:firstLine')), '0')
        self.assertTrue(result.metrics['code_content_preserved'])

    def test_multiline_json_and_python_not_plain_prose(self):
        self.assertTrue(is_code('def f(x):\n    return x+1'))
        self.assertFalse(is_code('We compare model: Qwen.\nResults are shown below.'))
        doc = Document()
        for text in ['{','  "x": [1, 2]','}']:
            doc.add_paragraph(text)
        self.assertEqual(len(code_nodes(doc.element.body)),3)

    def test_display_and_inline_equations_are_distinct(self):
        doc = Document()
        display = doc.add_paragraph()._p
        inline = doc.add_paragraph('Here x is a parameter.')._p
        for p in [display, inline]:
            p.append(etree.Element(qn('m:oMath')))
        display,inline = [etree.fromstring(etree.tostring(p)) for p in [display,inline]]
        self.assertTrue(is_display_equation(display))
        self.assertFalse(is_display_equation(inline))
        format_display_equation(display, SimpleNamespace(body_left_twips=2608,body_right_twips=0,column_width_twips=7800))
        self.assertEqual(display.find('w:pPr/w:jc',NS).get(qn('w:val')),'center')
        self.assertEqual(display.find('w:pPr/w:ind',NS).get(qn('w:firstLine')),'0')

    def test_native_template_defaults_do_not_inherit_source_spacing(self):
        doc=Document()
        doc.styles['Normal'].paragraph_format.space_after=None
        # Clear Word's docDefaults, as in the real JAMT template.
        defaults=doc.styles.element.find('w:docDefaults/w:pPrDefault',NS)
        for child in list(defaults): defaults.remove(child)
        p=doc.add_paragraph('Journal header')._p
        path=self.save(doc,'defaults')
        with ZipFile(path) as z:
            pp=NativeTemplateFormatting(z).properties(p)
        self.assertEqual(pp.find('w:spacing',NS).get(qn('w:after')),'0')

    def test_numbered_formula_manual_tabs_are_bounded(self):
        p=etree.Element(qn('w:p'),nsmap=NS)
        p.append(etree.Element(qn('m:oMath')))
        r=etree.SubElement(p,qn('w:r'))
        for _ in range(6):etree.SubElement(r,qn('w:tab'))
        etree.SubElement(r,qn('w:t')).text='(1)'
        layout=SimpleNamespace(body_left_twips=0,body_right_twips=0,column_width_twips=9600)
        format_display_equation(p,layout)
        self.assertEqual(len(p.xpath('./w:r/w:tab',namespaces=NS)),2)
        self.assertEqual(p.xpath('./w:pPr/w:tabs/w:tab/@w:pos',namespaces=NS),['4800','9600'])
        self.assertEqual(p.xpath('.//w:t/text()',namespaces=NS),['(1)'])
        format_display_equation(p,layout)
        self.assertEqual(len(p.xpath('./w:r/w:tab',namespaces=NS)),2)

    def test_hidden_mathtype_control_does_not_leak_or_hide_following_text(self):
        from apps.template_workspace.v2.editor.protected_blocks import hide_equation_control_fields
        doc=Document(); p=doc.add_paragraph()._p
        for kind, text, hidden in [('begin','',False),('instr',' MACROBUTTON MTEditEquationSection2 ',False),
                                   ('instr','Equation Chapter 1 Section 1',True),('begin','',False),
                                   ('instr',' SEQ MTEqn \\h ',False),('end','',False),('end','',False)]:
            r=etree.SubElement(p,qn('w:r'))
            if hidden:etree.SubElement(etree.SubElement(r,qn('w:rPr')),qn('w:vanish'))
            n=etree.SubElement(r,qn('w:instrText' if kind=='instr' else 'w:fldChar'))
            if kind=='instr':n.text=text
            else:n.set(qn('w:fldCharType'),kind)
        doc.paragraphs[0].add_run('УДК 004.9')
        body=etree.fromstring(etree.tostring(doc.element.body))
        p=body.find('w:p',NS)
        before=p.xpath('.//w:instrText/text()',namespaces=NS)
        self.assertEqual(hide_equation_control_fields(body),1)
        self.assertEqual(p.xpath('.//w:instrText/text()',namespaces=NS),before)
        self.assertEqual(len(p.xpath('./w:r/w:rPr/w:vanish',namespaces=NS)),7)
        self.assertIsNone(p.findall('w:r',NS)[-1].find('w:rPr/w:vanish',NS))

    def test_compact_table_keep_next_overrides_explicit_false(self):
        from apps.template_workspace.v2.editor.safe_word_editor import _keep_compact_table_together
        doc=Document(); table=doc.add_table(rows=11,cols=2)
        for row in table.rows:
            for c in row.cells:
                c.text='data';c.paragraphs[0].paragraph_format.keep_with_next=False
        root=etree.fromstring(etree.tostring(table._tbl))
        _keep_compact_table_together(root,None)
        self.assertEqual(root.find('w:tr/w:tc/w:p/w:pPr/w:keepNext',NS).get(qn('w:val')),'1')

    def test_long_identifier_column_has_room_and_index_stays_compact(self):
        doc=Document(); table=doc.add_table(rows=4,cols=4)
        rows=[['No','Model','Size','Quantization'],['1','DeepSeek-R1-Distill-Qwen-7B-Q4_K_M.gguf','7B','Q4_K_M'],
              ['2','NVIDIA-Nemotron-3-Nano-4B-Q8_0.gguf','4B','Q8_0'],['3','gpt-oss-20b-MXFP4.gguf','20B','MXFP4']]
        for r,values in zip(table.rows,rows):
            for cell,value in zip(r.cells,values):cell.text=value
        widths=_rebalance_data_column_widths([2500]*4,10000,etree.fromstring(etree.tostring(table._tbl)))
        self.assertEqual(sum(widths),10000)
        self.assertLess(widths[0],1000)
        self.assertGreater(widths[1],5000)

    def test_table_style_border_is_resolved(self):
        doc=Document(); table=doc.add_table(rows=2,cols=2); table.style='Table Grid'
        path=self.save(doc,'borders')
        with ZipFile(path) as z:
            props=NativeTemplateFormatting(z).table_properties(table._tbl,'tblPr')
        self.assertIsNotNone(props.find('w:tblBorders/w:top',NS))

    def test_converted_ole_equation_uses_stream_evidence_not_empty_progid(self):
        from io import BytesIO
        from apps.template_workspace.v2.editor.protected_blocks import legacy_equation_ids
        stream=BytesIO()
        with ZipFile(stream,'w') as z:
            z.writestr('word/_rels/document.xml.rels', '<Relationships><Relationship Id="eq" Type="urn:test/oleObject" Target="embeddings/eq.bin"/><Relationship Id="chart" Type="urn:test/oleObject" Target="embeddings/chart.bin"/></Relationships>')
            z.writestr('word/embeddings/eq.bin',b'MathType'+ 'Equation Native'.encode('utf-16le'))
            z.writestr('word/embeddings/chart.bin',b'Excel Chart')
        with ZipFile(stream) as z: ids=legacy_equation_ids(z)
        self.assertEqual(ids,{'eq'})
        p=etree.Element(qn('w:p'),nsmap=NS)
        ole=etree.SubElement(p,qn('o:OLEObject')); ole.set(qn('r:id'),'eq');ole.set('ProgID','')
        self.assertFalse(is_display_equation(p))
        self.assertTrue(is_display_equation(p,ids))
        ole.set(qn('r:id'),'chart')
        self.assertFalse(is_display_equation(p,ids))

    def test_dense_identifier_table_never_gets_negative_column_width(self):
        doc=Document(); table=doc.add_table(rows=2,cols=30)
        for i in range(30):
            table.cell(0,i).text='Column'
            table.cell(1,i).text='1'
        table.cell(1,0).text='Long-Identifier-That-Must-Remain-Intact'
        widths=_rebalance_data_column_widths([200]*30,6000,etree.fromstring(etree.tostring(table._tbl)))
        self.assertEqual(sum(widths),6000)
        self.assertGreater(min(widths),0)


class VisualReviewTests(SimpleTestCase):
    def test_model_cannot_supply_page_or_operations(self):
        issues=validate_issues({'issues':[{'page':999,'kind':'code_spacing','severity':'high',
                        'description':'Spacing is stretched','operation':'delete'}]},3)
        self.assertEqual(issues[0]['page'],3)
        self.assertNotIn('operation',issues[0])
        self.assertEqual(validate_issues({'issues':[{'kind':'rewrite','severity':'high','description':'change text'}]},1),[])
        with self.assertRaises(ValueError):validate_issues({'ok':True},1)

    def test_partial_coverage_and_failure_never_claim_all_pages(self):
        import pymupdf
        import tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'pages.pdf'
            with pymupdf.open() as doc:
                for i in range(3):doc.new_page().insert_text((40,40),f'Page {i+1}')
                doc.save(path)
            provider=SimpleNamespace(review=lambda *args,**kwargs:[])
            report=review_pdf(path,provider=provider,max_pages=1)
            self.assertEqual(report['status'],'partial')
            self.assertEqual(len(report['pages_not_checked']),2)
            with patch.object(provider,'review',side_effect=TimeoutError):
                report=review_pdf(path,provider=provider,max_pages=3)
            self.assertEqual(report['status'],'unavailable')
            self.assertEqual(report['pages_checked'],[])

    @override_settings(TEMPLATE_V2_VISUAL_REVIEW_ENABLED=False)
    def test_disabled_review_does_not_render_or_call_model(self):
        import tempfile
        with tempfile.TemporaryDirectory() as tmp:
            with patch('apps.submissions.document_preview.convert_word_path_to_pdf') as render:
                report=run_readability_review('missing.docx',tmp)
        self.assertEqual(report['status'],'disabled')
        render.assert_not_called()
