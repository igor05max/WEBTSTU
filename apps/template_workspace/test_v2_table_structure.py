from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace

from django.test import SimpleTestCase
from docx import Document
from lxml import etree

from apps.template_workspace.v2.editor.layout_fidelity import scientific_table_rules
from apps.template_workspace.v2.editor.table_structure import analyze_table_structure
from apps.template_workspace.v2.editor.protected_blocks import prop
from apps.template_workspace.v2.editor.safe_word_editor import _apply_template_table_evidence
from apps.template_workspace.v2.inspector.document import DocumentInspector
from apps.template_workspace.v2.ooxml.namespaces import NS, qn


INFO = SimpleNamespace(classification='DATA_TABLE')


def grouped(groups=(2, 3, 1), depth=2):
    doc = Document()
    table = doc.add_table(rows=depth+sum(groups), cols=6)
    table.style = 'Table Grid'
    table.cell(0, 0).merge(table.cell(depth-1, 0)).text = 'Material'
    table.cell(0, 1).merge(table.cell(depth-1, 1)).text = 'Parameter'
    table.cell(0, 2).merge(table.cell(0, 5)).text = 'Measured values'
    if depth == 3:
        table.cell(1, 2).merge(table.cell(1, 3)).text = 'Before'
        table.cell(1, 4).merge(table.cell(1, 5)).text = 'After'
    for c in range(2, 6):
        table.cell(depth-1, c).text = str(c)
    cursor = depth
    for gi, count in enumerate(groups):
        table.cell(cursor, 0).merge(table.cell(cursor+count-1, 0)).text = f'Material {gi}'
        for ri in range(cursor, cursor+count):
            table.cell(ri, 1).text = f'x{ri}'
            for c in range(2, 6):
                table.cell(ri, c).text = str(ri*10+c)
        cursor += count
    return etree.fromstring(etree.tostring(table._tbl))


def border_values(table, row, side):
    return table.xpath(f'./w:tr[{row+1}]/w:tc/w:tcPr/w:tcBorders/w:{side}/@w:val', namespaces=NS)


class TableStructureTests(SimpleTestCase):
    def test_unequal_groups_and_two_level_header(self):
        table = grouped()
        plan = analyze_table_structure(table)
        self.assertEqual(plan.header_rows, 2)
        self.assertEqual(plan.group_boundaries, [4, 7])
        self.assertTrue(plan.hierarchical)

    def test_three_level_header_and_other_group_lengths(self):
        plan = analyze_table_structure(grouped((3, 5), depth=3))
        self.assertEqual(plan.header_rows, 3)
        self.assertEqual(plan.group_boundaries, [6])

    def test_rules_only_at_whole_header_and_group_boundaries(self):
        table = grouped()
        words = table.xpath('.//w:t/text()', namespaces=NS)
        merges = table.xpath('.//w:vMerge/@w:val|.//w:gridSpan/@w:val', namespaces=NS)
        scientific_table_rules(table, INFO)
        for ri in range(8):
            self.assertTrue(all(v == ('single' if ri in {0, 2, 4, 7} else 'nil')
                                for v in border_values(table, ri, 'top')), ri)
            self.assertTrue(all(v == ('single' if ri+1 in {2, 4, 7, 8} else 'nil')
                                for v in border_values(table, ri, 'bottom')), ri)
        self.assertEqual(words, table.xpath('.//w:t/text()', namespaces=NS))
        self.assertEqual(merges, table.xpath('.//w:vMerge/@w:val|.//w:gridSpan/@w:val', namespaces=NS))
        headers=table.findall('w:tr/w:trPr/w:tblHeader',NS)
        self.assertEqual(len(headers),2)
        self.assertTrue(all(not n.attrib for n in headers))
        before = etree.tostring(table)
        scientific_table_rules(table, INFO)
        self.assertEqual(etree.tostring(table), before)

    def test_style_grid_and_row_exceptions_cannot_resurrect_lines(self):
        table = grouped()
        for parent in [table.find('w:tblPr',NS), prop(table.find('w:tr',NS), 'tblPrEx')]:
            borders = prop(parent, 'tblBorders')
            for side in ('top','bottom','insideH','insideV'):
                prop(borders, side, val='single', sz=12)
        for cp in table.findall('w:tr/w:tc/w:tcPr', NS):
            borders = prop(cp, 'tcBorders')
            for side in ('top','bottom','insideH','insideV'):
                prop(borders, side, val='single', sz=12)
        scientific_table_rules(table, INFO)
        self.assertEqual(table.find('w:tblPr/w:tblBorders/w:insideH',NS).get(qn('w:val')), 'nil')
        self.assertTrue(all(n.get(qn('w:val')) == 'nil' for n in table.findall('w:tr/w:tblPrEx/w:tblBorders/*',NS)))
        self.assertTrue(all(v == 'nil' for v in border_values(table,2,'bottom')))

    def test_outer_groups_do_not_get_inner_subgroup_rules(self):
        table = grouped((4, 3))
        rows = table.findall('w:tr',NS)
        for first, count in [(2,2),(4,2),(6,3)]:
            for ri in range(first, first+count):
                cell = rows[ri].findall('w:tc',NS)[1]
                prop(cell.find('w:tcPr',NS), 'vMerge', val='restart' if ri==first else 'continue')
        self.assertEqual(analyze_table_structure(table).group_boundaries, [6])

    def test_headerless_grouped_numeric_table(self):
        table = grouped((2,3))
        for row in table.findall('w:tr',NS)[:2]: table.remove(row)
        plan = analyze_table_structure(table)
        self.assertEqual(plan.header_rows, 0)
        self.assertEqual(plan.group_boundaries, [2])
        scientific_table_rules(table, INFO)
        self.assertTrue(all(v=='nil' for v in border_values(table,0,'bottom')))

    def test_blank_or_repeated_unmerged_labels_are_not_invented_groups(self):
        table = grouped()
        for vm in table.findall('w:tr/w:tc/w:tcPr/w:vMerge',NS): vm.getparent().remove(vm)
        self.assertEqual(analyze_table_structure(table).group_boundaries, [])

    def test_header_span_subdivision_without_vertical_merges(self):
        table = grouped()
        for row in table.findall('w:tr',NS)[:2]:
            for vm in row.findall('w:tc/w:tcPr/w:vMerge',NS): vm.getparent().remove(vm)
        self.assertEqual(analyze_table_structure(table).header_rows, 2)

    def test_invalid_continuation_does_not_get_guessed_borders(self):
        table = grouped()
        vm = table.find('w:tr/w:tc/w:tcPr/w:vMerge',NS)
        vm.set(qn('w:val'),'continue')
        before = etree.tostring(table)
        self.assertFalse(analyze_table_structure(table).valid)
        self.assertEqual(scientific_table_rules(table,INFO),0)
        self.assertEqual(etree.tostring(table),before)

    def test_grid_before_offsets_are_respected(self):
        table = grouped()
        for row in table.findall('w:tr',NS): prop(prop(row,'trPr'),'gridBefore',val=1)
        plan = analyze_table_structure(table)
        self.assertEqual(plan.rows[0][0].start,1)
        self.assertEqual(plan.group_boundaries,[4,7])

    def test_true_nested_container_not_flattened_or_styled_as_data(self):
        doc = Document(); outer = doc.add_table(rows=2,cols=1)
        outer.cell(0,0).text='Native layout caption'
        table = etree.fromstring(etree.tostring(outer._tbl))
        cell = table.find('w:tr/w:tc',NS)
        nested = grouped(); cell.insert(1,nested)
        outer_pr = etree.tostring(table.find('w:tblPr',NS))
        self.assertTrue(analyze_table_structure(table).nested)
        self.assertEqual(scientific_table_rules(table,SimpleNamespace(classification='LAYOUT_TABLE')),1)
        self.assertEqual(etree.tostring(table.find('w:tblPr',NS)),outer_pr)
        self.assertIs(cell.find('w:tbl',NS), nested)
        self.assertEqual(table.xpath('./w:tr/w:tc/w:tcPr/w:tcBorders',namespaces=NS),[])
        self.assertIsNone(_apply_template_table_evidence(table,None,None))

    def test_figure_container_is_untouched_even_if_it_has_merges(self):
        table=grouped(); before=etree.tostring(table)
        self.assertEqual(scientific_table_rules(table,SimpleNamespace(classification='FIGURE_CONTAINER')),0)
        self.assertEqual(etree.tostring(table),before)

    def test_inspector_counts_descendants_not_minus_one(self):
        with TemporaryDirectory() as tmp:
            doc=Document();table=doc.add_table(rows=2,cols=2)
            table.cell(0,0).text='Header';table.cell(1,0).text='Value'
            path=Path(tmp)/'table.docx';doc.save(path)
            info=DocumentInspector(path).inspect().tables[0]
            self.assertEqual(info.nested_table_count,0)
            self.assertEqual(info.classification,'DATA_TABLE')
            table.cell(1,0).add_table(rows=2,cols=2)
            doc.save(path)
            info=DocumentInspector(path).inspect().tables[0]
            self.assertEqual(info.nested_table_count,1)
            self.assertEqual(info.rows[1][0].nested_table_count,1)
            self.assertEqual(info.classification,'LAYOUT_TABLE')

    def test_single_body_group_has_no_internal_rule(self):
        table=grouped((5,))
        # Also works without a multi-row header or any inter-group boundary.
        for row in table.findall('w:tr',NS)[:2]: table.remove(row)
        self.assertTrue(analyze_table_structure(table).hierarchical)
        scientific_table_rules(table,INFO)
        self.assertTrue(all(v=='nil' for v in border_values(table,3,'bottom')))

    def test_alignment_uses_body_not_numeric_subheaders(self):
        table=grouped()
        row=table.findall('w:tr',NS)[2]
        cell=row.findall('w:tc',NS)[1]
        cell.find('w:p/w:r/w:t',NS).text='A descriptive parameter name requiring left alignment'
        scientific_table_rules(table,INFO)
        for ri in (0,1):
            self.assertTrue(all(v=='center' for v in table.xpath(f'./w:tr[{ri+1}]/w:tc/w:p/w:pPr/w:jc/@w:val',namespaces=NS)))
        self.assertEqual(cell.find('w:p/w:pPr/w:jc',NS).get(qn('w:val')),'left')
