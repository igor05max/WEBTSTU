from __future__ import annotations

import io
import tempfile
import unittest
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings, skipUnlessDBFeature
from django.urls import reverse
from docx import Document
from docx.enum.section import WD_SECTION
from lxml import etree

from apps.template_workspace.models import TemplateJob
from apps.template_workspace.v2.classification.roles import RoleClassifierV2
from apps.template_workspace.v2.comparison.document_diff import DocumentDiffBuilder
from apps.template_workspace.v2.editor.safe_word_editor import SafeWordEditor
from apps.template_workspace.v2.inspector.document import DocumentInspector
from apps.template_workspace.v2.mapping.preview import RoleMatcher
from apps.template_workspace.v2.planning.qwen_like import QwenLikePlanningEngine
from apps.template_workspace.v2.planning.qwen_provider import QwenPlanningProvider, build_planning_engine
from apps.template_workspace.v2.profile.template import TemplateProfileBuilder
from apps.template_workspace.v2.services import analysis_directory, result_docx_path, run_v2_job


def first_existing(*paths: str) -> Path:
    for path in paths:
        candidate = Path(path)
        if candidate.exists():
            return candidate
    return Path(paths[0])


FIXTURES = {
    "balabanov_source": first_existing(
        r"C:\Users\Igoryok\Downloads\Statya_Balabanov_trans (2).docx",
        r"C:\Users\Igoryok\Downloads\Statya_Balabanov_trans (1).docx",
        r"C:\Users\Igoryok\Downloads\Statya_Balabanov_trans.docx",
    ),
    "balabanov_formatted": first_existing(
        r"C:\Users\Igoryok\Downloads\03_Balabanov_Dyachkova_Gutnik_Chapaksov_Burakova_118-128 (1).docx",
        r"C:\Users\Igoryok\Downloads\03_Balabanov_Dyachkova_Gutnik_Chapaksov_Burakova_118-128.docx",
    ),
    "tyutyunnik_formatted": first_existing(
        r"C:\Users\Igoryok\Downloads\01_Tyutyunnik_98-103 (1).docx",
        r"C:\Users\Igoryok\Downloads\01_Tyutyunnik_98-103.docx",
    ),
}
REAL_FIXTURES_AVAILABLE = all(path.exists() for path in FIXTURES.values())


def file_hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def make_docx_with_core_objects(path: Path) -> None:
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Merged"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "B"
    document.add_paragraph("After table")
    section = document.sections[0]
    section.header.paragraphs[0].text = "Header PAGE"
    section.footer.paragraphs[0].text = "Footer"
    document.add_section(WD_SECTION.CONTINUOUS)
    buffer = io.BytesIO()
    document.save(buffer)

    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rIdLink" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://example.test" TargetMode="External"/>
</Relationships>"""
    paragraph_xml = """
<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
     xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">
  <w:hyperlink r:id="rIdLink"><w:r><w:t>Link text</w:t></w:r></w:hyperlink>
  <m:oMath><m:r><m:t>x=1</m:t></m:r></m:oMath>
</w:p>"""
    with ZipFile(io.BytesIO(buffer.getvalue())) as source, ZipFile(path, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                payload = payload.replace(b"</w:body>", paragraph_xml.encode("utf-8") + b"</w:body>")
            if info.filename != "word/_rels/document.xml.rels":
                target.writestr(info, payload)
        target.writestr("word/_rels/document.xml.rels", rels)


class TemplateV2InspectorTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.path = Path(self.tmp.name) / "core.docx"
        make_docx_with_core_objects(self.path)

    def test_inspector_does_not_modify_docx(self):
        before = file_hash(self.path)
        report = DocumentInspector(self.path).inspect()
        self.assertEqual(file_hash(self.path), before)
        self.assertIsNone(report.semantic_roles)

    def test_document_flow_preserves_paragraph_table_order(self):
        report = DocumentInspector(self.path).inspect()
        flow = [(block.kind, block.text_preview) for block in report.flow[:3]]
        self.assertEqual(flow[0], ("paragraph", "Before table"))
        self.assertEqual(flow[1][0], "table")
        self.assertEqual(flow[2], ("paragraph", "After table"))

    def test_detects_tables_merged_cells_hyperlinks_formulas_sections_and_stories(self):
        report = DocumentInspector(self.path).inspect()
        self.assertEqual(len(report.tables), 1)
        self.assertTrue(report.tables[0].merge_cells)
        self.assertGreaterEqual(report.tables[0].logical_column_count, 2)
        self.assertEqual(len(report.hyperlinks), 1)
        self.assertEqual(report.hyperlinks[0].target, "https://example.test")
        self.assertEqual(len(report.formulas), 1)
        self.assertGreaterEqual(len(report.sections), 2)
        self.assertGreaterEqual(len(report.headers), 1)
        self.assertGreaterEqual(len(report.footers), 1)

    def test_fingerprint_is_stable(self):
        first = DocumentInspector(self.path).inspect().fingerprint
        second = DocumentInspector(self.path).inspect().fingerprint
        self.assertEqual(first, second)

    def test_role_classifier_is_separate_from_raw_inspector(self):
        report = DocumentInspector(self.path).inspect()
        roles = RoleClassifierV2(use_ai=False).classify(report)
        self.assertEqual(roles.provider, "v2-context-rules")
        self.assertGreater(sum(roles.role_counts.values()), 0)

    @override_settings(AI_BASE_URL="http://192.0.2.10:8088/v1")
    def test_ai_settings_do_not_make_pass13_call_remote_model(self):
        report = DocumentInspector(self.path).inspect()
        roles = RoleClassifierV2().classify(report)
        self.assertEqual(roles.provider, "v2-context-rules")
        self.assertFalse(roles.warnings)

    def test_role_classifier_does_not_confuse_fig_sentence_with_caption_or_list(self):
        document = Document()
        document.add_paragraph("A Practical Article Title")
        document.add_paragraph("Ivan I. Author")
        document.add_paragraph("Abstract. This paper describes a reproducible experiment with enough body context.")
        document.add_paragraph("Keywords: experiment, samples, analysis")
        document.add_paragraph("1. Introduction")
        document.add_paragraph("Fig. 4 presents the results of the experiment and should stay body text.")
        document.add_paragraph("Sample preparation method")
        document.add_paragraph("Fig. 4. TG curves of the samples")
        path = Path(self.tmp.name) / "roles.docx"
        document.save(path)
        report = DocumentInspector(path).inspect()
        roles = RoleClassifierV2(use_ai=False).classify(report)
        by_text = {
            paragraph.normalized_text: role
            for paragraph, role in zip([p for p in report.paragraphs if p.normalized_text], roles.block_roles)
        }
        self.assertEqual(by_text["Fig. 4 presents the results of the experiment and should stay body text."]["role_hint"], "body")
        self.assertNotEqual(by_text["Sample preparation method"]["role_hint"], "list_item")
        self.assertEqual(by_text["Fig. 4. TG curves of the samples"]["role_hint"], "figure_caption")

    def test_template_profile_and_mapping_are_independent_from_content_diff(self):
        article_report = DocumentInspector(self.path).inspect()
        template_report = DocumentInspector(self.path).inspect()
        classifier = RoleClassifierV2(use_ai=False)
        article_structure = classifier.article_structure(article_report)
        template_profile = TemplateProfileBuilder(classifier=classifier).build(template_report)
        mapping = RoleMatcher().build_preview(article_structure, template_profile)
        self.assertTrue(template_profile.roles)
        self.assertGreater(mapping.summary["total_mappings"], 0)
        self.assertIn("analysed independently", " ".join(mapping.warnings))

    def test_safe_word_editor_preserves_article_objects(self):
        output = Path(self.tmp.name) / "result.docx"
        source = DocumentInspector(self.path).inspect()
        template = DocumentInspector(self.path).inspect()
        classifier = RoleClassifierV2(use_ai=False)
        structure = classifier.article_structure(source)
        profile = TemplateProfileBuilder(classifier=classifier).build(template)
        mapping = RoleMatcher().build_preview(structure, profile)
        result = SafeWordEditor(classifier=classifier).render(
            article_path=self.path,
            template_path=self.path,
            output_path=output,
            article_report=source,
            template_report=template,
            article_structure=structure,
            template_profile=profile,
            mapping_preview=mapping,
        )
        self.assertTrue(output.exists())
        self.assertIn("Preserved ARTICLE", " ".join(result.changes))
        rendered = DocumentInspector(output).inspect()
        self.assertEqual(rendered.fingerprint["table_count"], source.fingerprint["table_count"])
        self.assertEqual(rendered.fingerprint["drawing_count"], source.fingerprint["drawing_count"])
        self.assertEqual(rendered.fingerprint["formula_count"], source.fingerprint["formula_count"])

    def test_safe_word_editor_does_not_copy_template_footer_content_by_default(self):
        article = Document()
        article.add_paragraph("Article title")
        article.sections[0].footer.paragraphs[0].text = "ARTICLE FOOTER"
        article_path = Path(self.tmp.name) / "article-footer.docx"
        article.save(article_path)
        template = Document()
        template.add_paragraph("Template title")
        template.sections[0].footer.paragraphs[0].text = "TEMPLATE FOOTER"
        template_path = Path(self.tmp.name) / "template-footer.docx"
        template.save(template_path)
        output = Path(self.tmp.name) / "footer-result.docx"
        SafeWordEditor(classifier=RoleClassifierV2(use_ai=False)).render(
            article_path=article_path,
            template_path=template_path,
            output_path=output,
            copy_template_headers=False,
        )
        rendered = DocumentInspector(output).inspect()
        footer_text = " ".join(item.text for item in rendered.footers)
        self.assertIn("ARTICLE FOOTER", footer_text)
        self.assertNotIn("TEMPLATE FOOTER", footer_text)

    def test_safe_word_editor_replaces_story_relationships_without_duplicate_targets(self):
        article = Document()
        article.add_paragraph("A practical article title")
        article.add_paragraph("© Ivan I. Author")
        article.add_paragraph("Abstract. This paper contains enough text for deterministic role classification.")
        article.sections[0].footer.paragraphs[0].text = "ARTICLE FOOTER"
        article_path = Path(self.tmp.name) / "article-story.docx"
        article.save(article_path)

        template = Document()
        template.add_paragraph("A formatted article title")
        template.sections[0].header.paragraphs[0].text = "JOURNAL HEADER"
        template.sections[0].footer.paragraphs[0].text = "1"
        template.sections[0].footer.add_paragraph("TEMPLATE AUTHOR")
        template_path = Path(self.tmp.name) / "template-story.docx"
        template.save(template_path)

        output = Path(self.tmp.name) / "story-result.docx"
        SafeWordEditor(classifier=RoleClassifierV2(use_ai=False)).render(
            article_path=article_path,
            template_path=template_path,
            output_path=output,
            copy_template_headers=True,
        )

        rendered = DocumentInspector(output).inspect()
        self.assertIn("JOURNAL HEADER", " ".join(item.text for item in rendered.headers))
        footer_text = " ".join(item.text for item in rendered.footers)
        self.assertIn("Author", footer_text)
        self.assertNotIn("TEMPLATE AUTHOR", footer_text)

        with ZipFile(output) as archive:
            rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        story_relationships = [
            (item.get("Type"), item.get("Target"))
            for item in rels
            if item.get("Type", "").endswith(("/header", "/footer"))
        ]
        self.assertEqual(len(story_relationships), len(set(story_relationships)))

    def test_qwen_like_planner_starts_large_full_width_figures_on_fresh_pages(self):
        snapshot = {
            "article": {
                "blocks": [],
                "tables": [
                    {
                        "id": "block_0042",
                        "classification": "FIGURE_CONTAINER",
                        "rows": 4,
                        "cols": 2,
                        "grid": ["4932", "4932"],
                    }
                ],
            },
            "template": {
                "paragraphs": [],
                "front_sequence": [],
                "front_language_order": [],
                "roles": {},
                "layout": {},
            },
        }
        plan = QwenLikePlanningEngine()._local_plan(snapshot)
        self.assertTrue(plan.flow["page_break_before_large_full_width_figures"])
        self.assertEqual(plan.flow["large_figure_block_ids"], ["block_0042"])

    def test_qwen_planner_accepts_only_whitelisted_values_and_existing_ids(self):
        snapshot = {
            "article": {
                "blocks": [{"id": "block_0001"}],
                "tables": [{"id": "block_0042"}],
            },
            "template": {
                "front_sequence": ["title", "author"],
                "front_language_order": ["ru", "en"],
            },
        }
        local = QwenLikePlanningEngine()._local_plan({
            "article": {"blocks": [], "tables": []},
            "template": {
                "paragraphs": [], "front_sequence": ["title", "author"],
                "front_language_order": ["ru", "en"], "roles": {}, "layout": {},
            },
        })
        patch_payload = {
            "front": {"citation_expected_lines": 99, "language_order": ["en", "xx"]},
            "flow": {
                "large_figure_block_ids": ["block_0042", "invented"],
                "float_lead_after_figure_block_ids": ["block_0042"],
                "max_float_body_blocks": 50,
                "arbitrary_xml": "<w:del/>",
            },
            "replacement_text": "must be ignored",
        }
        result = QwenLikePlanningEngine()._merge_provider(local, patch_payload, snapshot)
        self.assertEqual(result.front["citation_expected_lines"], 4)
        self.assertEqual(result.front["language_order"], ["ru", "en"])
        self.assertEqual(result.flow["large_figure_block_ids"], ["block_0042"])
        self.assertEqual(result.flow["float_lead_after_figure_block_ids"], ["block_0042"])
        self.assertEqual(result.flow["max_float_body_blocks"], 3)
        self.assertNotIn("arbitrary_xml", result.flow)
        self.assertTrue(result.warnings)

    @override_settings(
        AI_BASE_URL="http://192.0.2.10:1234/v1",
        TEMPLATE_V2_QWEN_ENABLED=False,
        TEMPLATE_V2_QWEN_MODEL="qwen-test",
    )
    def test_qwen_planner_is_feature_gated_even_when_ai_endpoint_exists(self):
        planner = build_planning_engine()
        self.assertIsNone(planner.provider)

    def test_qwen_planner_provider_failure_keeps_local_plan(self):
        report = DocumentInspector(self.path).inspect()
        classifier = RoleClassifierV2(use_ai=False)
        structure = classifier.article_structure(report)
        profile = TemplateProfileBuilder(classifier=classifier).build(report)

        def unavailable(_snapshot):
            raise RuntimeError("vpn unavailable")

        result = QwenLikePlanningEngine(provider=unavailable).plan(
            article_report=report,
            template_report=report,
            article_structure=structure,
            template_profile=profile,
        )
        self.assertEqual(result.provider, "qwen-like-local-simulator")
        self.assertIn("vpn unavailable", " ".join(result.warnings))

    def test_qwen_provider_uses_one_model_and_compacts_ordinary_body_text(self):
        provider = QwenPlanningProvider(model="qwen-only", timeout=17)
        snapshot = {
            "article": {
                "blocks": [
                    {"id": "block_1", "kind": "paragraph", "role": "body", "zone": "body", "text": "ordinary prose"},
                    {"id": "block_2", "kind": "paragraph", "role": "body", "zone": "body", "text": "Fig. 2 presents results"},
                ],
                "tables": [],
            },
            "template": {"paragraphs": [], "roles": {}, "layout": {}},
        }
        response = {"candidates": [{"content": {"parts": [{"text": '{"flow": {}}'}]}}]}
        with patch(
            "apps.template_workspace.v2.planning.qwen_provider.generate_content",
            return_value=(response, "qwen-only"),
        ) as generate:
            self.assertEqual(provider(snapshot), {"flow": {}})
        kwargs = generate.call_args.kwargs
        self.assertEqual(kwargs["models"], [{"id": "qwen-only"}])
        sent = generate.call_args.args[0]["contents"][0]["parts"][0]["text"]
        self.assertNotIn("ordinary prose", sent)
        self.assertIn("Fig. 2 presents results", sent)


@skipUnlessDBFeature("supports_transactions")
class TemplateV2ViewTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        self.override = override_settings(MEDIA_ROOT=self.media.name)
        self.override.enable()
        self.addCleanup(self.override.disable)
        self.user = get_user_model().objects.create_user(username="v2_author")
        self.client.force_login(self.user)

    def test_v2_route_is_separate_from_legacy_route(self):
        response = self.client.get(reverse("template_workspace:v2_workspace"))
        self.assertContains(response, "Word-forensics V2")
        self.assertContains(response, "Шаблон два")
        legacy = self.client.get(reverse("template_workspace:workspace"))
        self.assertContains(legacy, "Статья по вашему шаблону")

    def test_run_v2_job_writes_json_reports(self):
        source = Path(self.media.name) / "source.docx"
        template = Path(self.media.name) / "template.docx"
        make_docx_with_core_objects(source)
        make_docx_with_core_objects(template)
        job = TemplateJob.objects.create(
            owner=self.user,
            kind="v2",
            article=SimpleUploadedFile("source.docx", source.read_bytes()),
            template=SimpleUploadedFile("template.docx", template.read_bytes()),
            article_name="source.docx",
            template_name="template.docx",
        )
        run_v2_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, "completed")
        self.assertTrue((analysis_directory(job) / "article_report.json").exists())
        self.assertTrue((analysis_directory(job) / "template_report.json").exists())
        self.assertTrue((analysis_directory(job) / "article_structure.json").exists())
        self.assertTrue((analysis_directory(job) / "template_profile.json").exists())
        self.assertTrue((analysis_directory(job) / "mapping_preview.json").exists())
        self.assertTrue((analysis_directory(job) / "planning_report.json").exists())
        self.assertTrue((analysis_directory(job) / "editor_report.json").exists())
        self.assertTrue(result_docx_path(job).exists())
        self.assertFalse((analysis_directory(job) / "document_diff.json").exists())

    def test_run_v2_job_converts_legacy_doc_uploads(self):
        source = Path(self.media.name) / "source.docx"
        template = Path(self.media.name) / "template.doc"
        make_docx_with_core_objects(source)
        template.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1legacy")
        job = TemplateJob.objects.create(
            owner=self.user,
            kind="v2",
            article=SimpleUploadedFile("source.docx", source.read_bytes()),
            template=SimpleUploadedFile("template.doc", template.read_bytes()),
            article_name="source.docx",
            template_name="template.doc",
        )
        with patch("apps.template_workspace.v2.services.convert_legacy_doc_to_docx", return_value=source.read_bytes()):
            run_v2_job(str(job.pk))
        job.refresh_from_db()
        self.assertEqual(job.status, "completed")
        self.assertTrue((analysis_directory(job) / "converted" / "template.docx").exists())
        self.assertTrue(result_docx_path(job).exists())


@skipUnlessDBFeature("supports_transactions")
class TemplateV2RealDocumentTests(TestCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        if not REAL_FIXTURES_AVAILABLE:
            raise unittest.SkipTest("Real DOCX fixtures are not available in Downloads.")

    def test_real_balabanov_pair_shows_layout_differences(self):
        source = DocumentInspector(FIXTURES["balabanov_source"]).inspect()
        formatted = DocumentInspector(FIXTURES["balabanov_formatted"]).inspect()
        diff = DocumentDiffBuilder().compare(source, formatted)
        self.assertEqual(source.fingerprint["table_count"], formatted.fingerprint["table_count"])
        self.assertGreater(formatted.fingerprint["section_count"], source.fingerprint["section_count"])
        self.assertGreater(formatted.fingerprint["header_count"], source.fingerprint["header_count"])
        self.assertTrue(diff.layout)

    def test_tyutyunnik_uses_same_generic_inspector(self):
        report = DocumentInspector(FIXTURES["tyutyunnik_formatted"]).inspect()
        self.assertGreater(report.fingerprint["section_count"], 1)
        self.assertGreater(report.fingerprint["drawing_count"], 0)
        self.assertGreater(report.fingerprint["header_count"], 0)
        self.assertIsNone(report.semantic_roles)

    def test_tyutyunnik_template_profile_has_layout_profile(self):
        report = DocumentInspector(FIXTURES["tyutyunnik_formatted"]).inspect()
        profile = TemplateProfileBuilder(classifier=RoleClassifierV2(use_ai=False)).build(report)
        self.assertGreaterEqual(profile.layout.default_body_column_count, 1)
        self.assertGreater(len(profile.layout.section_ranges), 0)
        self.assertIsInstance(profile.roles, dict)
