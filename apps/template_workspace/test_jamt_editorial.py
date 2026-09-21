from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from .v2.editorial import (annotate_manuscript, suspicious_spans, validate_model_findings,
                           highlight_spans, ImageOnlyManuscript)
from .v2.classification.roles import RoleClassifierV2
from .v2.inspector.document import DocumentInspector
from .v2.ooxml.namespaces import NS, qn


class EditorialMarksTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.source=self.root/'source.docx';self.result=self.root/'reviewed.docx'

    def draft(self):
        d=Document()
        for text in ['Thermal behaviour of porous materials','Ivan P. Petrov','Materials Research University, Moscow, Russia',
                     'petrov@example.org','Abstract: Porous materials were tested at constant temperature.',
                     'Keywords: porosity; heat transfer','1. Introduction',
                     'We measured the the values. This sentence contains a damaged symbol \ufffd.',
                     'В исследовании проверена прoчность образцов.', '2. Methods',
                     'The value is 12.5 ± 0.3 MPa and the alloy is Ti6Al4V.']:
            d.add_paragraph(text)
        d.save(self.source);return d

    def test_missing_slots_are_highlighted_and_original_text_is_unchanged(self):
        self.draft();report=annotate_manuscript(self.source,self.result)
        self.assertTrue(report['original_text_unchanged'])
        fields={x['field'] for x in report['missing_fields']}
        self.assertTrue({'doi','issue','received','accepted','published','title_ru','abstract_ru','citation_en'}<=fields)
        with ZipFile(self.result) as z:
            root=etree.fromstring(z.read('word/document.xml'))
            placeholders=root.xpath('//w:p[w:bookmarkStart[starts-with(@w:name,"JAMT_missing_")]]',namespaces=NS)
            self.assertEqual(len(placeholders),report['missing_count'])
            for p in placeholders:self.assertTrue(p.xpath('.//w:highlight[@w:val="yellow"]',namespaces=NS))
            self.assertIn('12.5 ± 0.3 MPa',''.join(root.xpath('//w:t/text()',namespaces=NS)))
        structure=RoleClassifierV2().article_structure(DocumentInspector(self.result).inspect())
        slot=next(b for b in structure.blocks if 'MISSING: article title' in b['text_preview']) if 'title_en' in fields else None
        if slot:self.assertEqual((slot['detected_role'],slot['language']),('title','en'))
        russian=next(b for b in structure.blocks if '[НЕ УКАЗАНО: название' in b['text_preview'])
        self.assertEqual((russian['detected_role'],russian['language'],russian['zone']),('title','ru','front_matter'))

    def test_repeated_processing_does_not_duplicate_missing_fields(self):
        self.draft();first=annotate_manuscript(self.source,self.result)
        second=self.root/'twice.docx';again=annotate_manuscript(self.result,second)
        with ZipFile(self.result) as a, ZipFile(second) as b:
            one=etree.fromstring(a.read('word/document.xml'));two=etree.fromstring(b.read('word/document.xml'))
            self.assertEqual(one.xpath('//w:bookmarkStart/@w:name',namespaces=NS),two.xpath('//w:bookmarkStart/@w:name',namespaces=NS))
        self.assertEqual(first['missing_count'],again['missing_count'])

    def test_reference_doi_does_not_satisfy_article_identifier(self):
        d=self.draft();d.add_paragraph('References');d.add_paragraph('1. Other work. DOI: 10.1234/other.2024');d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        self.assertIn('doi',[i['field'] for i in report['missing_fields']])

    def test_native_math_and_hyperlinks_survive_run_splitting(self):
        d=self.draft();p=d.add_paragraph();p.add_run('Start pr').italic=True
        run=p.add_run('oчность middle \ufffd end');run.bold=True
        math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x=12.5';r.append(t);math.append(r);p._p.append(math)
        link=OxmlElement('w:hyperlink')
        from docx.opc.constants import RELATIONSHIP_TYPE
        rid=d.part.relate_to('https://example.org/paper',RELATIONSHIP_TYPE.HYPERLINK,is_external=True)
        link.set(qn('r:id'),rid);lr=OxmlElement('w:r');lt=OxmlElement('w:t');lt.text='link \ufffd';lr.append(lt);link.append(lr);p._p.append(link)
        before=etree.tostring(math);d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        with ZipFile(self.result) as z:
            root=etree.fromstring(z.read('word/document.xml'));found=root.xpath('//m:oMath',namespaces=NS)[0]
            self.assertEqual(etree.tostring(found),before)
            self.assertEqual(root.xpath('//w:hyperlink/@r:id',namespaces=NS),[rid])
            self.assertEqual(''.join(root.xpath('//w:hyperlink//w:t/text()',namespaces=NS)),'link \ufffd')
            self.assertIn('https://example.org/paper',z.read('word/_rels/document.xml.rels').decode())
        self.assertTrue(report['source_binary_parts_unchanged'])

    def test_exact_spans_can_cross_multiple_text_runs(self):
        d=Document();p=d.add_paragraph();p.add_run('на пр');p.add_run('oчность и ещё');original=p.text
        highlight_spans(p._p,[(3,11)])
        self.assertEqual(p.text,original)
        marked=p._p.xpath('.//w:r[w:rPr/w:highlight]/w:t/text()')
        self.assertEqual(''.join(marked),'прoчност')

    def test_model_cannot_rewrite_or_mark_ambiguous_or_unsupplied_text(self):
        records=[{'id':'p1','text':'A typo occurs once. Word Word.'}]
        response={'issues':[{'id':'p1','quote':'typo','description':'Опечатка','code':'spelling'},
                            {'id':'p1','quote':'corrected text','description':'bad','code':'spelling'},
                            {'id':'p1','quote':'Word','description':'ambiguous','code':'spelling'},
                            {'id':'p2','quote':'typo','description':'wrong id','code':'spelling'}]}
        accepted,rejected=validate_model_findings(response,records)
        self.assertEqual(len(accepted),1);self.assertEqual(rejected,3)

    def test_model_failure_retains_deterministic_marks(self):
        self.draft()
        def fail(*args):raise TimeoutError()
        report=annotate_manuscript(self.source,self.result,reviewer=fail)
        self.assertEqual(report['ai']['status'],'unavailable')
        self.assertIn('suspicious_character',[i['code'] for i in report['issues']])
        self.assertTrue(self.result.exists())

    def test_only_suspicious_alphabet_mixture_is_marked_not_scientific_notation(self):
        self.assertEqual(suspicious_spans('Ti6Al4V β-SiC 12.5 ± 0.3 MPa; 20.09.2026'),[])
        self.assertIn('mixed_alphabets',[i['code'] for i in suspicious_spans('Измерена прoчность образца.')])

    def test_image_only_upload_gets_clear_error_instead_of_fake_editing(self):
        Document().save(self.source)
        with self.assertRaises(ImageOnlyManuscript):annotate_manuscript(self.source,self.result)

    def test_manuscript_starting_with_introduction_has_no_fake_title(self):
        d=Document();d.add_paragraph('1. Introduction');d.add_paragraph('The experiment used three samples.');d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        self.assertIn('title_en',[i['field'] for i in report['missing_fields']])

    def test_rubric_is_not_title_and_full_author_names_are_recognised(self):
        d=self.draft()
        first=d.paragraphs[0]._p
        for text in ('Original papers','Advanced structural materials, materials for extreme conditions'):
            p=OxmlElement('w:p');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=text;r.append(t);p.append(r);first.addprevious(p)
        d.paragraphs[3].text='Ivan P. Petrov, Anna S. Sokolova'
        d.save(self.source);report=annotate_manuscript(self.source,self.result)
        fields={x['field'] for x in report['missing_fields']}
        self.assertNotIn('title_en',fields);self.assertNotIn('author_en',fields)
        structure=RoleClassifierV2().article_structure(DocumentInspector(self.source).inspect())
        self.assertEqual(structure.blocks[1]['detected_role'],'rubric')

    def test_empty_existing_field_gets_prompt_at_its_location(self):
        d=self.draft();d.paragraphs[5].text='Keywords:';d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        values=[p.text for p in Document(self.result).paragraphs if p.text.startswith('Keywords:')]
        self.assertEqual(len(values),1);self.assertIn('[MISSING:',values[0])
        self.assertTrue(report['original_text_unchanged'])
        self.assertIn('keywords_en',[x['field'] for x in report['missing_fields']])

    def test_invalid_existing_date_is_marked_without_duplicate_received_slot(self):
        d=self.draft();d.add_paragraph('Received 31.02.2026');d.add_paragraph('References');d.add_paragraph('1. Test reference.');d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        self.assertNotIn('received',[x['field'] for x in report['missing_fields']])
        self.assertIn('invalid_date',[x['code'] for x in report['issues']])
        self.assertIn('31.02.2026',' '.join(p.text for p in Document(self.result).paragraphs))

    def test_month_first_date_and_empty_date_label_keep_their_locations(self):
        d=self.draft();d.add_paragraph('Received: September 20, 2026');d.add_paragraph('Accepted:');d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        self.assertNotIn('received',[x['field'] for x in report['missing_fields']])
        lines=[p.text for p in Document(self.result).paragraphs if p.text.startswith('Accepted:')]
        self.assertEqual(len(lines),1);self.assertIn('[MISSING:',lines[0])
        self.assertTrue(report['original_text_unchanged'])

    def test_invalid_doi_is_highlighted_and_not_corrected(self):
        self.assertIn('invalid_doi',[i['code'] for i in suspicious_spans('DOI: 1O.1234/paper')])
        self.assertEqual(suspicious_spans('DOI: https://doi.org/10.1234/paper'),[])
        d=self.draft();p=d.add_paragraph('DOI: 1O.1234/paper');d.paragraphs[0]._p.addprevious(p._p);d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        self.assertNotIn('doi',[x['field'] for x in report['missing_fields']])
        self.assertIn('invalid_doi',[i['code'] for i in report['issues']])

    def test_editorial_placeholder_is_not_treated_as_a_finished_doi(self):
        value='DOI: 10.17277/jamt.2026.02.pp.000-000 (предоставляется редакцией)'
        d=self.draft();p=d.add_paragraph(value);d.paragraphs[0]._p.addprevious(p._p);d.save(self.source)
        report=annotate_manuscript(self.source,self.result)
        field=next(x for x in report['missing_fields'] if x['field']=='doi')
        self.assertTrue(field['existing']);self.assertTrue(report['original_text_unchanged'])
        self.assertEqual(sum(p.text.startswith('DOI:') for p in Document(self.result).paragraphs),1)
        self.assertIn('unresolved_placeholder',[i['code'] for i in report['issues']])
        self.assertTrue(suspicious_spans('For citation: данные предоставляются редакцией.'))

    def test_heading_after_figure_caption_is_not_absorbed_and_bio_is_prose(self):
        d=self.draft()
        for text in ('Fig. 1. Measured values.','Рис. 1. Измеренные значения.','4. Conclusions',
                     'This paragraph explains the conclusions.','References','1. A source.',
                     'Information about the authors','Ivan P. Petrov — researcher; petrov@example.org.'):
            d.add_paragraph(text)
        d.save(self.source)
        blocks=RoleClassifierV2().article_structure(DocumentInspector(self.source).inspect()).blocks
        self.assertEqual(next(b for b in blocks if b['text_preview']=='4. Conclusions')['detected_role'],'heading_1')
        self.assertEqual(next(b for b in blocks if b['text_preview'].startswith('Ivan P. Petrov —'))['detected_role'],'author_bio')

    def test_table_reference_sentence_is_body_not_a_caption(self):
        d=self.draft();d.add_paragraph('Table 1 contains the measured values and their uncertainty.');d.save(self.source)
        blocks=RoleClassifierV2().article_structure(DocumentInspector(self.source).inspect()).blocks
        self.assertEqual(next(b for b in blocks if b['text_preview'].startswith('Table 1 contains'))['detected_role'],'body')
