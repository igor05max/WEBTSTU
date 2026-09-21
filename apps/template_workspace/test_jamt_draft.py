"""Regressions from raw manuscripts, distinct from finished reference articles."""
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

from django.test import SimpleTestCase
from docx import Document
from docx.shared import Pt
from lxml import etree
from PIL import Image

from .v2.classification.roles import RoleClassifierV2
from .v2.editor.safe_word_editor import SafeWordEditor
from .v2.editor.repair_actions import structural_issues
from .v2.editor.object_flow import vml_dimensions
from .v2.inspector.document import DocumentInspector
from .v2.ooxml.namespaces import NS, qn
from .v2.styles.jamt import prepare_jamt_style


def picture(doc):
    stream = BytesIO()
    Image.new('RGB', (40, 50), 'blue').save(stream, format='PNG')
    stream.seek(0)
    doc.add_picture(stream, width=Pt(110))


def bilingual_front(doc):
    for title, abstract, keywords, citation in [
        ('Исследование новых материалов', 'Аннотация', 'Ключевые слова: материал', 'Для цитирования: данные редакции'),
        ('A study of new materials', 'Abstract', 'Keywords: materials', 'For citation: editorial data'),
    ]:
        doc.add_paragraph(title, style='Title')
        doc.add_paragraph('© John Michael Smith')
        doc.add_paragraph('Technical University,')
        doc.add_paragraph('Research Street, 12, City')
        doc.add_paragraph('smith@example.test')
        doc.add_paragraph(abstract)
        doc.add_paragraph('Measurements are 42 ± 0.2 MPa; scientific content stays intact.')
        doc.add_paragraph(keywords)
        doc.add_paragraph(citation)


class DraftLayoutTests(SimpleTestCase):
    def setUp(self):
        self.tmp = TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def render(self, doc):
        source = self.root/'source.docx'; doc.save(source)
        report = DocumentInspector(source).inspect()
        structure = RoleClassifierV2(use_ai=False).article_structure(report)
        carrier, template, profile = prepare_jamt_style(self.root, article_report=report, article_structure=structure)
        result = SafeWordEditor().render(article_path=source, article_report=report, article_structure=structure,
            template_path=carrier, template_report=template, template_profile=profile, output_path=self.root/'result.docx')
        self.assertTrue(result.metrics['native_integrity']['passed'])
        with ZipFile(source) as before, ZipFile(self.root/'result.docx') as after:
            for name in before.namelist():
                if name.startswith(('word/media/', 'word/embeddings/')):
                    self.assertEqual(before.read(name), after.read(name))
        return Document(self.root/'result.docx'), result

    def test_inherited_list_numbers_and_bullets_survive_word_formatting(self):
        from paper_formatter.latex_lab.bridge import NativeBridge
        doc=Document();bilingual_front(doc);doc.add_paragraph('1. Introduction')
        for text in ('First operation.','Second operation.','Third operation.'):
            doc.add_paragraph(text,style='List Number')
        for text in ('First observation.','Second observation.'):
            doc.add_paragraph(text,style='List Bullet')
        self.render(doc)
        _,before=NativeBridge(self.root/'source.docx',self.root/'before-tex').build()
        _,after=NativeBridge(self.root/'result.docx',self.root/'after-tex').build()
        self.assertEqual([x['label'] for x in before['list_labels']],['1.','2.','3.','•','•'])
        self.assertEqual([(x['label'],x['text_anchor']) for x in before['list_labels']],
                         [(x['label'],x['text_anchor']) for x in after['list_labels']])

    def test_leading_body_picture_does_not_block_bilingual_front_layout(self):
        doc = Document(); bilingual_front(doc)
        picture(doc)
        doc.add_paragraph('Fig 1. Original portrait')
        doc.add_paragraph('The first body paragraph describes the original portrait.')
        output, result = self.render(doc)
        text = '\n'.join(output._element.xpath('.//w:t/text()'))
        self.assertLess(text.index('A study of'), text.index('Исследование'))
        self.assertLess(text.index('Исследование'), text.index('Fig 1.'))
        self.assertIn('Technical University, Research Street, 12, City', [p.text for p in output.paragraphs])
        self.assertGreater(result.metrics['front_blocks_reordered'], 0)
        self.assertGreaterEqual(result.metrics['front_paragraphs_merged'], 4)
        self.assertEqual(text.count('Measurements are 42 ± 0.2 MPa; scientific content stays intact.'), 2)
        self.assertTrue(output._element.xpath('.//w:tbl[.//w:drawing][.//w:t[contains(., "Original portrait")]]'))
        self.assertTrue(any(p.text.startswith('Abstract Measurements') for p in output.paragraphs))

    def test_wide_picture_retains_both_caption_languages_in_the_same_span(self):
        doc = Document(); bilingual_front(doc)
        doc.add_paragraph('1. Introduction', style='Heading 1')
        picture(doc)
        for extent in doc._element.xpath('.//wp:extent'):
            extent.set('cx', str(470*12700))
        doc.add_paragraph('Fig 5. Стенд в музее')
        doc.add_paragraph('A stand made for the museum')
        doc.add_paragraph('This long paragraph is ordinary scientific prose describing the original measurements, their units, uncertainties and independent observations.')
        output, _ = self.render(doc)
        groups = output._element.xpath('.//w:tbl[.//w:drawing][.//w:t[contains(., "Стенд")]]')
        self.assertEqual(len(groups), 1)
        self.assertIn('A stand made for the museum', ''.join(groups[0].xpath('.//w:t/text()')))

    def test_genuine_front_logo_still_prevents_unsafe_sorting(self):
        doc = Document(); picture(doc); bilingual_front(doc)
        doc.add_paragraph('1. Introduction', style='Heading 1')
        doc.add_paragraph('Original text after the front logo.')
        output, result = self.render(doc)
        self.assertEqual(result.metrics['front_blocks_reordered'], 0)
        self.assertTrue(output._element.body[0].xpath('.//w:drawing'))

    def test_bilingual_rubrics_have_one_inter_block_gap(self):
        doc = Document()
        doc.add_paragraph('Нобелистика'); doc.add_paragraph('Nobelistics')
        bilingual_front(doc)
        doc.add_paragraph('1. Introduction', style='Heading 1')
        doc.add_paragraph('Original scientific measurements remain unchanged.')
        output, _ = self.render(doc)
        rubrics = [p for p in output.paragraphs if p.text in {'Нобелистика', 'Nobelistics'}]
        self.assertEqual(len(rubrics), 2)
        self.assertEqual(rubrics[0]._p.xpath('./w:pPr/w:spacing/@w:after'), ['0'])
        self.assertEqual(rubrics[1]._p.xpath('./w:pPr/w:spacing/@w:after'), ['230'])

    def test_native_caption_box_uses_caption_typography_without_touching_license(self):
        doc = Document(); bilingual_front(doc)
        doc.add_paragraph('1. Introduction', style='Heading 1')
        p = doc.add_paragraph('The source measurements remain 42 ± 0.2 MPa.')
        for target, text, size in [(p, 'Fig. 2. Original book caption', 24),
                                   (doc.add_paragraph(), 'Copyright: Original license.', 16)]:
            pict = etree.SubElement(target.add_run()._r, qn('w:pict'))
            shape = etree.SubElement(pict, qn('v:shape'), style='width:130pt;height:50pt')
            box = etree.SubElement(etree.SubElement(shape, qn('v:textbox')), qn('w:txbxContent'))
            inner = etree.SubElement(box, qn('w:p')); run = etree.SubElement(inner, qn('w:r'))
            etree.SubElement(etree.SubElement(run, qn('w:rPr')), qn('w:sz')).set(qn('w:val'), str(size))
            etree.SubElement(run, qn('w:t')).text = text
        output, result = self.render(doc)
        boxes = output._element.xpath('.//w:txbxContent')
        for box in boxes:
            expected = '20' if 'Fig.' in ''.join(box.xpath('.//w:t/text()', namespaces=NS)) else '16'
            self.assertEqual(set(box.xpath('.//w:rPr/w:sz/@w:val', namespaces=NS)), {expected})
            frame_width = vml_dimensions(box.getparent().getparent())[1]
            if expected == '20':
                self.assertGreater(frame_width, 230)
            else:
                self.assertEqual(frame_width, 130)
        self.assertEqual(result.metrics['display_objects_detached'], 1)
        caption = next(box for box in boxes if 'Fig.' in ''.join(box.xpath('.//w:t/text()', namespaces=NS)))
        caption.find('.//w:rPr/w:sz', NS).set(qn('w:val'), '28')
        damaged = self.root/'damaged.docx'; output.save(damaged)
        issues = structural_issues(damaged, result.metrics['layout_targets'])
        self.assertTrue(any(issue['kind'] == 'text_typography' for issue in issues))

    def test_caption_without_abbreviation_dot_keeps_its_translation(self):
        doc = Document(); bilingual_front(doc)
        doc.add_paragraph('1. Introduction', style='Heading 1')
        for text in ['Fig 5. Стенд в музее', 'A stand made for the museum', 'Fig 5 presents the measured values.']:
            p = doc.add_paragraph(text); p.runs[0].font.size = Pt(12)
        source = self.root/'captions.docx'; doc.save(source)
        blocks = RoleClassifierV2(use_ai=False).article_structure(DocumentInspector(source).inspect()).blocks
        captions = [b for b in blocks if b['text_preview'] in ['Fig 5. Стенд в музее', 'A stand made for the museum']]
        self.assertEqual([b['detected_role'] for b in captions], ['figure_caption']*2)
        self.assertEqual(captions[0]['group_id'], captions[1]['group_id'])
        prose = next(b for b in blocks if b['text_preview'].startswith('Fig 5 presents'))
        self.assertEqual(prose['detected_role'], 'body')

    def test_compiled_journal_header_keeps_reference_black_rules(self):
        doc = Document(); bilingual_front(doc)
        doc.add_paragraph('1. Introduction', style='Heading 1')
        doc.add_paragraph('Source manuscript body.')
        self.render(doc)
        with ZipFile(self.root/'result.docx') as archive:
            headers = [etree.fromstring(archive.read(n)) for n in archive.namelist()
                       if n.startswith('word/header') and n.endswith('.xml')]
        self.assertTrue(headers)
        for header in headers:
            self.assertEqual(set(header.xpath('.//w:rPr/w:color/@w:val|.//w:pBdr/*/@w:color', namespaces=NS)), {'000000'})
