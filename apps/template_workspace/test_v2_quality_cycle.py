from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase, override_settings
from docx import Document
from docx.shared import Pt
from lxml import etree
import pymupdf

from .v2.editor.safe_word_editor import (SafeWordEditorResult, _apply_run_profile, _NodeMeta,
                                       _normalise_front_matter, _resize_table_to_width, _wide_object_spans)
from .v2.editor.object_flow import detach_display_objects, normalize_display_geometry
from .v2.editor.repair_actions import (apply_actions, ground_targets, select_actions,
                                     structural_issues, inventory, rendered_target_ids)
from .v2.ooxml.namespaces import NS, qn
from .v2.quality_cycle import run_quality_cycle, improvement, quality_status, repairs_covered
from .v2.readability import validate_issues
from .v2.timeouts import job_timeout_seconds
from .v2.editor.layout_fidelity import wrap_picture_captions


def reviewed(issues=(), **kwargs):
    return dict(status='reviewed', pages_total=1, pages_checked=[1], pages_not_checked=[],
                issues=list(issues), warnings=[], **kwargs)


ISSUE = dict(page=1, kind='text_typography', severity='medium', description='Mixed font sizes',
             source='qwen-vision', target_id='layout_0001', action='restore_typography')


@override_settings(TEMPLATE_V2_EDIT_CYCLE_ENABLED=True, TEMPLATE_V2_EDIT_CYCLE_MAX_PASSES=3,
                   TEMPLATE_V2_EDIT_CYCLE_BUDGET=2100)
class QualityCycleTests(SimpleTestCase):
    def test_missing_rendered_caption_rejects_apparent_visual_improvement(self):
        before = reviewed([ISSUE], rendered_target_ids=['body', 'caption'])
        after = reviewed([], rendered_target_ids=['body'])
        self.assertEqual(improvement(before, after, [], []), (False, 'rendered_content_regression'))

    def test_rendered_text_across_pages_is_not_a_loss(self):
        pdf = [SimpleNamespace(get_text=lambda:'Figure 6. Several original '),
               SimpleNamespace(get_text=lambda:'panels and their common caption.')]
        targets = [{'id':'caption', 'text':'Figure 6. Several original panels and their common caption.'},
                   {'id':'missing', 'text':'Another caption absent from this rendered document.'}]
        self.assertEqual(rendered_target_ids(pdf, targets), ['caption'])

    def setUp(self):
        self.tmp = TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.path = self.root/'result.docx'
        doc = Document(); doc.add_paragraph('Original scientific finding: 12.5 ± 0.2 mg.').runs[0].font.size = Pt(20)
        doc.save(self.path)
        self.original = self.path.read_bytes()
        self.target = dict(id='layout_0001', path=[0, 0], role='body', text='Original scientific finding: 12.5 ± 0.2 mg.',
                           allowed_actions=['restore_typography'], run_format={'size':24}, paragraph_format={})
        self.editor = SafeWordEditorResult(str(self.path), metrics={'layout_targets':[self.target]})

    def run_cycle(self, reviewer):
        return run_quality_cycle(self.path, self.root, editor_result=self.editor, reviewer=reviewer)

    def test_repair_rerender_accept_and_promote_best(self):
        calls = []
        def review(path, folder, **kwargs):
            calls.append(Path(path))
            issues = structural_issues(path, [self.target])
            return reviewed([ISSUE] if issues else [])
        review, report = self.run_cycle(review)
        self.assertEqual(report['accepted_repairs'], 1)
        self.assertEqual(report['best_iteration'], 1)
        self.assertEqual(report['status'], 'checked')
        self.assertEqual(len(calls), 2)
        self.assertNotEqual(self.original, self.path.read_bytes())
        self.assertEqual((self.root/'quality_cycle/pass-0/candidate.docx').read_bytes(), self.original)
        with ZipFile(self.path) as z, ZipFile(self.root/'quality_cycle/pass-0/candidate.docx') as before:
            for name in before.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(z.read(name), before.read(name))
        self.assertIn('12.5 ± 0.2 mg.', Document(self.path).paragraphs[0].text)
        self.assertEqual(review['quality_status'], 'checked')

    def test_saved_style_is_used_for_baseline_and_repair_recheck(self):
        calls = []
        def reviewer(path, folder, **kwargs):
            calls.append(kwargs)
            return reviewed([ISSUE] if structural_issues(path, [self.target]) else [])
        _, report = run_quality_cycle(self.path, self.root, editor_result=self.editor,
                                     style_id='jamt', reviewer=reviewer)
        self.assertEqual(report['accepted_repairs'], 1)
        self.assertEqual(len(calls), 2)
        self.assertTrue(all(call['style_id'] == 'jamt' for call in calls))

    def test_visual_regression_rolls_back_even_if_font_fixed(self):
        results = iter([reviewed([ISSUE]), reviewed([dict(ISSUE, severity='high', kind='overlap')])])
        _, report = self.run_cycle(lambda *a, **k: next(results))
        self.assertEqual(report['accepted_repairs'], 0)
        self.assertEqual(report['rejected_candidates'], 1)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_failed_recheck_keeps_valid_download(self):
        results = iter([reviewed([ISSUE]), dict(status='unavailable', pages_checked=[], issues=[], warnings=['queue timeout'])])
        review, report = self.run_cycle(lambda *a, **k: next(results))
        self.assertEqual(report['status'], 'needs_review')
        self.assertEqual(self.path.read_bytes(), self.original)
        self.assertEqual(review['pages_checked'], [1])

    def test_exception_in_reviewer_retains_base_and_reports_unverified(self):
        def fail(*a, **k): raise TimeoutError()
        _, report = self.run_cycle(fail)
        self.assertEqual(self.path.read_bytes(), self.original)
        self.assertTrue(report['warnings'])
        self.assertNotEqual(report['status'], 'checked')

    def test_no_model_targets_do_not_enable_arbitrary_action(self):
        bad = dict(ISSUE, target_id='other', action='delete_text')
        self.assertEqual(select_actions([bad], [self.target]), [])
        self.assertEqual(select_actions([dict(ISSUE, action='center_picture')], [self.target]), [])
        self.assertEqual(select_actions([ISSUE], [self.target], {('layout_0001','restore_typography')}), [])

    def test_stale_document_rejected(self):
        with self.assertRaisesMessage(ValueError, 'Stale'):
            apply_actions(self.path, self.root/'candidate.docx', [self.target], [ISSUE], 'bad hash')
        self.assertFalse((self.root/'candidate.docx').exists())

    def test_inline_text_mutation_rolls_back_transaction(self):
        def corrupt(p, *args): p.find('.//w:t', NS).text = 'Fabricated finding'
        with patch('apps.template_workspace.v2.editor.safe_word_editor._apply_run_profile', side_effect=corrupt):
            with self.assertRaisesMessage(ValueError, 'native content'):
                apply_actions(self.path, self.root/'candidate.docx', [self.target], [ISSUE], sha256(self.original).hexdigest())
        self.assertFalse((self.root/'candidate.docx').exists())

    def test_disabled_cycle_performs_one_review_without_repair(self):
        with override_settings(TEMPLATE_V2_EDIT_CYCLE_ENABLED=False):
            _, report = self.run_cycle(lambda *a, **k: reviewed([ISSUE]))
        self.assertEqual(len(report['iterations']), 1)
        self.assertEqual(report['accepted_repairs'], 0)
        self.assertEqual(self.path.read_bytes(), self.original)

    def test_partial_coverage_cannot_be_quality_pass(self):
        r = reviewed(); r.update(status='partial', pages_total=3, pages_not_checked=[2,3])
        self.assertEqual(quality_status(r, []), 'partially_checked')
        r['issues'] = [ISSUE]
        self.assertEqual(quality_status(r, []), 'needs_review')

    def test_repagination_with_partial_coverage_rejected(self):
        a, b = reviewed([ISSUE]), reviewed()
        b.update(status='partial', pages_total=2, pages_not_checked=[2])
        self.assertFalse(improvement(a, b, [], [])[0])

    def test_missing_checked_page_is_not_improvement(self):
        a = reviewed([ISSUE]); a.update(status='partial', pages_total=3, pages_checked=[1,2])
        b = reviewed(); b.update(status='partial', pages_total=3, pages_checked=[1,3])
        self.assertFalse(improvement(a, b, [], [])[0])

    def test_partial_review_requires_target_and_neighbours(self):
        r = reviewed(); r.update(status='partial', pages_total=12, pages_checked=[1,2,6,7],
                                grounded_targets={2:[{'id':'front'}],7:[{'id':'body'}]})
        self.assertFalse(repairs_covered(r,[{'target_id':'body'}]))
        r['pages_checked'].append(8)
        self.assertTrue(repairs_covered(r,[{'target_id':'body'}]))
        self.assertFalse(repairs_covered(r,[{'target_id':'unmatched'}]))
        self.assertFalse(repairs_covered(r,[{'target_id':'front'}]))  # page 3 absent

    def test_review_issue_actions_only_with_same_page_whitelist(self):
        issue = dict(ISSUE, description='Visible issue')
        self.assertNotIn('target_id', validate_issues({'issues':[issue]}, 1)[0])
        self.assertIn('target_id', validate_issues({'issues':[issue]}, 1, [self.target])[0])
        issue['kind'] = 'header_alignment'
        self.assertNotIn('target_id', validate_issues({'issues':[issue]}, 1, [self.target])[0])

    def test_ambiguous_text_is_never_grounded(self):
        with pymupdf.open() as pdf:
            for _ in range(2): pdf.new_page().insert_text((40,40), 'Original scientific finding: 12.5 +/- 0.2 mg.')
            t = dict(self.target, text='Original scientific finding')
            self.assertEqual(ground_targets(pdf, [t]), {})
            pdf[1].insert_text((40,80), 'Unique known figure caption with enough characters.')
            t['text'] = 'Unique known figure caption with enough characters.'
            self.assertEqual(list(ground_targets(pdf, [t])), [2])

    @override_settings(TEMPLATE_V2_QWEN_ENABLED=True, TEMPLATE_V2_QWEN_TIMEOUT=300,
                       TEMPLATE_V2_VISUAL_REVIEW_ENABLED=True, TEMPLATE_V2_VISUAL_REVIEW_BUDGET=900)
    def test_worker_encloses_whole_cycle_budget(self):
        self.assertEqual(job_timeout_seconds(), 3120)


class NativeObjectFlowTests(SimpleTestCase):
    def test_three_inline_panels_use_full_width_with_their_native_caption(self):
        body = etree.Element(qn('w:body')); p = etree.SubElement(body,qn('w:p'))
        for _ in range(3):
            inline = etree.SubElement(etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:drawing')),qn('wp:inline'))
            etree.SubElement(inline,qn('wp:extent'),cx='2000000',cy=str(150*12700))
        caption = etree.SubElement(body,qn('w:p'))
        etree.SubElement(etree.SubElement(caption,qn('w:r')),qn('w:t')).text='Figure 6. Native panel group.'
        metadata={caption:_NodeMeta(None,'figure_caption',group_id='fig6')}
        section = etree.Element(qn('w:sectPr'))
        etree.SubElement(section,qn('w:cols')).set(qn('w:num'),'2')
        layout = SimpleNamespace(column_width_twips=4500,printable_width_twips=10000,body_left_twips=0,body_section=section)
        self.assertEqual(_wide_object_spans(body,metadata,{},layout), [[p,caption]])
        self.assertEqual(wrap_picture_captions(body,metadata,layout,{p,caption}),1)
        self.assertIs(p.getparent(),caption.getparent())
        self.assertIsNotNone(body.find('.//w:cantSplit',NS))

    def test_tall_panel_stack_in_narrow_column_is_not_wrapped_in_atomic_row(self):
        body = etree.Element(qn('w:body')); metadata = {}
        p = etree.SubElement(body,qn('w:p'))
        for _ in range(3):
            inline = etree.SubElement(etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:drawing')),qn('wp:inline'))
            etree.SubElement(inline,qn('wp:extent'),cx='2000000',cy=str(130*12700))
        p = etree.SubElement(body,qn('w:p'))
        etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:t')).text='Fig. 6. Three related panels.'
        metadata[p] = _NodeMeta(None,'figure_caption',group_id='fig6')
        layout = SimpleNamespace(column_width_twips=4500,body_left_twips=0)
        self.assertEqual(wrap_picture_captions(body,metadata,layout,set()),0)
        self.assertEqual(len(body.findall('w:tbl',NS)),0)
        self.assertEqual(len(body.xpath('.//w:drawing',namespaces=NS)),3)

    def test_front_shell_is_not_a_column_sized_article_figure(self):
        root = etree.Element(qn('w:document')); body = etree.SubElement(root, qn('w:body'))
        p = etree.SubElement(body, qn('w:p')); r = etree.SubElement(p, qn('w:r'))
        pict = etree.SubElement(r, qn('w:pict'))
        etree.SubElement(pict, qn('v:shape'),style='width:480pt;height:40pt')
        targets = inventory(root,{p:_NodeMeta(None,'front_shell',zone='front_matter')},SimpleNamespace(roles={}),None,set())
        self.assertEqual(targets, [])

    def test_actual_template_variant_is_exported_for_repair(self):
        root = etree.Element(qn('w:document')); body = etree.SubElement(root, qn('w:body'))
        p = etree.SubElement(body, qn('w:p')); etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:t')).text='Abstract'
        role = SimpleNamespace(typical_paragraph_formatting={},typical_run_formatting={'size':'20'})
        layout = SimpleNamespace(column_width_twips=4500,body_left_twips=0,body_right_twips=0)
        target = inventory(root,{p:_NodeMeta(None,'abstract')},SimpleNamespace(roles={'abstract':role}),layout,set(),
                           role_formats={p:({}, {'size':'22','bold':True})})[0]
        self.assertEqual(target['run_format']['size'],'22')

    def test_local_title_merge_survives_unrelated_front_object(self):
        body = etree.Element(qn('w:body')); metadata = {}
        for text, role in [('A scientific title', 'title'), ('with a continuation', 'title'), ('and another line', 'title'), ('', 'figure'), ('Body text', 'body')]:
            p = etree.SubElement(body, qn('w:p')); r = etree.SubElement(p, qn('w:r'))
            if text: etree.SubElement(r, qn('w:t')).text = text
            else: etree.SubElement(r, qn('w:drawing'))
            metadata[p] = _NodeMeta('b', role, zone='body' if role=='body' else 'front_matter', group_id='front_1')
        result = _normalise_front_matter(body, metadata, SimpleNamespace())
        self.assertEqual(result['front_paragraphs_merged'], 2)
        self.assertEqual(''.join(body[0].xpath('.//w:t/text()',namespaces=NS)), 'A scientific title with a continuation and another line')
        self.assertEqual(len(body.xpath('.//w:drawing',namespaces=NS)), 1)

    def test_parallel_text_gutter_not_enlarged_like_numeric_column(self):
        doc = Document(); table = doc.add_table(rows=1,cols=3)
        table.cell(0,0).text = 'Left biography with native details. '*8
        table.cell(0,2).text = 'Right biography with native details. '*8
        native = etree.fromstring(etree.tostring(table._tbl))
        for col, width in zip(native.findall('w:tblGrid/w:gridCol',NS), (4800,300,4800)):
            col.set(qn('w:w'),str(width))
        self.assertTrue(_resize_table_to_width(native,9900,SimpleNamespace(classification='DATA_TABLE')))
        self.assertEqual([c.get(qn('w:w')) for c in native.findall('w:tblGrid/w:gridCol',NS)],['4800','300','4800'])

    def test_mixed_run_text_gets_font_without_destroying_drawing(self):
        p = etree.Element(qn('w:p')); r = etree.SubElement(p, qn('w:r'))
        etree.SubElement(r, qn('w:t')).text = 'Scientific prose'
        drawing = etree.SubElement(r, qn('w:drawing'))
        _apply_run_profile(p, 'body', {'size':22})
        self.assertIs(r.find('w:drawing', NS), drawing)
        self.assertEqual(r.find('w:rPr/w:sz', NS).get(qn('w:val')), '22')

    def test_detach_native_anchor_preserves_objects_and_prose(self):
        body = etree.Element(qn('w:body')); p = etree.SubElement(body, qn('w:p'))
        r = etree.SubElement(p, qn('w:r')); etree.SubElement(r, qn('w:t')).text = 'Original description.'
        drawing = etree.SubElement(r, qn('w:drawing')); anchor = etree.SubElement(drawing, qn('wp:anchor'))
        etree.SubElement(anchor, qn('wp:extent'), cx='2000000', cy='2000000')
        metadata = {p:_NodeMeta('block_1','body', zone='body')}
        self.assertEqual(detach_display_objects(body, metadata, set()), 1)
        self.assertEqual(p.find('.//w:t', NS).text, 'Original description.')
        self.assertIs(body[1].find('.//w:drawing', NS), drawing)
        self.assertEqual(detach_display_objects(body, metadata, set()), 0)

    def test_actual_legacy_equation_not_resized(self):
        p = etree.Element(qn('w:p')); r = etree.SubElement(p, qn('w:r')); obj = etree.SubElement(r, qn('w:object'))
        shape = etree.SubElement(obj, qn('v:shape'), style='width:400pt;height:50pt')
        etree.SubElement(shape, qn('v:imagedata'))
        ole = etree.SubElement(obj, qn('o:OLEObject')); ole.set(qn('r:id'), 'equation')
        original = etree.tostring(p)
        self.assertEqual(normalize_display_geometry(p, 2000, {'equation'}), 0)
        self.assertEqual(etree.tostring(p), original)

    def test_non_equation_ole_preview_fits_proportionally(self):
        p = etree.Element(qn('w:p')); r = etree.SubElement(p, qn('w:r')); obj = etree.SubElement(r, qn('w:object'))
        shape = etree.SubElement(obj, qn('v:shape'), style='width:400pt;height:200pt;position:absolute;margin-left:100pt')
        etree.SubElement(shape, qn('v:imagedata'))
        self.assertGreater(normalize_display_geometry(p, 4000), 0)
        self.assertEqual(shape.get('style'), 'width:196.000pt;height:98.000pt')
