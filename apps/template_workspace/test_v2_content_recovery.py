from copy import deepcopy
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from django.test import SimpleTestCase
from docx import Document
from lxml import etree

from apps.template_workspace.v2.editor.content_recovery import (
    LocalContentGuard, ContentPreservationError, insert_rich_line_breaks,
)
from apps.template_workspace.v2.editor.safe_word_editor import SafeWordEditor, _balance_title_lines
from apps.template_workspace.v2.ooxml.namespaces import NS, qn


def native(doc):
    return etree.fromstring(etree.tostring(doc.element.body))


class ContentRecoveryTests(SimpleTestCase):
    def test_title_breaks_keep_mixed_indices_and_run_formatting(self):
        doc = Document(); p = doc.add_paragraph('Chain reactions in chemistry and physics: 130')
        p.add_run('t').font.superscript = True
        p.add_run('h').font.superscript = True
        p.add_run(' birth anniversary of a distinguished scientist and researcher').italic = True
        p = native(doc)[0]
        before = p.xpath('.//w:rPr/w:vertAlign/@w:val', namespaces=NS)
        self.assertEqual(_balance_title_lines(p), 1)
        self.assertEqual(p.xpath('.//w:rPr/w:vertAlign/@w:val', namespaces=NS), before)
        self.assertEqual(p.xpath('.//w:r[w:rPr/w:vertAlign]/w:t/text()', namespaces=NS), ['t', 'h'])
        self.assertTrue(p.xpath('.//w:i', namespaces=NS))
        self.assertEqual(_balance_title_lines(p), 0)

    def test_rich_breaks_across_whitespace_runs_keep_hyperlink_and_bookmarks(self):
        doc = Document(); p = doc.add_paragraph('Alpha '); p.add_run('  '); p.add_run(' beta')
        p = native(doc)[0]
        link = etree.Element(qn('w:hyperlink')); link.set(qn('r:id'), 'rId9'); link.append(p[-1]); p.append(link)
        mark = etree.SubElement(p, qn('w:bookmarkStart')); mark.set(qn('w:id'), '1'); mark.set(qn('w:name'), 'test')
        self.assertTrue(insert_rich_line_breaks(p, ['Alpha', 'beta']))
        self.assertEqual(p.xpath('.//w:t/text()', namespaces=NS), ['Alpha', '', '', 'beta'])
        self.assertEqual(len(p.xpath('.//w:br', namespaces=NS)), 1)
        self.assertIs(link.getparent(), p); self.assertIs(mark.getparent(), p)

    def test_complex_title_is_left_to_word_wrapping(self):
        for tag in ['w:fldChar', 'w:instrText', 'w:fldSimple', 'm:oMath', 'w:drawing']:
            doc = Document(); p = doc.add_paragraph('Alpha beta'); p = native(doc)[0]
            etree.SubElement(p, qn(tag)); before = etree.tostring(p)
            self.assertFalse(insert_rich_line_breaks(p, ['Alpha', 'beta']))
            self.assertEqual(etree.tostring(p), before)

    def test_repair_is_local_keeps_final_geometry_and_section(self):
        doc = Document(); p = doc.add_paragraph('130'); p.add_run('th').font.superscript = True
        doc.add_paragraph('Another paragraph'); body = native(doc); p, other = body[:2]
        guard = LocalContentGuard(body, {})
        prop = p.find('.//w:vertAlign', NS); prop.getparent().remove(prop)
        pp = etree.Element(qn('w:pPr')); p.insert(0, pp); etree.SubElement(pp, qn('w:sectPr'))
        other.set(qn('w:rsidR'), '12345678'); other_final = etree.tostring(other)
        events = guard.recover()
        self.assertEqual(events[0]['losses'], ['positioned_text'])
        self.assertTrue(p.xpath('.//w:vertAlign', namespaces=NS))
        self.assertIs(p.find('w:pPr', NS), pp)
        self.assertEqual(etree.tostring(other), other_final)
        self.assertEqual(guard.recover(), [])

    def test_cell_number_and_equation_order_and_relationship_are_restored(self):
        for change in ['number', 'math', 'link']:
            doc = Document(); p = doc.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
            p.add_run('12.5')
            math = etree.SubElement(p._p, qn('m:oMath'))
            for token in ['a', '-', 'b']:
                etree.SubElement(etree.SubElement(math, qn('m:r')), qn('m:t')).text = token
            link = etree.SubElement(p._p, qn('w:hyperlink')); link.set(qn('r:id'), 'rId5')
            body = native(doc); p = body.find('.//w:p', NS); original = etree.tostring(p)
            guard = LocalContentGuard(body, {})
            if change == 'number': p.find('.//w:t', NS).text = '125'
            elif change == 'math':
                math = p.find('m:oMath', NS); math.insert(0, math[-1])
            else: p.find('w:hyperlink', NS).set(qn('r:id'), 'rId6')
            self.assertEqual(len(guard.recover()), 1)
            self.assertEqual(etree.tostring(p), original)

    def test_exact_code_whitespace_restored(self):
        doc = Document(); doc.add_paragraph('    return 5')
        body = native(doc); p = body[0]; guard = LocalContentGuard(body, {}, {p})
        p.find('.//w:t', NS).text = 'return 5'
        self.assertIn('code_whitespace', guard.recover()[0]['losses'])
        self.assertEqual(p.find('.//w:t', NS).text, '    return 5')

    def test_local_formatting_exception_rolls_back_and_processing_continues(self):
        doc = Document(); doc.add_paragraph('Original'); doc.add_paragraph('Next')
        body = native(doc); guard = LocalContentGuard(body, {}); p = body[0]
        original = etree.tostring(p)
        with guard.formatting(p):
            p.find('.//w:t', NS).text = 'Damaged'
            raise ValueError('invalid local formatting value')
        self.assertEqual(etree.tostring(p), original)
        self.assertEqual(guard.recover()[0]['losses'], ['formatting_exception:ValueError'])
        with self.assertRaises(OSError):
            with guard.formatting(p): raise OSError('disk failure')

    def test_plain_text_formula_sign_change_is_restored(self):
        doc = Document(); doc.add_paragraph('a-b = 1.5')
        body = native(doc); guard = LocalContentGuard(body, {})
        body[0].find('.//w:t', NS).text = 'a+b = 1,5'
        self.assertEqual(guard.recover()[0]['losses'], ['scientific_text'])

    def test_removing_blank_section_separator_is_not_content_loss(self):
        doc = Document(); p = doc.add_paragraph()._p
        sect = etree.SubElement(p.get_or_add_pPr(), qn('w:sectPr'))
        etree.SubElement(sect, qn('w:headerReference')).set(qn('r:id'), 'rId7')
        body = native(doc); guard = LocalContentGuard(body, {}); body.remove(body[0])
        self.assertEqual(guard.recover(), [])

    def test_removed_nonempty_paragraph_fails_but_blank_spacer_is_allowed(self):
        doc = Document(); doc.add_paragraph('Important'); doc.add_paragraph()
        body = native(doc); guard = LocalContentGuard(body, {}); body.remove(body[1])
        self.assertEqual(guard.recover(), [])
        body.remove(body[0])
        with self.assertRaises(ContentPreservationError): guard.recover()

    def test_allowed_caption_and_author_conventions_do_not_trigger_recovery(self):
        doc = Document(); doc.add_paragraph('Figure 1. Data'); p = doc.add_paragraph('Author')
        p.add_run('*').font.superscript = True
        body = native(doc); p, author = body[:2]
        guard = LocalContentGuard(body, {p: SimpleNamespace(role='figure_caption'), author: SimpleNamespace(role='author')})
        p.find('.//w:t', NS).text = 'Fig. 1. Data'
        author[-1].find('w:t', NS).text = ''
        etree.SubElement(author[-1], qn('w:sym')).set(qn('w:char'), 'F02A')
        self.assertEqual(guard.recover(), [])

    def test_full_editor_recovers_injected_title_loss_and_writes_valid_result(self):
        with TemporaryDirectory() as tmp:
            tmp = Path(tmp); source = Document(); p = source.add_paragraph('Scientific 130')
            p.add_run('th').font.superscript = True
            source.add_paragraph('A. Author'); source.add_paragraph('Introduction')
            source.add_paragraph('Reliable measurements of 42 samples are retained.')
            source.save(tmp/'article.docx'); source.save(tmp/'template.docx')
            from apps.template_workspace.v2.editor import safe_word_editor as editor
            original = editor._apply_role_format
            def damage(p, *args, **kwargs):
                original(p, *args, **kwargs)
                for node in p.xpath('.//w:vertAlign', namespaces=NS): node.getparent().remove(node)
            with patch.object(editor, '_apply_role_format', side_effect=damage):
                result = SafeWordEditor().render(article_path=tmp/'article.docx', template_path=tmp/'template.docx', output_path=tmp/'result.docx')
            self.assertTrue(result.metrics['native_integrity']['passed'])
            self.assertGreater(result.metrics['local_content_recovery_count'], 0)
            self.assertTrue(any('восстановлено' in warning for warning in result.warnings))
            with ZipFile(tmp/'result.docx') as archive:
                root = etree.fromstring(archive.read('word/document.xml'))
                self.assertEqual(root.xpath('//w:r[w:rPr/w:vertAlign]/w:t/text()', namespaces=NS), ['th'])
