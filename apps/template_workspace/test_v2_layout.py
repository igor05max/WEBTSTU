from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile
from unittest.mock import patch
import base64
from docx import Document
from docx.shared import Pt
from django.test import SimpleTestCase, override_settings
from lxml import etree
from apps.template_workspace.v2.ooxml.namespaces import NS, qn
from apps.template_workspace.v2.editor.layout_fidelity import (
    front_gap_evidence, apply_front_gaps, align_standalone_picture,
    scientific_table_rules, math_typography, descriptions_before_figures, wrap_picture_captions,
    heading_gap_evidence, apply_heading_gaps,
)
from apps.template_workspace.v2.editor.safe_word_editor import _guard_inline_figure_captions
from apps.template_workspace.v2.editor.integrity import native_integrity
from apps.template_workspace.v2.readability import QwenReadabilityProvider, review_pdf, validate_issues, reference_page_index


def xml(node):
    return etree.fromstring(etree.tostring(node))


def package(doc):
    stream = BytesIO(); doc.save(stream)
    return ZipFile(stream)


def meta(role):
    return SimpleNamespace(role=role, group_id='front')


class LayoutFidelityTests(SimpleTestCase):
    def test_reference_skips_instructions_before_sample_title(self):
        pages=[SimpleNamespace(get_text=lambda:'Instructions: formatting requirements'),
               SimpleNamespace(get_text=lambda:'A scientific\nsample title\nAuthors')]
        self.assertEqual(reference_page_index(pages,'A scientific sample title'),1)

    def test_short_template_placeholder_must_be_its_own_line(self):
        pages=[SimpleNamespace(get_text=lambda:'Put the Title here'),
               SimpleNamespace(get_text=lambda:'Title\nFirstname Lastname')]
        self.assertEqual(reference_page_index(pages,'Title'),1)

    def test_unmatched_reference_title_does_not_guess_page_one(self):
        pages=[SimpleNamespace(get_text=lambda:'Submission instructions')]
        self.assertIsNone(reference_page_index(pages,'Missing title'))
        self.assertIsNone(reference_page_index(pages,''))

    def test_front_blank_line_geometry_is_role_specific_and_idempotent(self):
        doc = Document()
        doc.styles['Normal'].font.size = Pt(10)
        doc.styles['Normal'].paragraph_format.space_after = Pt(0)
        doc.add_paragraph('Title'); doc.add_paragraph(); doc.add_paragraph('Authors')
        structure = SimpleNamespace(blocks=[{'id':'block_0001','detected_role':'title'}, {'id':'block_0003','detected_role':'author'}])
        with package(doc) as z: gaps = front_gap_evidence(z, structure)
        self.assertGreater(gaps['title>author'], 0)
        body = xml(doc.element.body); a,b,c = list(body)[:3]
        metadata = {a:meta('title'), c:meta('author')}
        apply_front_gaps(body, metadata, gaps)
        expected = etree.tostring(body)
        apply_front_gaps(body, metadata, gaps)
        self.assertEqual(etree.tostring(body), expected)
        self.assertNotIn(b, list(body))
        self.assertEqual(a.find('w:pPr/w:spacing',NS).get(qn('w:after')), str(gaps['title>author']))

    def test_unknown_front_neighbour_does_not_crash(self):
        doc=Document();doc.add_paragraph('Title');doc.add_paragraph('Unknown')
        body=xml(doc.element.body);a,b=list(body)[:2]
        self.assertEqual(apply_front_gaps(body,{a:meta('title'),b:meta(None)},{}),0)

    def test_front_field_and_section_separator_are_not_removed(self):
        doc=Document();doc.add_paragraph('Title');p=doc.add_paragraph()._p
        etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:fldChar')).set(qn('w:fldCharType'),'begin')
        doc.add_paragraph('Authors');body=xml(doc.element.body)
        a,b,c=list(body)[:3]
        apply_front_gaps(body,{a:meta('title'),c:meta('author')},{'title>author':240})
        self.assertIn(b,list(body))

    def test_standalone_image_loses_source_offset_but_inline_prose_does_not(self):
        doc=Document();p=doc.add_paragraph()._p
        p.get_or_add_pPr().append(etree.Element(qn('w:ind')))
        p.find('w:pPr/w:ind',NS).set(qn('w:firstLine'),'709')
        drawing=etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:drawing'))
        anchor=etree.SubElement(drawing,qn('wp:anchor'))
        etree.SubElement(anchor,qn('wp:extent')).set('cx','1000')
        etree.SubElement(anchor,qn('a:graphic'))
        p=xml(p);layout=SimpleNamespace(body_left_twips=0,body_right_twips=0)
        self.assertEqual(align_standalone_picture(p,layout),1)
        self.assertEqual(p.find('w:pPr/w:ind',NS).get(qn('w:firstLine')),'0')
        self.assertEqual(p.find('w:pPr/w:jc',NS).get(qn('w:val')),'center')
        self.assertEqual(len(p.xpath('.//wp:inline',namespaces=NS)),1)
        etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:t')).text='inline prose'
        snapshot=etree.tostring(p)
        self.assertEqual(align_standalone_picture(p,layout),0)
        self.assertEqual(etree.tostring(p),snapshot)

    def test_table_rules_keep_native_cells_and_apply_consistent_alignment(self):
        doc=Document();table=doc.add_table(rows=3,cols=2);table.style='Table Grid'
        for i,row in enumerate(table.rows):
            row.cells[0].text='A very long descriptive model identifier' if i else 'Model'
            row.cells[1].text=str(i)
        table=xml(table._tbl);before=table.xpath('.//w:t/text()',namespaces=NS)
        self.assertEqual(scientific_table_rules(table,SimpleNamespace(classification='DATA_TABLE')),1)
        self.assertEqual(before,table.xpath('.//w:t/text()',namespaces=NS))
        self.assertEqual(table.xpath('./w:tr[2]/w:tc/w:p/w:pPr/w:jc/@w:val',namespaces=NS),['left','center'])
        self.assertTrue(all(n.get(qn('w:val'))=='nil' for n in table.xpath('.//w:tcBorders/w:left|.//w:tcBorders/w:insideV|./w:tblPr/w:tblBorders/w:right',namespaces=NS)))
        snapshot=etree.tostring(table)
        self.assertEqual(scientific_table_rules(table,SimpleNamespace(classification='FIGURE_CONTAINER')),0)
        self.assertEqual(etree.tostring(table),snapshot)

    def test_math_formatting_is_allowed_but_reordered_math_is_rejected(self):
        doc=Document();p=doc.add_paragraph()._p
        math=etree.SubElement(p,qn('m:oMath'))
        for token in ['a','-','b']:
            r=etree.SubElement(math,qn('m:r'));etree.SubElement(r,qn('m:t')).text=token
        root=xml(doc.element);body=root.find('w:body',NS)
        profile=SimpleNamespace(typical_run_formatting={'size':'22','fonts':{qn('w:ascii'):'Times New Roman'}})
        with package(doc) as z:
            math_typography(body,z,profile)
            self.assertTrue(native_integrity(z,root,{})['passed'])
            changed=root.find('.//m:oMath',NS); changed.insert(0,changed[-1])
            self.assertIn('math_structures',native_integrity(z,root,{})['losses'])

    def test_caption_keep_chain_overrides_false_across_blank(self):
        doc=Document();pic=doc.add_paragraph();pic.paragraph_format.keep_with_next=False
        etree.SubElement(etree.SubElement(pic._p,qn('w:r')),qn('w:drawing'))
        doc.add_paragraph();cap=doc.add_paragraph('Fig. 1. Example');cap.paragraph_format.keep_together=False
        body=xml(doc.element.body);a,b,c=list(body)[:3]
        _guard_inline_figure_captions(body=body,meta_by_node={c:meta('figure_caption')})
        self.assertEqual(a.find('w:pPr/w:keepNext',NS).get(qn('w:val')),'1')
        self.assertEqual(b.find('w:pPr/w:keepNext',NS).get(qn('w:val')),'1')
        self.assertEqual(c.find('w:pPr/w:keepLines',NS).get(qn('w:val')),'1')

    def test_figure_moves_after_nearby_mention_not_other_prose(self):
        doc=Document();p=doc.add_paragraph()._p
        etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:drawing'))
        doc.add_paragraph('Рис. 2. Example');doc.add_paragraph('Как показано на рис. 2, результат.')
        body=xml(doc.element.body);a,b,c=list(body)[:3]
        md={b:meta('figure_caption'),c:meta('body')}
        self.assertEqual(descriptions_before_figures(body,md),1)
        self.assertEqual(list(body)[:3],[c,a,b])
        self.assertEqual(descriptions_before_figures(body,md),0)

    @override_settings(TEMPLATE_V2_QWEN_MODEL='local',TEMPLATE_V2_QWEN_BASE_URL='http://local/v1')
    def test_reference_image_is_labelled_and_only_used_for_front_pages(self):
        provider=QwenReadabilityProvider(b'reference')
        with patch('apps.template_workspace.v2.readability._request_json',return_value={'choices':[{'message':{'content':'{"issues":[]}'}}]}) as request:
            provider.review([b'top',b'bottom'],1,30)
            content=request.call_args.kwargs['payload']['messages'][1]['content']
            self.assertEqual(request.call_count,2)
            self.assertEqual(len([i for i in content if i['type']=='image_url']),2)
            self.assertIn('chunk 2/2',content[0]['text'])
            provider.review([b'top',b'bottom'],5,30)
            content=request.call_args.kwargs['payload']['messages'][1]['content']
            self.assertEqual(len([i for i in content if i['type']=='image_url']),2)

    def test_new_visual_findings_are_advisory_not_edit_operations(self):
        for kind in ['front_spacing','figure_alignment','equation_typography','header_alignment','table_rules']:
            item=validate_issues({'issues':[{'kind':kind,'severity':'high','description':'visible defect','operation':'rewrite'}]},2)[0]
            self.assertTrue(item['advisory']);self.assertNotIn('operation',item)

    @override_settings(TEMPLATE_V2_QWEN_MODEL='local',TEMPLATE_V2_QWEN_BASE_URL='http://local/v1')
    def test_incomplete_front_comparison_is_not_a_checked_page(self):
        import pymupdf, tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'front.pdf'
            with pymupdf.open() as doc:
                doc.new_page().insert_text((40,40),'Title')
                doc.save(path)
            provider=QwenReadabilityProvider(b'reference')
            with patch.object(provider,'_review_request',side_effect=[[],TimeoutError]):
                report=review_pdf(path,provider=provider,max_pages=1)
            self.assertEqual(report['pages_checked'],[])
            self.assertEqual(report['pages_not_checked'],[1])
            self.assertEqual(report['status'],'unavailable')

    def test_malformed_page_answer_does_not_abort_other_pages(self):
        import pymupdf, tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'pages.pdf'
            with pymupdf.open() as doc:
                for i in range(3): doc.new_page().insert_text((40,40),f'Page {i+1}')
                doc.save(path)
            with patch.object(QwenReadabilityProvider,'review',side_effect=[[],ValueError('Bad JSON'),[]]):
                report=review_pdf(path,provider=QwenReadabilityProvider(),max_pages=3)
            self.assertEqual(report['pages_checked'],[1,3])
            self.assertEqual(report['pages_not_checked'],[2])
            self.assertEqual(report['status'],'partial')

    @override_settings(TEMPLATE_V2_QWEN_MODEL='local',TEMPLATE_V2_QWEN_BASE_URL='http://local/v1')
    def test_truncated_model_output_cannot_be_a_quality_pass(self):
        with patch('apps.template_workspace.v2.readability._request_json',return_value={
                'choices':[{'finish_reason':'length','message':{'content':'{"issues":[]}'}}]}):
            with self.assertRaises(ValueError):
                QwenReadabilityProvider().review([b'image'],3,10)

    def test_atomic_picture_caption_row_is_borderless_and_preserves_nodes(self):
        doc=Document();p=doc.add_paragraph()._p
        drawing=etree.SubElement(etree.SubElement(p,qn('w:r')),qn('w:drawing'))
        holder=etree.SubElement(drawing,qn('wp:inline'))
        extent=etree.SubElement(holder,qn('wp:extent'));extent.set('cx','1000');extent.set('cy','1000')
        doc.add_paragraph('Fig. 1. Native caption')
        body=xml(doc.element.body);p,c=list(body)[:2]
        layout=SimpleNamespace(printable_width_twips=9000,column_width_twips=4000,body_left_twips=0)
        self.assertEqual(wrap_picture_captions(body,{c:meta('figure_caption')},layout,set()),1)
        table=body[0]
        self.assertIs(table.find('w:tr/w:tc/w:p',NS),p)
        self.assertIs(p.getnext(),c)
        self.assertIsNotNone(table.find('w:tr/w:trPr/w:cantSplit',NS))
        self.assertIsNone(table.find('w:tr/w:trPr/w:cantSplit',NS).get(qn('w:val')))
        self.assertTrue(all(n.get(qn('w:val'))=='nil' for n in table.findall('w:tblPr/w:tblBorders/*',NS)))

    def test_jamt_header_override_does_not_change_other_journal(self):
        from apps.template_workspace.v2.editor.safe_word_editor import _merge_template_header_footer
        for label,expected in [('Journal of Advanced Materials and Technologies. 2026. Vol. 11','right'),
                               ('Journal of Testing. 2026. Vol. 1','left')]:
            source=Document();source.add_paragraph('Author')
            template=Document();template.sections[0].header.paragraphs[0].text=label
            root=xml(source.element)
            with package(source) as a, package(template) as t:
                parts=_merge_template_header_footer(article_zip=a,template_zip=t,document_root=root,author_shortline='Author')
            header=next(etree.fromstring(data) for name,data in parts.items() if name.startswith('word/header') and name.endswith('.xml'))
            p=next(p for p in header.findall('w:p',NS) if p.xpath('.//w:t',namespaces=NS))
            self.assertEqual(p.find('w:pPr/w:jc',NS).get(qn('w:val')),expected)

    def test_front_pages_cannot_be_starved_by_dense_body_pages(self):
        import pymupdf, tempfile
        from pathlib import Path
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'pages.pdf'
            with pymupdf.open() as doc:
                for i in range(4):
                    page=doc.new_page();page.insert_text((40,40),'Title' if i < 2 else '"dense": 1.2, "keys": 3.4')
                doc.save(path)
            report=review_pdf(path,provider=SimpleNamespace(review=lambda *a,**k:[]),max_pages=2)
            self.assertEqual(report['pages_checked'],[1,2])

    def test_synthetic_heading_gap_uses_reference_evidence(self):
        doc=Document();doc.add_paragraph();doc.add_paragraph('References');doc.add_paragraph()
        structure=SimpleNamespace(blocks=[{'id':'block_0002','detected_role':'references_heading'}])
        with package(doc) as z: evidence=heading_gap_evidence(z,structure)
        body=xml(doc.element.body);p=body[1]
        self.assertEqual(apply_heading_gaps(body,{p:meta('heading_2')},evidence),2)
        self.assertEqual(apply_heading_gaps(body,{p:meta('heading_2')},evidence),0)
        self.assertEqual(int(p.find('w:pPr/w:spacing',NS).get(qn('w:before'))),evidence['references_heading']['before'])
