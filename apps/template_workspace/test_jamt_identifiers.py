from pathlib import Path
from types import SimpleNamespace
import tempfile
from unittest import TestCase

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from .v2.editor.safe_word_editor import SafeWordEditor
from .v2.classification.roles import RoleClassifierV2
from .v2.inspector.document import DocumentInspector
from .v2.ooxml.namespaces import NS, qn
from .v2.styles.jamt import prepare_jamt_style, load_style
from .v2.styles.identifiers import arrange_identifiers, repair_existing_identifiers
from paper_formatter.latex_lab.bridge import NativeBridge
from paper_formatter.latex_lab.typesetter import default_plan, render


class JamtIdentifierTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def format(self, identifiers):
        doc = Document()
        for text in ['Original papers', 'Nobelistics', *identifiers,
                     'Analysis of a demonstration material', 'Ivan I. Petrov',
                     'Demonstration Institute, Moscow', 'Abstract. Original scientific text.',
                     'Keywords: model; sample', '1. Introduction',
                     'The body mentions DOI: 10.9999/body and retains its original text.',
                     'References', '1. Author A. DOI: 10.9999/reference']:
            doc.add_paragraph(text)
        source = self.root/'source.docx'; doc.save(source)
        carrier, report, profile = prepare_jamt_style(self.root/'style')
        result = self.root/'result.docx'
        metadata = SafeWordEditor().render(article_path=source, output_path=result,
            template_path=carrier, template_report=report, template_profile=profile)
        self.assertTrue(metadata.metrics['native_integrity']['passed'])
        blocks, manifest = NativeBridge(result, self.root/'latex').build()
        self.assertTrue(manifest['text_transfer']['exact'])
        tex = render(blocks, self.root/'latex', default_plan(blocks)).read_text(encoding='utf-8')
        return Document(result), tex.replace(r'\allowbreak{}', '')

    def test_separate_reversed_identifiers_share_right_tab_in_word_and_latex(self):
        doc, tex = self.format(['DOI: 10.9999/demo', 'UDC 620.3'])
        p = next(p for p in doc.paragraphs if p.text.startswith('UDC'))
        self.assertEqual(p.text, 'UDC 620.3\tDOI: 10.9999/demo')
        tab = p._p.find('w:pPr/w:tabs/w:tab', NS)
        self.assertEqual(tab.get(qn('w:val')), 'right')
        self.assertEqual(tab.get(qn('w:pos')), '9865')
        self.assertIn(r'\JAMTIdentifierLine{UDC 620.3}{DOI: 10.9999/demo}', tex)
        self.assertNotIn('{doi_metadata}{1. Author', tex)
        frames = [p._p.find('w:pPr/w:framePr', NS) for p in doc.paragraphs[:2]]
        self.assertTrue(all(f is not None and f.get(qn('w:y')) == '1478' for f in frames))
        self.assertEqual(doc.sections[0].top_margin.twips, 1701)
        self.assertIn(r'\JAMTFrontStart', tex)
        header = doc.sections[0].header.paragraphs[0]._p.find('w:pPr/w:framePr', NS)
        self.assertEqual(header.get(qn('w:y')), '1083')

    def test_standalone_doi_is_right_aligned_but_body_and_reference_are_not(self):
        doc, tex = self.format(['DOI: 10.9999/demo'])
        p = next(p for p in doc.paragraphs if p.text.startswith('DOI:'))
        self.assertEqual(p._p.find('w:pPr/w:jc', NS).get(qn('w:val')), 'right')
        self.assertIn(']{doi_metadata}{DOI: 10.9999/demo}', tex)
        body = next(p for p in doc.paragraphs if p.text.startswith('The body'))
        self.assertEqual(body._p.find('w:pPr/w:jc', NS).get(qn('w:val')), 'both')

    def test_combined_one_run_gets_tab_without_rewriting_characters(self):
        doc, tex = self.format(['УДК 620.3 DOI: 10.9999/demo'])
        p = next(p for p in doc.paragraphs if p.text.startswith('УДК'))
        self.assertEqual(p.text.replace('\t', ''), 'УДК 620.3 DOI: 10.9999/demo')
        self.assertEqual(len(p._p.xpath('.//w:r/w:tab')), 1)
        self.assertIn(r'\JAMTIdentifierLine{', tex)

    def test_merge_keeps_yellow_runs_hyperlinks_and_targets(self):
        doc = Document()
        doi = doc.add_paragraph('DOI: ')
        hyperlink = OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'), 'rIdImportant')
        run = OxmlElement('w:r'); props = OxmlElement('w:rPr')
        highlight = OxmlElement('w:highlight'); highlight.set(qn('w:val'), 'yellow')
        props.append(highlight); run.append(props)
        text = OxmlElement('w:t'); text.text = '[НЕ УКАЗАНО: DOI]'; run.append(text)
        hyperlink.append(run); doi._p.append(hyperlink)
        udc = doc.add_paragraph('УДК: [НЕ УКАЗАНО: индекс]')
        body = etree.fromstring(etree.tostring(doc._element.body))
        meta = {p: SimpleNamespace(role='editorial_metadata', zone='front_matter', subtype=None) for p in body.findall('w:p', NS)}
        before = sorted(body.xpath('.//w:t/text()', namespaces=NS))
        self.assertEqual(arrange_identifiers(body, meta), 1)
        self.assertEqual(before, sorted(body.xpath('.//w:t/text()', namespaces=NS)))
        self.assertEqual(body.xpath('.//w:hyperlink/@r:id', namespaces=NS), ['rIdImportant'])
        self.assertEqual(body.xpath('.//w:highlight/@w:val', namespaces=NS), ['yellow'])
        self.assertEqual(arrange_identifiers(body, meta), 0)

    def test_duplicate_identifiers_are_not_silently_merged_or_dropped(self):
        doc = Document()
        for t in ('УДК 620.3', 'DOI: 10.9999/one', 'DOI: 10.9999/two'): doc.add_paragraph(t)
        body = etree.fromstring(etree.tostring(doc._element.body))
        meta = {p: SimpleNamespace(role='editorial_metadata', zone='front_matter', subtype=None) for p in body.findall('w:p', NS)}
        before = etree.tostring(body)
        self.assertEqual(arrange_identifiers(body, meta), 0)
        self.assertEqual(before, etree.tostring(body))

    def test_conservative_path_repairs_wrong_tab_and_then_is_idempotent(self):
        doc = Document(); doc.add_paragraph('УДК 620.3\tDOI: 10.9999/demo')
        doc.add_paragraph('Demonstration title', 'Title')
        source = self.root/'source.docx'; doc.save(source)
        def repair():
            report = DocumentInspector(source).inspect()
            structure = RoleClassifierV2(use_ai=False).article_structure(report)
            root = etree.fromstring(etree.tostring(Document(source)._element))
            result = repair_existing_identifiers(root, report, structure, load_style())
            return result, root
        changes, root = repair()
        self.assertEqual(len(changes), 1)
        from zipfile import ZipFile, ZIP_DEFLATED
        with ZipFile(source) as z: parts = {n:z.read(n) for n in z.namelist()}
        parts['word/document.xml'] = etree.tostring(root)
        with ZipFile(source, 'w', ZIP_DEFLATED) as z:
            for n, data in parts.items(): z.writestr(n, data)
        changes, _ = repair()
        self.assertFalse(changes)

    def test_bilingual_biographies_move_as_original_nodes_without_editing_text(self):
        from .v2.styles.author_columns import arrange_author_columns
        doc = Document()
        labels = [('Information about the authors','author_information','en'),
                  ('Alex Smith; original email: a@example.org.','author_bio','en'),
                  ('Информация об авторах','author_information','ru'),
                  ('Алекс Смит; исходные сведения об авторе.','author_bio','ru')]
        for text,_,_ in labels: doc.add_paragraph(text)
        body = etree.fromstring(etree.tostring(doc._element.body))
        meta = {p:SimpleNamespace(role=role,language=lang) for p,(_,role,lang) in zip(body,labels)}
        before = body.xpath('.//w:t/text()',namespaces=NS)
        self.assertEqual(arrange_author_columns(body,meta),1)
        self.assertEqual(sorted(before),sorted(body.xpath('.//w:t/text()',namespaces=NS)))
        cells = body.find('w:tbl',NS).findall('w:tr/w:tc',NS)
        self.assertIn('Alex Smith',''.join(cells[0].xpath('.//w:t/text()',namespaces=NS)))
        self.assertIn('Алекс Смит',''.join(cells[2].xpath('.//w:t/text()',namespaces=NS)))
        self.assertFalse(cells[1].xpath('.//w:t',namespaces=NS))
        self.assertEqual(arrange_author_columns(body,meta),0)
