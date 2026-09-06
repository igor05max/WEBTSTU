import io
import tempfile
from datetime import timedelta
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from docx import Document

from paper_formatter.compiler import LatexCompiler
from paper_formatter.parsers.latex_parser import LatexParser
from .models import TemplateJob
from .services import output_directory, run_job


def document_upload(name="article.docx", template=False):
    doc = Document()
    doc.add_heading("Образец оформления" if template else "Исследование материалов", 0)
    doc.add_heading("Введение", 1)
    doc.add_paragraph("Сохраняемый научный текст: значение 10% и A_B.")
    doc.add_heading("Результаты", 1)
    doc.add_paragraph("Получены воспроизводимые результаты исследования.")
    stream = io.BytesIO()
    doc.save(stream)
    return SimpleUploadedFile(name, stream.getvalue())


class WorkspaceTests(TestCase):
    def setUp(self):
        self.media = tempfile.TemporaryDirectory()
        self.addCleanup(self.media.cleanup)
        self.settings_override = override_settings(MEDIA_ROOT=self.media.name, AI_BASE_URL="")
        self.settings_override.enable()
        self.addCleanup(self.settings_override.disable)
        self.user = get_user_model().objects.create_user(username="template_author")
        self.client.force_login(self.user)
        self.url = reverse("template_workspace:workspace")

    def create_job(self):
        return TemplateJob.objects.create(owner=self.user, article=document_upload(),
            template=document_upload("template.docx", True), article_name="article.docx", template_name="template.docx")

    def test_login_and_navigation(self):
        response = self.client.get(self.url)
        self.assertContains(response, "Статья по вашему шаблону")
        self.assertContains(response, 'aria-current="page"')
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)

    @patch("apps.template_workspace.views.launch_job")
    def test_upload_persists_inputs_and_launches_after_commit(self, launch):
        with self.captureOnCommitCallbacks(execute=True):
            response = self.client.post(self.url, {"article": document_upload(), "template": document_upload("template.docx", True)})
        self.assertEqual(response.status_code, 302)
        job = TemplateJob.objects.get()
        self.assertTrue(Path(job.article.path).is_file())
        launch.assert_called_once()

    def test_rejects_unsupported_input(self):
        response = self.client.post(self.url, {"article": SimpleUploadedFile("bad.exe", b"x"), "template": document_upload()})
        self.assertContains(response, "Выберите DOCX")
        self.assertFalse(TemplateJob.objects.exists())

    def test_other_users_cannot_access_job_or_files(self):
        job = self.create_job()
        other = get_user_model().objects.create_user(username="other_author")
        self.client.force_login(other)
        for name, args in [("detail", [job.pk]), ("progress", [job.pk]), ("download", [job.pk, "latex"])]:
            self.assertEqual(self.client.get(reverse("template_workspace:" + name, args=args)).status_code, 404)

    @patch("apps.template_workspace.views.launch_job")
    def test_prevents_duplicate_active_work(self, launch):
        self.create_job()
        response = self.client.post(self.url, {"article": document_upload(), "template": document_upload("template.docx", True)})
        self.assertContains(response, "Дождитесь завершения")
        self.assertEqual(TemplateJob.objects.count(), 1)
        launch.assert_not_called()

    @patch.object(LatexCompiler, "_select_engine", return_value=None)
    def test_real_pipeline_retains_tex_when_compiler_unavailable(self, select):
        job = self.create_job()
        run_job(job.pk)
        job.refresh_from_db()
        self.assertEqual(job.status, "partial")
        self.assertTrue(job.plan)
        archive = output_directory(job) / "result/latex_project.zip"
        with ZipFile(archive) as z:
            self.assertIn("main.tex", z.namelist())
            self.assertIn("Сохраняемый", z.read("body.tex").decode())
        self.assertFalse((output_directory(job) / "result/result.docx").exists())
        response = self.client.get(reverse("template_workspace:download", args=[job.pk, "latex"]))
        self.assertEqual(response.status_code, 200)
        response.close()
        self.assertEqual(self.client.get(reverse("template_workspace:download", args=[job.pk, "pdf"])).status_code, 404)

    @patch.object(LatexCompiler, "compile")
    def test_pdf_download_uses_compiler_artifact_without_rebuilding(self, compile_mock):
        def compile_pdf(project, log):
            self.assertTrue((project / "main.tex").exists())
            pdf = project / "main.pdf"
            pdf.write_bytes(b"%PDF-1.4\ncontrolled compiler output")
            return pdf, []
        compile_mock.side_effect = compile_pdf
        job = self.create_job()
        run_job(job.pk)
        job.refresh_from_db()
        self.assertIn(job.status, ["completed", "partial"])
        for _ in range(2):
            response = self.client.get(reverse("template_workspace:download", args=[job.pk, "pdf"]) + "?preview=1")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response["X-Frame-Options"], "SAMEORIGIN")
            self.assertIn(b"controlled compiler output", b"".join(response.streaming_content))
            response.close()
        compile_mock.assert_called_once()

    def test_compiler_disables_shell_and_local_latexmk_configuration(self):
        command = LatexCompiler._command("latexmk", "latexmk")
        self.assertIn("-no-shell-escape", command)
        self.assertIn("-norc", command)

    def test_abandoned_worker_becomes_failed(self):
        job = self.create_job()
        TemplateJob.objects.filter(pk=job.pk).update(updated_at=timezone.now() - timedelta(minutes=21))
        response = self.client.get(reverse("template_workspace:progress", args=[job.pk]))
        self.assertEqual(response.json()["status"], "failed")
        self.assertFalse(response.json()["pending"])

    def test_latex_parser_does_not_read_outside_project(self):
        root = Path(self.media.name)
        (root / "private.tex").write_text("PRIVATE_SENTINEL", encoding="utf-8")
        project = root / "project"
        project.mkdir()
        source = project / "main.tex"
        source.write_text(r"\documentclass{article}\begin{document}\input{../private.tex}\end{document}", encoding="utf-8")
        article = LatexParser(source, project / "assets").parse()
        self.assertNotIn("PRIVATE_SENTINEL", article.model_dump_json())
        self.assertTrue(any("вне LaTeX-проекта" in w for w in article.warnings))
