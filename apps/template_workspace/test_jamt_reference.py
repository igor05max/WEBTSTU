from pathlib import Path
from types import SimpleNamespace
import tempfile
from unittest import TestCase
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Pt, Twips
from lxml import etree

from apps.template_workspace.v2.inspector.document import DocumentInspector
from apps.template_workspace.v2.ooxml.namespaces import NS, qn
from apps.template_workspace.v2.styles.jamt import load_style
from apps.template_workspace.v2.styles.reference_fidelity import assess_existing_layout, preserve_existing_layout


class ReferenceFidelityTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def source(self):
        doc = Document()
        doc.styles['Normal'].font.size = Pt(11)
        doc.styles['Normal'].font.name = 'Times New Roman'
        roles = []
        def paragraph(text, role, size):
            p = doc.add_paragraph(text); p.runs[0].font.size = Pt(size)
            roles.append({'id': f'block_{len(doc._element.body)-1:04d}', 'detected_role': role})
            return p
        for title, abstract in [('Reference article title', 'Abstract. Original scientific text.'),
                                ('Название исследования', 'Аннотация. Исходный научный текст.')]:
            paragraph(title, 'title', 14)
            paragraph(abstract, 'abstract', 10)
        section = doc.add_section(WD_SECTION.CONTINUOUS)
        section._sectPr.find(qn('w:cols')).set(qn('w:num'), '2')
        for i in range(24):
            paragraph(f'Original measurement {i}: 42 ± 0.2 MPa; DOI 10.1234/example.', 'body', 11)
        paragraph('References', 'references_heading', 11)
        paragraph('1. Source A. Original research. 2026.', 'reference_item', 10)
        paragraph('Information about the authors / Информация об авторах', 'author_information', 12)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Original author'; table.cell(0, 1).text = 'Исходные сведения'
        style = load_style()
        for s in doc.sections:
            s.page_width = Twips(style['page_twips']['w']); s.page_height = Twips(style['page_twips']['h'])
            for side in ('top', 'bottom', 'left', 'right'):
                setattr(s, side+'_margin', Twips(style['margins_twips'][side]))
            s._sectPr.find(qn('w:cols')).set(qn('w:space'), '340')
        doc.sections[0].header.paragraphs[0].text = style['journal']
        source = self.root/'source.docx'; doc.save(source)
        return source, SimpleNamespace(blocks=roles)

    def test_matching_document_is_byte_identical_without_hash_allowlist(self):
        source, structure = self.source()
        report = DocumentInspector(source).inspect()
        output = self.root/'result.docx'
        result = preserve_existing_layout(source, output, report, structure)
        self.assertIsNotNone(result)
        self.assertTrue(result.metrics['native_integrity']['passed'])
        self.assertEqual(source.read_bytes(), output.read_bytes())
        self.assertFalse(result.metrics['reference_fidelity']['visual_quality_confirmed'])

    def test_isolated_size_damage_repaired_without_rebuilding_layout(self):
        source, structure = self.source()
        doc = Document(source)
        body = next(p for p in doc.paragraphs if p.text.startswith('Original measurement 7:'))
        body.runs[0].font.size = Pt(16)
        doc.save(source)
        output = self.root/'result.docx'
        result = preserve_existing_layout(source, output, DocumentInspector(source).inspect(), structure)
        self.assertEqual(result.metrics['formatted_paragraphs'], 1)
        with ZipFile(source) as a, ZipFile(output) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != 'word/document.xml': self.assertEqual(a.read(name), b.read(name))
            before, after = (etree.fromstring(z.read('word/document.xml')) for z in (a, b))
            for query in ('//w:t/text()', '//w:sectPr', '//w:tbl', '//w:br'):
                def values(root):
                    return [etree.tostring(n, method='c14n') if isinstance(n, etree._Element) else n
                            for n in root.xpath(query, namespaces=NS)]
                self.assertEqual(values(before), values(after))
        second = self.root/'second.docx'
        preserve_existing_layout(output, second, DocumentInspector(output).inspect(), structure)
        self.assertEqual(output.read_bytes(), second.read_bytes())

    def test_geometry_mismatch_does_not_skip_normal_formatter(self):
        source, structure = self.source()
        doc = Document(source); doc.sections[0].left_margin = Twips(2000); doc.save(source)
        assessment = assess_existing_layout(DocumentInspector(source).inspect(), structure)
        self.assertFalse(assessment['eligible'])
        self.assertIn('page_margins', assessment['reasons'])

    def test_separate_identifiers_cannot_bypass_row_repair_in_matching_layout(self):
        source, structure = self.source()
        doc = Document(source)
        first = doc.paragraphs[0]
        first.insert_paragraph_before('UDC 620.3')
        first.insert_paragraph_before('DOI: 10.9999/layout-test')
        doc.save(source)
        for block in structure.blocks:
            block['id'] = f"block_{int(block['id'].split('_')[-1]) + 2:04d}"
        structure.blocks[:0] = [
            {'id': f'block_{i:04d}', 'detected_role': 'editorial_metadata', 'zone': 'front_matter'}
            for i in (1, 2)
        ]
        report = DocumentInspector(source).inspect()
        assessment = assess_existing_layout(report, structure)
        self.assertEqual(assessment['reasons'], ['identifier_fields_separate'])
        self.assertFalse(assessment['eligible'])
        self.assertIsNone(preserve_existing_layout(source, self.root/'result.docx', report, structure))

    def test_isolated_font_family_change_is_repaired(self):
        source, structure = self.source()
        doc = Document(source)
        paragraph = next(p for p in doc.paragraphs if p.text.startswith('Original measurement 7:'))
        paragraph.runs[0].font.name = 'Arial'; doc.save(source)
        output = self.root/'result.docx'
        result = preserve_existing_layout(source, output, DocumentInspector(source).inspect(), structure)
        self.assertEqual(result.metrics['formatted_paragraphs'], 1)
        paragraph = next(p for p in Document(output).paragraphs if p.text.startswith('Original measurement 7:'))
        self.assertEqual(paragraph.runs[0].font.name, 'Times New Roman')
