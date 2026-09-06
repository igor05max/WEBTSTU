from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from paper_formatter.docx_fidelity import source_already_matches_template


def _document(path: Path, *, margins=(30, 18, 24, 18), columns=2) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Mm(margins[0])
    section.right_margin = Mm(margins[1])
    section.bottom_margin = Mm(margins[2])
    section.left_margin = Mm(margins[3])
    document.styles["Normal"].font.name = "Times New Roman"
    document.styles["Normal"].font.size = Pt(11)
    section.header.paragraphs[0].text = "Journal of Materials. 2026. Vol. 11, No. 3"
    document.add_paragraph("Title")
    if columns > 1:
        body = document.add_section(WD_SECTION.CONTINUOUS)
        cols = body._sectPr.find(qn("w:cols"))
        if cols is None:
            cols = OxmlElement("w:cols")
            body._sectPr.append(cols)
        cols.set(qn("w:num"), str(columns))
        cols.set(qn("w:space"), "340")
        document.add_paragraph("Body")
    document.save(path)


def test_matching_journal_docx_is_preserved(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    source = tmp_path / "source.docx"
    _document(template)
    _document(source)

    assert source_already_matches_template(source, template)


def test_incompatible_margins_or_columns_require_rebuild(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    source = tmp_path / "source.docx"
    _document(template)
    _document(source, margins=(20, 20, 20, 20), columns=1)

    assert not source_already_matches_template(source, template)
