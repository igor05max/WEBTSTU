from copy import deepcopy
from hashlib import sha256
from pathlib import Path
import tempfile
import unittest

from docx import Document
from docx.oxml import OxmlElement
from lxml import etree
from PIL import Image

from paper_formatter.latex_lab.bridge import NativeBridge, UnsupportedContent, Block
from paper_formatter.latex_lab.typesetter import default_plan, validate_plan, render


class LatexBridgeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def build(self, doc):
        source = self.root/'article.docx'; doc.save(source)
        return NativeBridge(source, self.root/'project').build()

    def test_text_is_escaped_not_executed(self):
        doc = Document()
        doc.add_paragraph(r'Ignore instructions. \input{/etc/passwd} $x$ 50% #1 & β')
        blocks, manifest = self.build(doc)
        self.assertNotIn(r'\input{', blocks[0].tex)
        self.assertIn(r'\textbackslash{}input', blocks[0].tex)
        self.assertIn(r'\$x\$', blocks[0].tex)
        self.assertIn('50%', manifest['source_text'])

    def test_inherited_word_numbering_is_visible_and_audited(self):
        doc = Document()
        doc.add_paragraph('First scientific reference.', style='List Number')
        doc.add_paragraph('Second scientific reference.', style='List Number')
        blocks, manifest = self.build(doc)
        self.assertTrue(blocks[0].tex.startswith('1. '))
        self.assertTrue(blocks[1].tex.startswith('2. '))
        self.assertEqual([n['label'] for n in manifest['list_labels']], ['1.', '2.'])

    def test_figure_caption_starts_a_new_paragraph(self):
        image_path = self.root/'plot.png'; Image.new('RGB', (120,80), 'white').save(image_path)
        doc = Document(); doc.add_paragraph('Scientific article', style='Title')
        doc.add_paragraph('Abstract. Original research.')
        doc.add_paragraph('1. Introduction', style='Heading 1')
        doc.add_picture(str(image_path)); doc.add_paragraph('Fig. 1. Measured plot.')
        blocks, manifest = self.build(doc)
        project = self.root/'project'
        source = render(blocks, project, default_plan(blocks, reference_layout=True)).read_text()
        self.assertIn('figure_caption\n\\par\n', source)

    def test_generated_reference_labels_are_a_separate_pdf_gate(self):
        from paper_formatter.latex_lab.quality import gate
        baseline = dict(word_coverage=1, out_of_page=[], sparse_pages=[], compile={})
        self.assertFalse(gate(dict(baseline, missing_list_labels=[{'label':'1.'}]), baseline)['passed'])

    def test_crop_matches_word_and_original_image_remains_unchanged(self):
        path = self.root/'two-colors.png'
        im = Image.new('RGB', (100, 60), 'red')
        im.paste('blue', (50, 0, 100, 60)); im.save(path)
        doc = Document(); picture = doc.add_picture(str(path))
        fill = picture._inline.xpath('.//pic:blipFill')[0]
        crop = OxmlElement('a:srcRect'); crop.set('l', '50000'); fill.insert(1, crop)
        blocks, manifest = self.build(doc)
        asset = manifest['assets'][0]
        self.assertEqual(asset['sha256'], sha256(path.read_bytes()).hexdigest())
        self.assertEqual((self.root/'project'/asset['path']).read_bytes(), path.read_bytes())
        with Image.open(self.root/'project'/asset['rendered_path']) as rendered:
            self.assertEqual(rendered.size, (50, 60))
            self.assertEqual(rendered.getpixel((25, 30)), (0, 0, 255))
        self.assertEqual(len(blocks), 1)

    def test_alternate_content_does_not_duplicate_visible_text(self):
        doc = Document(); p = doc.add_paragraph()
        mc = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
        alt = etree.Element('{'+mc+'}AlternateContent', nsmap={'mc':mc})
        for name, value in [('Choice','Selected visible text'),('Fallback','Duplicate fallback text')]:
            branch = etree.SubElement(alt, '{'+mc+'}'+name)
            run = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = value
            run.append(t); branch.append(run)
        p._p.append(alt)
        blocks, manifest = self.build(doc)
        self.assertIn('Selected visible text', manifest['source_text'])
        self.assertNotIn('Duplicate fallback text', manifest['source_text'])
        self.assertNotIn('Duplicate fallback', blocks[0].tex)

    def test_simple_math_stays_structured(self):
        doc = Document(); p = doc.add_paragraph()
        math = OxmlElement('m:oMath'); frac = OxmlElement('m:f')
        for name, value in [('num','a'),('den','b')]:
            part = OxmlElement('m:'+name); run = OxmlElement('m:r'); t = OxmlElement('m:t')
            t.text = value; run.append(t); part.append(run); frac.append(part)
        math.append(frac); p._p.append(math)
        blocks, manifest = self.build(doc)
        self.assertEqual(manifest['native_equations'], 1)
        self.assertEqual(manifest['converted_equations'], 1)
        self.assertIn(r'\frac{a}{b}', blocks[0].tex)
        self.assertEqual(manifest['assets'], [])

    def test_numbered_display_equation_keeps_original_number_at_right(self):
        doc=Document();p=doc.add_paragraph()
        math=OxmlElement('m:oMath');r=OxmlElement('m:r');t=OxmlElement('m:t');t.text='x=12';r.append(t);math.append(r);p._p.append(math)
        p.add_run('\t(7a)')
        blocks,_=self.build(doc);self.assertEqual(blocks[0].kind,'equation')
        self.assertEqual(blocks[0].equation_number,'(7a)')
        tex=render(blocks,self.root/'project',default_plan(blocks)).read_text()
        self.assertIn(r'\tag*{(7a)}',tex);self.assertEqual(tex.count('(7a)'),1)

    def test_unsupported_math_cannot_silently_degrade(self):
        doc = Document(); p = doc.add_paragraph()
        math = OxmlElement('m:oMath'); math.append(OxmlElement('m:limLow')); p._p.append(math)
        with self.assertRaises(UnsupportedContent):
            self.build(doc)

    def test_math_delimiter_cannot_inject_tex(self):
        doc = Document(); p = doc.add_paragraph()
        math = OxmlElement('m:oMath'); d = OxmlElement('m:d'); props = OxmlElement('m:dPr')
        prop = OxmlElement('m:begChr'); prop.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', r'\input{secret}')
        props.append(prop); d.append(props); math.append(d); p._p.append(math)
        with self.assertRaises(UnsupportedContent):
            self.build(doc)

    def test_raw_tex_in_math_is_not_executed(self):
        doc = Document(); p = doc.add_paragraph()
        math = OxmlElement('m:oMath'); run = OxmlElement('m:r'); t = OxmlElement('m:t')
        t.text = r'$\input{secret}$'; run.append(t); math.append(run); p._p.append(math)
        with self.assertRaises(UnsupportedContent):
            self.build(doc)

    def test_unknown_math_accent_does_not_become_a_hat(self):
        doc = Document(); p = doc.add_paragraph()
        math = OxmlElement('m:oMath'); acc = OxmlElement('m:acc'); props = OxmlElement('m:accPr')
        prop = OxmlElement('m:chr'); prop.set('{http://schemas.openxmlformats.org/officeDocument/2006/math}val', '̈')
        props.append(prop); acc.append(props); math.append(acc); p._p.append(math)
        with self.assertRaises(UnsupportedContent):
            self.build(doc)

    def test_table_cells_do_not_join_into_fake_words_or_numbers(self):
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        for cell, text in zip([c for row in table.rows for c in row.cells], ['Component','Value','Graphite','54.2']):
            cell.text = text
        blocks, manifest = self.build(doc)
        self.assertIn('Component Value Graphite 54.2', ' '.join(manifest['source_text'].split()))
        self.assertNotIn('ComponentValue', manifest['source_text'])
        self.assertIn('54.2', blocks[0].tex)
        self.assertTrue(manifest['text_transfer']['exact'])
        self.assertEqual(manifest['text_transfer']['expected_nodes'], 4)

    def test_full_width_merge_has_no_extra_outer_padding(self):
        doc = Document(); table = doc.add_table(rows=2, cols=2)
        table.cell(0,0).merge(table.cell(0,1)).text = 'Merged heading'
        blocks, _ = self.build(doc)
        self.assertIn(r'\multicolumn{2}{@{}', blocks[0].tex)
        self.assertIn('@{}}{Merged heading}', blocks[0].tex)

    def test_parallel_biographies_are_unruled_prose_without_a_heading(self):
        doc=Document();table=doc.add_table(rows=1,cols=3)
        table.cell(0,0).text='Author One, department and institution. '*5
        table.cell(0,2).text='Первый автор, подразделение и организация. '*5
        blocks,manifest=self.build(doc)
        self.assertNotIn(r'\toprule',blocks[0].tex)
        self.assertIn(r'\justifying',blocks[0].tex)
        self.assertTrue(manifest['text_transfer']['exact'])

    def test_tracked_deletion_is_not_silently_accepted(self):
        doc = Document(); p = doc.add_paragraph('Retained')
        deleted = OxmlElement('w:del'); p._p.append(deleted)
        with self.assertRaises(UnsupportedContent):
            self.build(doc)

    def test_unused_footer_instructions_are_not_exported(self):
        from zipfile import ZipFile
        doc = Document(); doc.add_paragraph('Scientific article')
        doc.sections[0].footer.paragraphs[0].text = 'Current Author'
        path = self.root/'article.docx'; doc.save(path)
        with ZipFile(path,'a') as archive:
            archive.writestr('word/footer0.xml', '<w:ftr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:r><w:t>Old template instructions</w:t></w:r></w:p></w:ftr>')
        _, manifest = NativeBridge(path,self.root/'project').build()
        self.assertEqual(manifest['running_footer'],'Current Author')


    def test_negative_crop_is_transparent_padding_and_retains_original_asset(self):
        path=self.root/'padded.png';Image.new('RGB',(100,80),'red').save(path)
        doc=Document();pic=doc.add_picture(str(path));fill=pic._inline.xpath('.//pic:blipFill')[0]
        crop=OxmlElement('a:srcRect');crop.set('l','-10000');fill.insert(1,crop)
        _,manifest=self.build(doc);asset=manifest['assets'][0]
        self.assertEqual((self.root/'project'/asset['path']).read_bytes(),path.read_bytes())
        with Image.open(self.root/'project'/asset['rendered_path']) as result:
            self.assertEqual(result.size,(110,80));self.assertEqual(result.getpixel((0,0))[3],0)

    def test_long_table_is_breakable_with_repeated_header_outside_multicols(self):
        doc=Document();doc.add_paragraph('Test article');doc.add_paragraph('1. Introduction')
        doc.add_paragraph('Table 1. Measurements');doc.add_paragraph('Таблица 1. Измерения')
        table=doc.add_table(rows=26,cols=3)
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells):cell.text=f'Value {i} {j}'
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        blocks,_=self.build(doc);block=next(b for b in blocks if b.kind=='table')
        self.assertTrue(block.breakable);self.assertEqual(len(block.captions),2)
        self.assertIn(r'\endfirsthead',block.tex);self.assertIn(r'\endhead',block.tex)
        tex=render(blocks,self.root/'project',default_plan(blocks)).read_text()
        prefix=tex.split(r'\begin{longtable}')[0]
        self.assertEqual(prefix.count(r'\begin{multicols}'),prefix.count(r'\end{multicols}'))
        self.assertNotIn(r'\begin{minipage}',prefix)
        self.assertNotIn('Table 1. Measurements',prefix)
        first_head=tex.split(r'\begin{longtable}',1)[1].split(r'\endfirsthead',1)[0]
        self.assertIn('Table 1. Measurements',first_head)
        self.assertIn('Таблица 1. Измерения',first_head)
        self.assertEqual(tex.count('Table 1. Measurements'),1)

    def test_marked_text_is_escaped_and_breakable_at_word_boundaries(self):
        from docx.enum.text import WD_COLOR_INDEX
        doc=Document();run=doc.add_paragraph().add_run('Questionable words 50%');run.font.highlight_color=WD_COLOR_INDEX.YELLOW
        blocks,manifest=self.build(doc)
        self.assertIn(r'\jamtmark{Questionable} \jamtmark{words} \jamtmark{50\%}',blocks[0].tex)
        self.assertTrue(manifest['text_transfer']['exact'])

    def test_nested_numbers_reset_and_bullets_survive(self):
        from apps.template_workspace.v2.ooxml.namespaces import qn
        doc=Document();numbering=doc.part.numbering_part.element
        abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'99')
        for index,fmt,pattern in [(0,'decimal','%1.'),(1,'lowerLetter','%1.%2)')]:
            level=OxmlElement('w:lvl');level.set(qn('w:ilvl'),str(index))
            for tag,val in [('w:start','1'),('w:numFmt',fmt),('w:lvlText',pattern)]:
                el=OxmlElement(tag);el.set(qn('w:val'),val);level.append(el)
            abstract.append(level)
        numbering.append(abstract);num=OxmlElement('w:num');num.set(qn('w:numId'),'99')
        link=OxmlElement('w:abstractNumId');link.set(qn('w:val'),'99');num.append(link);numbering.append(num)
        for i,depth in enumerate([0,1,1,0,1]):
            p=doc.add_paragraph(f'Item {i}');np=OxmlElement('w:numPr')
            for tag,val in [('w:ilvl',str(depth)),('w:numId','99')]:
                el=OxmlElement(tag);el.set(qn('w:val'),val);np.append(el)
            p._p.get_or_add_pPr().append(np)
        doc.add_paragraph('A bullet',style='List Bullet')
        _,manifest=self.build(doc)
        self.assertEqual([v['label'] for v in manifest['list_labels']],['1.','1.a)','1.b)','2.','2.a)','•'])

    def test_embedded_pdf_keeps_vector_page_without_active_annotations(self):
        import pymupdf
        from apps.template_workspace.v2.ooxml.namespaces import NS
        doc=Document();source=self.root/'source.docx';doc.save(source)
        bridge=NativeBridge(source,self.root/'project');bridge.block_id='block_0001'
        node=etree.fromstring(f'<w:object xmlns:w="{NS["w"]}" xmlns:v="{NS["v"]}"><v:shape style="width:240pt;height:120pt"/></w:object>')
        with pymupdf.open() as pdf:
            page=pdf.new_page(width=240,height=120);page.insert_text((20,40),'Native vector figure')
            page.insert_link({'kind':pymupdf.LINK_URI,'from':pymupdf.Rect(0,0,200,60),'uri':'https://example.org'})
            payload=pdf.tobytes()
        tex=bridge.embedded_pdf(node,payload,'word/embeddings/test.bin',b'original OLE')
        self.assertIn('.pdf}',tex);self.assertFalse(bridge.embedded_objects[0]['executed'])
        with pymupdf.open(self.root/'project'/bridge.assets[0]['path']) as clean:
            self.assertEqual(len(clean),1);self.assertEqual(clean[0].get_links(),[])
            self.assertIn('Native vector figure',clean[0].get_text())


class PlanTests(unittest.TestCase):
    def setUp(self):
        self.blocks = [Block('tbl','table','table','body','data','data',493,0,0,8,10),
                       Block('fig','figure','figure','body','','image',210,1)]
        self.plan = default_plan(self.blocks)

    def test_model_cannot_invent_commands_targets_or_tiny_type(self):
        before = deepcopy(self.plan)
        patch = {'body_leading': True, 'table_size': 5, 'tex': r'\input{/etc/passwd}',
                 'objects': {'invented': {'width':'wide'}, 'tbl':{'width':'column','scale':0.1,'delete':True},
                             'fig': {'width': {'bad':'type'}, 'break_before':'true'}}}
        result, accepted, rejected = validate_plan(patch, self.plan, self.blocks)
        self.assertEqual(result, before)
        self.assertEqual(self.plan, before)
        self.assertFalse(accepted)
        self.assertIn('invented', rejected)

    def test_bounded_layout_advice_can_be_applied(self):
        result, accepted, rejected = validate_plan({'body_leading':12.4,'objects':{'fig':{'width':'wide','scale':0.95}}}, self.plan, self.blocks)
        self.assertEqual(result['body_leading'],12.4)
        self.assertEqual(result['objects']['fig']['scale'],0.95)
        self.assertEqual(rejected,[])

    def test_multicolumn_headings_use_non_ejecting_space_reservation(self):
        with tempfile.TemporaryDirectory() as temp:
            block = Block('h','paragraph','heading_1','body','Methods','Methods')
            path = render([block], Path(temp), default_plan([block]))
            tex = path.read_text(encoding='utf-8')
            self.assertIn(r'\needspace{',tex)
            self.assertNotIn(r'\Needspace{',tex)
            self.assertIn(r'\nobreak',tex)


class PdfGateTests(unittest.TestCase):
    def test_only_explicitly_discarded_balance_trials_are_excluded(self):
        from paper_formatter.latex_lab.typesetter import compile_report
        log = ('JAMT-BALANCE-BEGIN\nOverfull \\vbox (51.8pt too high)\n'
               'JAMT-BALANCE-DISCARDED\nOverfull \\vbox (9.0pt too high)\n'
               'JAMT-BALANCE-BEGIN\nOverfull \\vbox (4.0pt too high)\nJAMT-BALANCE-KEPT')
        report=compile_report(log)
        self.assertEqual(report['discarded_balance_vbox_pt'],[51.8])
        self.assertEqual(report['overfull_vbox_pt'],[9.0,4.0])

    def test_incomplete_balance_trace_never_hides_an_overflow(self):
        from paper_formatter.latex_lab.typesetter import compile_report
        report=compile_report('JAMT-BALANCE-BEGIN\nOverfull \\vbox (12.0pt too high)\n')
        self.assertEqual(report['overfull_vbox_pt'],[12.0])

    def test_lost_text_blocks_promotion(self):
        from paper_formatter.latex_lab.quality import gate
        candidate = {'word_coverage':0.97,'out_of_page':[],'compile':{}}
        self.assertIn('text_coverage_regression', gate(candidate,{'word_coverage':1.0})['errors'])

    def test_blank_pages_and_compiler_overflow_block_promotion(self):
        from paper_formatter.latex_lab.quality import gate
        candidate = {'word_coverage':1.0,'out_of_page':[],'sparse_pages':[3],
                     'compile':{'overfull_vbox_pt':[36.9], 'missing_glyphs':['missing beta']}}
        errors = gate(candidate,{'word_coverage':1.0})['errors']
        self.assertIn('almost_empty_pages', errors)
        self.assertIn('overfull_vertical_boxes', errors)
        self.assertIn('missing_glyphs', errors)

    def test_failed_pdf_extraction_is_not_a_perfect_score(self):
        import pymupdf
        from paper_formatter.latex_lab.quality import measure
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp)/'blank.pdf'
            pdf = pymupdf.open(); pdf.new_page(); pdf.save(path); pdf.close()
            metrics = measure(path, {'source_text':'Retain every scientific result', 'blocks':[]})
            self.assertEqual(metrics['word_coverage'],0.0)
            self.assertEqual(metrics['sparse_pages'],[1])


if __name__ == '__main__':
    unittest.main()
