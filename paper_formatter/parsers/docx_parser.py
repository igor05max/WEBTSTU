from __future__ import annotations

import mimetypes
import re
import zipfile
from difflib import SequenceMatcher
from pathlib import Path
from statistics import median
from typing import Iterator, Sequence

from docx import Document
from docx.document import Document as _Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from paper_formatter.exceptions import DocxParseError
from paper_formatter.models import (
    ArticleIR,
    Asset,
    Author,
    Affiliation,
    CitationOccurrence,
    CrossReference,
    EquationBlock,
    FigureBlock,
    ListItemBlock,
    LocalizedText,
    MergedTableCell,
    NoteEntry,
    ParagraphBlock,
    ReferenceEntry,
    SectionBlock,
    SourceTrace,
    TableBlock,
    TextRun,
)
from paper_formatter.parsers.omml_parser import omml_to_latex
from paper_formatter.utils.files import sha256_file
from paper_formatter.utils.images import convert_metafile_to_png
from paper_formatter.utils.text import clean_text, safe_filename
from paper_formatter.semantic.classifier import HybridSemanticClassifier
from paper_formatter.semantic.models import SemanticAnalysis, SemanticBlock, SemanticDecision
from paper_formatter.semantic.rules import RuleSemanticClassifier


_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


class DocxParser:
    def __init__(
        self,
        source_path: Path,
        assets_dir: Path,
        semantic_classifier: HybridSemanticClassifier | None = None,
    ) -> None:
        self.source_path = Path(source_path)
        self.assets_dir = Path(assets_dir)
        self.assets_dir.mkdir(parents=True, exist_ok=True)
        self.semantic_classifier = semantic_classifier
        self.semantic_analysis: SemanticAnalysis | None = None
        self.semantic_blocks: list[SemanticBlock] = []
        self._counters: dict[str, int] = {}
        self._used_asset_names: set[str] = set()
        self._ole_formula_count = 0
        self._ole_formula_converted = 0
        self._ole_formula_failed = 0
        self._ole_conversion_errors: list[str] = []
        self._bookmark_targets: dict[str, str] = {}

    def parse(self) -> ArticleIR:
        self._validate_docx()
        try:
            document = Document(self.source_path)
        except Exception as exc:
            raise DocxParseError(f"Не удалось открыть DOCX: {exc}") from exc

        items = list(self._iter_block_items(document))
        semantic_blocks = self._build_semantic_blocks(items)
        self.semantic_blocks = semantic_blocks
        if self.semantic_classifier is not None:
            self.semantic_analysis = self.semantic_classifier.analyze_document(
                semantic_blocks,
                document_name=self.source_path.name,
            )
        else:
            self.semantic_analysis = RuleSemanticClassifier().analyze_document(
                semantic_blocks,
                document_name=self.source_path.name,
            )
        semantic_by_id = self.semantic_analysis.by_id()
        semantic_blocks_by_id = {
            block.block_id: block for block in semantic_blocks
        }

        article = ArticleIR(semantic_provider=self.semantic_analysis.provider)
        self._extract_notes(article)
        article.warnings.extend(self.semantic_analysis.warnings)
        article.semantic_low_confidence = [
            decision.block_id
            for decision in self.semantic_analysis.decisions
            if decision.confidence < 0.60 or decision.role == "unknown"
        ]

        pending_author_lines: list[str] = []
        pending_caption: str | None = None
        in_references = False
        in_secondary_metadata = False

        for order, item in enumerate(items, start=1):
            location = f"word/document.xml:block[{order}]"
            block_id = f"b-{order}"
            decision = semantic_by_id.get(block_id)

            if isinstance(item, Paragraph):
                text = clean_text(item.text)
                style_name = (item.style.name if item.style is not None else "") or ""
                style_lower = style_name.lower()
                role = decision.role if decision else "paragraph"
                confidence = decision.confidence if decision else 0.5
                has_ole = bool(item._p.xpath(".//w:object | .//*[local-name()='OLEObject']"))

                image_blocks = self._extract_images(item, document, article, location)
                equation_blocks = self._extract_equations(item, location)
                runs = (
                    self._paragraph_runs_with_objects(item, document, article, location)
                    if has_ole
                    else self._paragraph_runs_rich(item, document)
                )
                has_run_content = any(
                    run.text or run.asset_id or run.math_latex is not None
                    for run in runs
                )

                if pending_caption and image_blocks:
                    image_blocks[0].caption = pending_caption
                    pending_caption = None

                if role == "title":
                    if text:
                        article.metadata.titles.append(
                            LocalizedText(language=self._guess_language(text), text=text)
                        )
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if role == "subtitle":
                    if text:
                        article.metadata.subtitles.append(
                            LocalizedText(language=self._guess_language(text), text=text)
                        )
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if role == "author":
                    if text:
                        pending_author_lines.append(
                            self._strip_prefix(text, ("авторы", "автор", "authors", "author"))
                        )
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if role == "affiliation":
                    if text:
                        article.metadata.affiliations.append(
                            self._make_affiliation(article, text)
                        )
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if role == "abstract_heading":
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if role == "abstract":
                    abstract = self._strip_prefix(text, ("аннотация", "abstract"))
                    if abstract:
                        self._append_abstract(article, abstract)
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if role == "keywords":
                    keywords_text = self._strip_prefix(
                        text,
                        ("ключевые слова", "keywords", "key words"),
                    )
                    for keyword in self._split_keywords(keywords_text):
                        if keyword not in article.metadata.keywords:
                            article.metadata.keywords.append(keyword)
                    article.body.extend(equation_blocks)
                    article.body.extend(image_blocks)
                    continue

                if text:
                    udc_match = re.match(r"^\s*УДК\s*[:.]?\s*(.+)$", text, flags=re.IGNORECASE)
                    if udc_match:
                        article.metadata.udc = udc_match.group(1).strip()
                        continue
                    doi_match = re.search(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+", text, flags=re.IGNORECASE)
                    if doi_match and text.lower().startswith("doi"):
                        article.metadata.doi = doi_match.group(0)
                        continue

                if in_references and self._is_secondary_title_paragraph(item, text):
                    article.metadata.titles.append(
                        LocalizedText(language=self._guess_language(text), text=text)
                    )
                    in_references = False
                    in_secondary_metadata = True
                    continue

                if in_secondary_metadata and self._looks_like_secondary_author_line(text):
                    article.metadata.author_variants.append(
                        LocalizedText(language=self._guess_language(text), text=text)
                    )
                    continue

                if role == "references_heading":
                    in_references = True
                    in_secondary_metadata = False
                    continue

                if role == "reference" or (in_references and text and role not in {"section", "subsection", "subsubsection"}):
                    if text:
                        reference_text = re.sub(r"^\s*\d+[.)]\s*", "", text).strip()
                        article.references.append(
                            ReferenceEntry(
                                id=f"ref-{len(article.references) + 1}",
                                text=reference_text or text,
                                doi=self._extract_doi(text),
                                citation_key=f"ref{len(article.references) + 1}",
                            )
                        )
                    continue

                if role in {"section", "subsection", "subsubsection"}:
                    in_references = False
                    in_secondary_metadata = False
                    level = decision.heading_level if decision and decision.heading_level else {
                        "section": 1,
                        "subsection": 2,
                        "subsubsection": 3,
                    }[role]
                    if text:
                        section_id = self._next_id("sec")
                        semantic_block = semantic_blocks_by_id.get(block_id)
                        article.body.append(
                            SectionBlock(
                                id=section_id,
                                title=decision.normalized_text or self._clean_heading_title(text),
                                number=(
                                    semantic_block.numbered_prefix
                                    if semantic_block is not None
                                    else None
                                ),
                                level=level,
                                source=SourceTrace(
                                    format="docx",
                                    location=location,
                                    confidence=confidence,
                                ),
                            )
                        )
                        self._record_bookmarks(item, section_id)
                elif role in {"figure_caption", "table_caption"} or self._is_caption(style_lower, text):
                    caption_text = self._clean_caption(text)
                    if not self._attach_caption(article, caption_text):
                        pending_caption = caption_text
                elif role == "list_item" or self._is_list_item(item):
                    list_runs = runs
                    ordered = self._is_ordered_list(item)
                    if not item._p.xpath("./w:pPr/w:numPr") and re.match(r"^\s*\d+[.)]\s+", text):
                        ordered = True
                        list_runs = self._strip_number_from_runs(runs)
                    list_id = self._next_id("li")
                    article.body.append(
                        ListItemBlock(
                            id=list_id,
                            runs=list_runs,
                            ordered=ordered,
                            level=self._list_level(item),
                            source=SourceTrace(
                                format="docx",
                                location=location,
                                confidence=confidence,
                            ),
                        )
                    )
                    self._record_bookmarks(item, list_id)
                    self._record_run_links(article, list_id, runs, location)
                elif text or has_run_content:
                    paragraph_id = self._next_id("p")
                    article.body.append(
                        ParagraphBlock(
                            id=paragraph_id,
                            runs=runs,
                            source=SourceTrace(
                                format="docx",
                                location=location,
                                confidence=confidence,
                            ),
                        )
                    )
                    self._record_bookmarks(item, paragraph_id)
                    self._record_run_links(article, paragraph_id, runs, location)

                article.body.extend(equation_blocks)
                article.body.extend(image_blocks)

            elif isinstance(item, Table):
                table_figures = self._extract_table_images(
                    item,
                    document,
                    article,
                    location,
                )
                if table_figures:
                    group_label = f"figure-group-{order}"
                    for figure in table_figures:
                        figure.group_id = group_label
                    if pending_caption:
                        table_figures[0].caption = pending_caption
                        pending_caption = None
                    article.body.extend(table_figures)
                    continue
                equation = self._equation_from_table(item, location)
                if equation is not None:
                    article.body.append(equation)
                    continue
                rows = [[clean_text(cell.text) for cell in row.cells] for row in item.rows]
                merged_cells, header_rows = self._table_geometry(item)
                table = TableBlock(
                    id=self._next_id("tbl"),
                    rows=rows,
                    caption=pending_caption,
                    header_rows=header_rows,
                    merged_cells=merged_cells,
                    source=SourceTrace(format="docx", location=location),
                )
                pending_caption = None
                article.body.append(table)

        self._save_authors(article, pending_author_lines)
        self._postprocess_metadata(article)
        self._deduplicate_article(article)
        self._append_ole_summary(article)
        self._resolve_bookmarks(article)
        self._link_numeric_citations(article)

        if pending_caption:
            article.warnings.append(
                f"Подпись не удалось связать с рисунком или таблицей: {pending_caption}"
            )
        if not article.metadata.titles:
            article.warnings.append("Название статьи не распознано автоматически.")
        if not article.metadata.authors:
            article.warnings.append("Авторы статьи не распознаны автоматически.")
        if article.semantic_low_confidence:
            article.warnings.append(
                "Остались спорные структурные блоки: "
                + ", ".join(article.semantic_low_confidence[:20])
                + (" …" if len(article.semantic_low_confidence) > 20 else "")
            )
        if not article.body:
            article.warnings.append("В DOCX не обнаружены содержательные блоки.")

        return article

    def _build_semantic_blocks(
        self,
        items: Sequence[Paragraph | Table],
    ) -> list[SemanticBlock]:
        blocks: list[SemanticBlock] = []
        for order, item in enumerate(items, start=1):
            if not isinstance(item, Paragraph):
                continue
            text = clean_text(item.text)
            style = (item.style.name if item.style is not None else "") or ""
            font_size = self._paragraph_font_size(item)
            bold_ratio, italic_ratio = self._format_ratios(item)
            alignment = self._alignment_name(item)
            outline_level = self._outline_level(item)
            has_numbering, numbering_level = self._numbering_info(item)
            prefix_match = re.match(r"^\s*(\d+(?:\.\d+)*)[.)]?\s+\S", text)
            blocks.append(
                SemanticBlock(
                    block_id=f"b-{order}",
                    order=order,
                    text=text,
                    style=style,
                    font_size_pt=font_size,
                    bold_ratio=bold_ratio,
                    italic_ratio=italic_ratio,
                    alignment=alignment,
                    outline_level=outline_level,
                    numbered_prefix=prefix_match.group(1) if prefix_match else None,
                    numbering_level=numbering_level,
                    has_numbering=has_numbering,
                )
            )

        for index, block in enumerate(blocks):
            block.previous_text = blocks[index - 1].text if index > 0 else ""
            block.next_text = blocks[index + 1].text if index + 1 < len(blocks) else ""

        self._mark_numbered_sequences(blocks)
        return blocks

    def _mark_numbered_sequences(self, blocks: list[SemanticBlock]) -> None:
        index = 0
        while index < len(blocks):
            current = blocks[index]
            if not current.numbered_prefix or "." in current.numbered_prefix:
                index += 1
                continue
            run = [current]
            expected = int(current.numbered_prefix) + 1
            cursor = index + 1
            while cursor < len(blocks):
                candidate = blocks[cursor]
                if (
                    candidate.order - run[-1].order > 2
                    or candidate.style.lower() != current.style.lower()
                    or not candidate.numbered_prefix
                    or "." in candidate.numbered_prefix
                ):
                    break
                try:
                    numeric = int(candidate.numbered_prefix)
                except ValueError:
                    break
                if numeric != expected:
                    break
                if len(candidate.text) > 260:
                    break
                run.append(candidate)
                expected += 1
                cursor += 1
            if len(run) >= 3:
                for item in run:
                    item.is_in_numbered_sequence = True
                index = cursor
            else:
                index += 1

    def _paragraph_font_size(self, paragraph: Paragraph) -> float | None:
        sizes: list[float] = []
        for run in paragraph.runs:
            if run.text and run.font.size is not None:
                sizes.append(float(run.font.size.pt))
        if sizes:
            return float(median(sizes))
        if paragraph.style is not None and paragraph.style.font.size is not None:
            return float(paragraph.style.font.size.pt)
        return None

    def _format_ratios(self, paragraph: Paragraph) -> tuple[float, float]:
        total = 0
        bold = 0
        italic = 0
        style_bold = bool(paragraph.style.font.bold) if paragraph.style is not None else False
        style_italic = bool(paragraph.style.font.italic) if paragraph.style is not None else False
        for run in paragraph.runs:
            length = len(run.text or "")
            if length == 0:
                continue
            total += length
            if run.bold is True or (run.bold is None and style_bold):
                bold += length
            if run.italic is True or (run.italic is None and style_italic):
                italic += length
        if total == 0:
            return (1.0 if style_bold else 0.0, 1.0 if style_italic else 0.0)
        return bold / total, italic / total

    def _alignment_name(self, paragraph: Paragraph) -> str | None:
        value = paragraph.alignment
        if value is None:
            return None
        name = getattr(value, "name", None)
        if name:
            return str(name).lower()
        numeric = int(value)
        return {0: "left", 1: "center", 2: "right", 3: "justify"}.get(numeric, str(value).lower())

    def _outline_level(self, paragraph: Paragraph) -> int | None:
        p_pr = paragraph._p.pPr
        if p_pr is None:
            return None
        node = p_pr.find(qn("w:outlineLvl"))
        if node is None:
            return None
        raw = node.get(qn("w:val"))
        try:
            return max(1, min(6, int(raw) + 1))
        except (TypeError, ValueError):
            return None

    def _numbering_info(self, paragraph: Paragraph) -> tuple[bool, int | None]:
        p_pr = paragraph._p.pPr
        if p_pr is None or p_pr.numPr is None:
            return False, None
        level: int | None = None
        if p_pr.numPr.ilvl is not None:
            try:
                level = int(p_pr.numPr.ilvl.val)
            except (TypeError, ValueError):
                level = None
        return True, level

    def _make_affiliation(self, article: ArticleIR, text: str) -> Affiliation:
        return Affiliation(
            id=f"affiliation-{len(article.metadata.affiliations) + 1}",
            name=text,
        )

    def _append_abstract(self, article: ArticleIR, text: str) -> None:
        language = self._guess_language(text)
        if article.metadata.abstracts and article.metadata.abstracts[-1].language == language:
            current = article.metadata.abstracts[-1]
            if self._normalized_text(text) not in self._normalized_text(current.text):
                current.text = clean_text(current.text + " " + text)
            return
        article.metadata.abstracts.append(LocalizedText(language=language, text=text))

    def _strip_number_from_runs(self, runs: list[TextRun]) -> list[TextRun]:
        result = [run.model_copy(deep=True) for run in runs]
        prefix_pending = True
        for run in result:
            if not prefix_pending or not run.text:
                continue
            cleaned, count = re.subn(r"^\s*\d+[.)]\s+", "", run.text, count=1)
            if count:
                run.text = cleaned
                prefix_pending = False
                break
            if run.text.strip():
                prefix_pending = False
        return [run for run in result if run.text or run.asset_id]

    def _deduplicate_article(self, article: ArticleIR) -> None:
        abstract_values = [self._normalized_text(item.text) for item in article.metadata.abstracts]
        if abstract_values:
            kept: list = []
            removed = 0
            for block in article.body:
                if isinstance(block, ParagraphBlock):
                    value = self._normalized_text(block.text)
                    if len(value) >= 180 and any(
                        value == abstract
                        or value in abstract
                        or abstract in value
                        or SequenceMatcher(None, value, abstract).ratio() >= 0.88
                        for abstract in abstract_values
                        if abstract
                    ):
                        removed += 1
                        continue
                kept.append(block)
            if removed:
                article.body = kept
                article.warnings.append(
                    f"Удалено повторов аннотации из основного текста: {removed}."
                )

        seen: set[str] = set()
        unique_references: list[ReferenceEntry] = []
        for reference in article.references:
            key = self._normalized_text(reference.text)
            if key and key not in seen:
                seen.add(key)
                unique_references.append(reference)
        if len(unique_references) != len(article.references):
            article.warnings.append("Повторяющиеся источники в списке литературы объединены.")
            article.references = unique_references
            for index, reference in enumerate(article.references, start=1):
                reference.id = f"ref-{index}"

    @staticmethod
    def _normalized_text(text: str) -> str:
        return re.sub(r"[^0-9a-zа-яё]+", " ", (text or "").lower()).strip()

    def _validate_docx(self) -> None:
        if not self.source_path.exists():
            raise DocxParseError(f"Файл не найден: {self.source_path}")
        if self.source_path.suffix.lower() != ".docx":
            raise DocxParseError("Ожидается файл с расширением .docx")
        if not zipfile.is_zipfile(self.source_path):
            raise DocxParseError("DOCX не является корректным OOXML ZIP-пакетом")
        with zipfile.ZipFile(self.source_path) as archive:
            required = {"[Content_Types].xml", "word/document.xml"}
            missing = required.difference(archive.namelist())
            if missing:
                raise DocxParseError(f"В DOCX отсутствуют обязательные файлы: {sorted(missing)}")
            bad_file = archive.testzip()
            if bad_file:
                raise DocxParseError(f"Повреждён файл внутри DOCX: {bad_file}")

    def _iter_block_items(self, parent: _Document | _Cell) -> Iterator[Paragraph | Table]:
        if isinstance(parent, _Document):
            parent_element = parent.element.body
            parent_object = parent
        elif isinstance(parent, _Cell):
            parent_element = parent._tc
            parent_object = parent
        else:
            raise TypeError("Неподдерживаемый родитель DOCX")

        for child in parent_element.iterchildren():
            local_name = etree.QName(child).localname
            if local_name == "p":
                yield Paragraph(child, parent_object)
            elif local_name == "tbl":
                yield Table(child, parent_object)

    def _next_id(self, prefix: str) -> str:
        self._counters[prefix] = self._counters.get(prefix, 0) + 1
        return f"{prefix}-{self._counters[prefix]}"

    def _paragraph_runs(self, paragraph: Paragraph) -> list[TextRun]:
        runs: list[TextRun] = []
        for run in paragraph.runs:
            text = run.text
            if not text:
                continue
            font = run.font
            runs.append(
                TextRun(
                    text=text,
                    bold=bool(run.bold),
                    italic=bool(run.italic),
                    underline=bool(run.underline),
                    superscript=bool(font.superscript),
                    subscript=bool(font.subscript),
                )
            )
        if not runs and paragraph.text:
            runs.append(TextRun(text=paragraph.text))
        return runs

    def _paragraph_runs_rich(
        self,
        paragraph: Paragraph,
        document: _Document,
    ) -> list[TextRun]:
        """Читает также hyperlink и поля REF/CITATION/HYPERLINK."""
        runs: list[TextRun] = []
        field_instruction: list[str] | None = None
        field_display: list[str] = []
        in_field_display = False

        for child in paragraph._p.iterchildren():
            local_name = etree.QName(child).localname
            if local_name == "hyperlink":
                rel_id = child.get(qn("r:id"))
                anchor = child.get(qn("w:anchor"))
                target: str | None = None
                if rel_id and rel_id in document.part.rels:
                    target = str(document.part.rels[rel_id].target_ref)
                elif anchor:
                    target = f"#{anchor}"
                for run_element in child.xpath(".//*[local-name()='r']"):
                    rich_run = self._text_run_from_element(run_element, paragraph)
                    if rich_run:
                        rich_run.hyperlink = target
                        if anchor:
                            rich_run.reference_target = anchor
                        runs.append(rich_run)
                continue
            if local_name == "fldSimple":
                instruction = child.get(qn("w:instr"), "")
                display = "".join(child.xpath(".//*[local-name()='t']/text()"))
                runs.append(self._field_run(instruction, display))
                continue
            if local_name == "oMath":
                latex = omml_to_latex(child)
                if latex:
                    runs.append(TextRun(math_latex=latex))
                continue
            if local_name != "r":
                if local_name in {"oMathPara", "object", "drawing", "pict"}:
                    continue
                for nested in child.xpath(
                    ".//*[local-name()='r' or local-name()='oMath']"
                ):
                    nested_name = etree.QName(nested).localname
                    if nested_name == "oMath":
                        latex = omml_to_latex(nested)
                        if latex:
                            runs.append(TextRun(math_latex=latex))
                    elif not nested.xpath("ancestor::*[local-name()='oMath']"):
                        rich_run = self._text_run_from_element(nested, paragraph)
                        if rich_run:
                            runs.append(rich_run)
                continue

            field_types = [
                element.get(qn("w:fldCharType"), "")
                for element in child.xpath("./*[local-name()='fldChar']")
            ]
            if "begin" in field_types:
                field_instruction = []
                field_display = []
                in_field_display = False
                continue
            instruction_text = "".join(
                child.xpath("./*[local-name()='instrText']/text()")
            )
            if field_instruction is not None and instruction_text:
                field_instruction.append(instruction_text)
                continue
            if "separate" in field_types and field_instruction is not None:
                in_field_display = True
                continue
            if "end" in field_types and field_instruction is not None:
                runs.append(
                    self._field_run(
                        "".join(field_instruction),
                        "".join(field_display),
                    )
                )
                field_instruction = None
                field_display = []
                in_field_display = False
                continue
            if field_instruction is not None and in_field_display:
                field_display.extend(child.xpath(".//*[local-name()='t']/text()"))
                continue

            rich_run = self._text_run_from_element(child, paragraph)
            if rich_run:
                runs.append(rich_run)
            for kind in ("footnote", "endnote"):
                note_elements = child.xpath(
                    f".//*[local-name()='{kind}Reference']"
                )
                for note_element in note_elements:
                    note_id = note_element.get(qn("w:id"))
                    if note_id is None:
                        continue
                    runs.append(
                        TextRun(
                            text=f"[{note_id}]",
                            reference_target=f"{kind}-{note_id}",
                            superscript=True,
                        )
                    )
        if field_instruction is not None:
            runs.append(self._field_run("".join(field_instruction), "".join(field_display)))
        if not runs and paragraph.text:
            runs.append(TextRun(text=paragraph.text))
        return runs

    @staticmethod
    def _text_run_from_element(
        run_element: etree._Element,
        paragraph: Paragraph,
    ) -> TextRun | None:
        fragments: list[str] = []
        for child in run_element.iterchildren():
            local_name = etree.QName(child).localname
            if local_name in {"t", "delText"} and child.text:
                fragments.append(child.text)
            elif local_name == "tab":
                fragments.append("    ")
            elif local_name in {"br", "cr"}:
                fragments.append("\n")
        text = "".join(fragments)
        if not text:
            return None
        properties = run_element.xpath("./*[local-name()='rPr']")
        property_node = properties[0] if properties else None

        def has_property(name: str) -> bool:
            if property_node is None:
                return False
            nodes = property_node.xpath(f"./*[local-name()='{name}']")
            if not nodes:
                return False
            value = nodes[0].get(qn("w:val"))
            return value not in {"0", "false", "off", "none"}

        vertical = None
        if property_node is not None:
            nodes = property_node.xpath("./*[local-name()='vertAlign']")
            if nodes:
                vertical = nodes[0].get(qn("w:val"))
        return TextRun(
            text=text,
            bold=has_property("b"),
            italic=has_property("i"),
            underline=has_property("u"),
            superscript=vertical == "superscript",
            subscript=vertical == "subscript",
        )

    @staticmethod
    def _field_run(instruction: str, display: str) -> TextRun:
        instruction = re.sub(r"\s+", " ", instruction).strip()
        display = display or instruction
        hyperlink = re.search(r'HYPERLINK\s+"([^"]+)"', instruction, re.IGNORECASE)
        reference = re.search(
            r"\b(?:REF|PAGEREF)\s+([^\s\\]+)", instruction, re.IGNORECASE
        )
        citation = re.search(r"\bCITATION\s+([^\s\\]+)", instruction, re.IGNORECASE)
        return TextRun(
            text=display,
            hyperlink=hyperlink.group(1) if hyperlink else None,
            reference_target=reference.group(1) if reference else None,
            citation_keys=[citation.group(1)] if citation else [],
        )

    def _record_run_links(
        self,
        article: ArticleIR,
        block_id: str,
        runs: list[TextRun],
        location: str,
    ) -> None:
        for run in runs:
            if run.citation_keys:
                article.citations.append(
                    CitationOccurrence(
                        id=f"citation-{len(article.citations) + 1}",
                        keys=run.citation_keys,
                        raw_text=run.text,
                        source_block_id=block_id,
                        source=SourceTrace(format="docx_field", location=location),
                    )
                )
            if run.reference_target:
                article.cross_references.append(
                    CrossReference(
                        id=f"xref-{len(article.cross_references) + 1}",
                        target_id=run.reference_target,
                        raw_text=run.text,
                        source_block_id=block_id,
                        source=SourceTrace(format="docx_field", location=location),
                    )
                )

    def _record_bookmarks(self, paragraph: Paragraph, block_id: str) -> None:
        for bookmark in paragraph._p.xpath(".//w:bookmarkStart"):
            name = bookmark.get(qn("w:name"))
            if name and not name.startswith("_"):
                self._bookmark_targets[name] = block_id

    def _resolve_bookmarks(self, article: ArticleIR) -> None:
        blocks = {block.id: block for block in article.body}
        for reference in article.cross_references:
            if reference.target_id in self._bookmark_targets:
                reference.target_id = self._bookmark_targets[reference.target_id]
            target = blocks.get(reference.target_id or "")
            if isinstance(target, EquationBlock):
                reference.target_kind = "equation"
            elif isinstance(target, FigureBlock):
                reference.target_kind = "figure"
            elif isinstance(target, TableBlock):
                reference.target_kind = "table"
            elif isinstance(target, SectionBlock):
                reference.target_kind = "section"

    def _link_numeric_citations(self, article: ArticleIR) -> None:
        if not article.references:
            return
        for index, reference in enumerate(article.references, start=1):
            if not reference.citation_key:
                reference.citation_key = f"ref{index}"
        maximum = len(article.references)
        pattern = re.compile(
            r"\[(\s*\d+(?:\s*(?:[,;–—-])\s*\d+)*)\]"
        )
        for block in article.body:
            if not isinstance(block, (ParagraphBlock, ListItemBlock)):
                continue
            updated_runs: list[TextRun] = []
            for run in block.runs:
                if (
                    run.asset_id
                    or run.math_latex is not None
                    or run.citation_keys
                    or run.reference_target
                ):
                    updated_runs.append(run)
                    continue
                cursor = 0
                changed = False
                for match in pattern.finditer(run.text):
                    numbers = self._citation_numbers(match.group(1), maximum)
                    if not numbers:
                        continue
                    changed = True
                    if match.start() > cursor:
                        updated_runs.append(
                            run.model_copy(
                                update={"text": run.text[cursor : match.start()]}
                            )
                        )
                    keys = [f"ref{number}" for number in numbers]
                    citation_run = run.model_copy(
                        update={
                            "text": match.group(0),
                            "citation_keys": keys,
                            "hyperlink": None,
                            "reference_target": None,
                        }
                    )
                    updated_runs.append(citation_run)
                    article.citations.append(
                        CitationOccurrence(
                            id=f"citation-{len(article.citations) + 1}",
                            keys=keys,
                            raw_text=match.group(0),
                            source_block_id=block.id,
                            source=block.source,
                        )
                    )
                    cursor = match.end()
                if changed:
                    if cursor < len(run.text):
                        updated_runs.append(
                            run.model_copy(update={"text": run.text[cursor:]})
                        )
                else:
                    updated_runs.append(run)
            block.runs = updated_runs

    @staticmethod
    def _citation_numbers(value: str, maximum: int) -> list[int]:
        normalized = value.replace("–", "-").replace("—", "-")
        result: list[int] = []
        for part in re.split(r"[,;]", normalized):
            part = part.strip()
            if not part:
                continue
            if "-" in part:
                left, right = (item.strip() for item in part.split("-", 1))
                if not left.isdigit() or not right.isdigit():
                    return []
                start, end = int(left), int(right)
                if start > end or end - start > 50:
                    return []
                result.extend(range(start, end + 1))
            elif part.isdigit():
                result.append(int(part))
            else:
                return []
        if not result or any(number < 1 or number > maximum for number in result):
            return []
        return list(dict.fromkeys(result))

    def _extract_notes(self, article: ArticleIR) -> None:
        parts = {
            "word/footnotes.xml": ("footnote", "footnote"),
            "word/endnotes.xml": ("endnote", "endnote"),
            "word/comments.xml": ("comment", "comment"),
        }
        try:
            with zipfile.ZipFile(self.source_path) as archive:
                names = set(archive.namelist())
                for part_name, (element_name, kind) in parts.items():
                    if part_name not in names:
                        continue
                    root = etree.fromstring(archive.read(part_name))
                    for node in root.xpath(f"./*[local-name()='{element_name}']"):
                        note_id = node.get(qn("w:id"))
                        if note_id is None or note_id.startswith("-"):
                            continue
                        text = clean_text(" ".join(node.xpath(".//*[local-name()='t']/text()")))
                        if text:
                            article.notes.append(
                                NoteEntry(
                                    id=f"{kind}-{note_id}",
                                    kind=kind,
                                    text=text,
                                    source=SourceTrace(
                                        format="docx",
                                        part=part_name,
                                        location=f"{part_name}:{element_name}[{note_id}]",
                                    ),
                                )
                            )
        except Exception as exc:
            article.warnings.append(f"Не удалось прочитать сноски/комментарии: {exc}")

    @staticmethod
    def _table_geometry(table: Table) -> tuple[list[MergedTableCell], int]:
        merged: list[MergedTableCell] = []
        active_vertical: dict[int, MergedTableCell] = {}
        header_rows = 0
        for row_index, row in enumerate(table._tbl.tr_lst):
            if row.xpath("./w:trPr/w:tblHeader") and header_rows == row_index:
                header_rows += 1
            column = 0
            for cell in row.tc_lst:
                grid_span_nodes = cell.xpath("./w:tcPr/w:gridSpan")
                grid_span = 1
                if grid_span_nodes:
                    value = grid_span_nodes[0].get(qn("w:val"))
                    if value and value.isdigit():
                        grid_span = max(1, int(value))
                vertical_nodes = cell.xpath("./w:tcPr/w:vMerge")
                vertical_value = (
                    vertical_nodes[0].get(qn("w:val"), "continue")
                    if vertical_nodes
                    else None
                )
                record: MergedTableCell | None = None
                if grid_span > 1:
                    record = MergedTableCell(
                        row=row_index,
                        column=column,
                        column_span=grid_span,
                    )
                    merged.append(record)
                if vertical_value == "restart":
                    record = record or MergedTableCell(row=row_index, column=column)
                    if record not in merged:
                        merged.append(record)
                    active_vertical[column] = record
                elif vertical_value == "continue" and column in active_vertical:
                    active_vertical[column].row_span += 1
                elif vertical_value is None:
                    active_vertical.pop(column, None)
                column += grid_span
        return merged, header_rows

    def _paragraph_runs_with_objects(
        self,
        paragraph: Paragraph,
        document: _Document,
        article: ArticleIR,
        location: str,
    ) -> list[TextRun]:
        runs: list[TextRun] = []
        formula_display = not bool(clean_text(paragraph.text))

        for run_element in paragraph._p.xpath(".//w:r"):
            word_run = Run(run_element, paragraph)
            font = word_run.font
            formatting = {
                "bold": bool(word_run.bold),
                "italic": bool(word_run.italic),
                "underline": bool(word_run.underline),
                "superscript": bool(font.superscript),
                "subscript": bool(font.subscript),
            }

            for child in run_element.iterchildren():
                local_name = etree.QName(child).localname
                if local_name in {"t", "instrText"}:
                    if child.text:
                        runs.append(TextRun(text=child.text, **formatting))
                elif local_name == "tab":
                    runs.append(TextRun(text="    ", **formatting))
                elif local_name in {"br", "cr"}:
                    runs.append(TextRun(text="\n", **formatting))
                elif local_name == "object":
                    formula_run = self._extract_formula_image_run(
                        child,
                        document,
                        article,
                        location,
                        display=formula_display,
                    )
                    if formula_run is not None:
                        runs.append(formula_run)

        if not runs and paragraph.text:
            runs.append(TextRun(text=paragraph.text))
        return runs

    def _extract_formula_image_run(
        self,
        object_element: etree._Element,
        document: _Document,
        article: ArticleIR,
        location: str,
        display: bool,
    ) -> TextRun | None:
        self._ole_formula_count += 1
        image_nodes = object_element.xpath(".//*[local-name()='imagedata']")
        if not image_nodes:
            self._ole_formula_failed += 1
            return None

        rel_id = image_nodes[0].get(qn("r:id"))
        if not rel_id or rel_id not in document.part.rels:
            self._ole_formula_failed += 1
            return None

        relationship = document.part.rels[rel_id]
        try:
            part = relationship.target_part
            blob = part.blob
            original_name = Path(str(part.partname)).name
        except Exception as exc:
            self._ole_formula_failed += 1
            self._ole_conversion_errors.append(f"{location}: {exc}")
            return None

        raw_name = self._unique_asset_name(original_name)
        raw_target = self.assets_dir / raw_name
        raw_target.write_bytes(blob)
        final_target = raw_target
        final_name = raw_name
        media_type = getattr(part, "content_type", None) or mimetypes.guess_type(raw_name)[0]
        raw_media_type = media_type

        ole_object_xml, ole_object_asset_id, ole_preview_asset_id = self._preserved_math_ole(
            object_element,
            document,
            article,
            raw_target=raw_target,
            raw_media_type=raw_media_type,
        )
        preserved_preview = next(
            (
                candidate
                for candidate in article.assets
                if candidate.id == ole_preview_asset_id
            ),
            None,
        )

        if raw_target.suffix.lower() in {".wmf", ".emf"}:
            if ole_object_asset_id and preserved_preview is not None:
                # The original MathType payload plus its native WMF/EMF preview
                # is the most faithful and portable representation.  Avoid an
                # expensive best-effort raster conversion that Linux often
                # cannot perform and that the renderer no longer needs.
                self._ole_formula_converted += 1
            else:
                png_name = self._unique_asset_name(f"{raw_target.stem}.png")
                png_target = self.assets_dir / png_name
                converted, error = convert_metafile_to_png(raw_target, png_target)
                if converted:
                    final_target = png_target
                    final_name = png_name
                    media_type = "image/png"
                    self._ole_formula_converted += 1
                else:
                    self._ole_formula_failed += 1
                    if error and len(self._ole_conversion_errors) < 5:
                        self._ole_conversion_errors.append(f"{location}: {error}")
        else:
            self._ole_formula_converted += 1

        if final_target == raw_target and preserved_preview is not None:
            asset = preserved_preview
        else:
            asset = Asset(
                id=self._next_id("asset"),
                path=f"assets/{final_name}",
                media_type=media_type,
                original_name=original_name,
                sha256=sha256_file(final_target),
            )
            article.assets.append(asset)
        if final_target != raw_target and ole_preview_asset_id is None:
            raw_target.unlink(missing_ok=True)
        width_pt, height_pt = self._shape_size_points(object_element)
        return TextRun(
            asset_id=asset.id,
            formula_image=True,
            ole_object_xml=ole_object_xml,
            ole_object_asset_id=ole_object_asset_id,
            ole_preview_asset_id=ole_preview_asset_id,
            display=display,
            width_pt=width_pt,
            height_pt=height_pt,
        )

    def _preserved_math_ole(
        self,
        object_element: etree._Element,
        document: _Document,
        article: ArticleIR,
        *,
        raw_target: Path,
        raw_media_type: str | None,
    ) -> tuple[str | None, str | None, str | None]:
        """Preserve a known MathType stream for Word-native equation rendering."""
        ole_nodes = object_element.xpath(".//*[local-name()='OLEObject']")
        if not ole_nodes:
            return None, None, None
        ole_node = ole_nodes[0]
        prog_id = (ole_node.get("ProgID") or "").lower()
        if not (prog_id.startswith("equation.") or "mathtype" in prog_id):
            return None, None, None
        rel_id = ole_node.get(qn("r:id"))
        if not rel_id or rel_id not in document.part.rels:
            return None, None, None
        try:
            part = document.part.rels[rel_id].target_part
            original_name = Path(str(part.partname)).name
            target = self.assets_dir / self._unique_asset_name(original_name)
            target.write_bytes(part.blob)
        except Exception as exc:
            if len(self._ole_conversion_errors) < 5:
                self._ole_conversion_errors.append(f"OLE payload: {exc}")
            return None, None, None
        ole_asset = Asset(
            id=self._next_id("asset"),
            path=f"assets/{target.name}",
            media_type=getattr(part, "content_type", None)
            or "application/vnd.openxmlformats-officedocument.oleObject",
            original_name=original_name,
            sha256=sha256_file(target),
        )
        article.assets.append(ole_asset)
        preview_asset = Asset(
            id=self._next_id("asset"),
            path=f"assets/{raw_target.name}",
            media_type=raw_media_type,
            original_name=raw_target.name,
            sha256=sha256_file(raw_target),
        )
        article.assets.append(preview_asset)
        return etree.tostring(object_element, encoding="unicode"), ole_asset.id, preview_asset.id

    def _shape_size_points(self, object_element: etree._Element) -> tuple[float | None, float | None]:
        shapes = object_element.xpath(".//*[local-name()='shape']")
        if not shapes:
            return None, None
        style = shapes[0].get("style", "")

        def read_value(name: str) -> float | None:
            match = re.search(rf"(?:^|;)\s*{name}\s*:\s*([0-9.]+)pt", style, flags=re.IGNORECASE)
            return float(match.group(1)) if match else None

        return read_value("width"), read_value("height")

    def _paragraph_block(self, paragraph: Paragraph, location: str) -> ParagraphBlock:
        return ParagraphBlock(
            id=self._next_id("p"),
            runs=self._paragraph_runs(paragraph),
            source=SourceTrace(format="docx", location=location),
        )

    def _extract_equations(self, paragraph: Paragraph, location: str) -> list[EquationBlock]:
        equations: list[EquationBlock] = []
        root = paragraph._p
        # Формулы в смешанном абзаце остаются inline-элементами TextRun. В отдельные
        # EquationBlock выносим только display-формулы и полностью формульные абзацы.
        candidates = list(root.xpath("./m:oMathPara"))
        if not clean_text(paragraph.text):
            candidates.extend(
                root.xpath("./m:oMath | ./w:r/m:oMath | ./w:smartTag/m:oMath")
            )
        seen: set[int] = set()
        for element in candidates:
            marker = id(element)
            if marker in seen:
                continue
            seen.add(marker)
            latex = omml_to_latex(element)
            if not latex:
                continue
            equations.append(
                EquationBlock(
                    id=self._next_id("eq"),
                    latex=latex,
                    display=etree.QName(element).localname == "oMathPara" or not clean_text(paragraph.text),
                    source=SourceTrace(
                        format="docx_omml",
                        location=location,
                        confidence=0.9,
                    ),
                )
            )
        return equations

    def _extract_table_images(
        self,
        table: Table,
        document: _Document,
        article: ArticleIR,
        location: str,
    ) -> list[FigureBlock]:
        figures: list[FigureBlock] = []
        seen_cells: set[int] = set()
        for row_index, row in enumerate(table.rows, start=1):
            for column_index, cell in enumerate(row.cells, start=1):
                marker = id(cell._tc)
                if marker in seen_cells:
                    continue
                seen_cells.add(marker)
                for paragraph_index, paragraph in enumerate(
                    cell.paragraphs,
                    start=1,
                ):
                    figures.extend(
                        self._extract_images(
                            paragraph,
                            document,
                            article,
                            (
                                f"{location}:cell[{row_index},{column_index}]"
                                f":paragraph[{paragraph_index}]"
                            ),
                        )
                    )
        return figures

    def _equation_from_table(
        self,
        table: Table,
        location: str,
    ) -> EquationBlock | None:
        if len(table.rows) != 1 or len(table.columns) not in {1, 2}:
            return None
        first_cell = table.cell(0, 0)
        display_nodes = first_cell._tc.xpath(".//m:oMathPara")
        math_nodes = display_nodes or first_cell._tc.xpath(".//m:oMath")
        if not math_nodes:
            return None
        latex = omml_to_latex(math_nodes[0])
        if not latex:
            return None
        number: str | None = None
        if len(table.columns) == 2:
            match = re.fullmatch(
                r"\s*\(([^()]+)\)\s*",
                clean_text(table.cell(0, 1).text),
            )
            if match is None:
                return None
            number = match.group(1).strip()
        return EquationBlock(
            id=self._next_id("eq"),
            latex=latex,
            number=number,
            display=True,
            source=SourceTrace(
                format="docx_omml_table",
                location=location,
                confidence=0.95,
            ),
        )

    def _extract_images(
        self,
        paragraph: Paragraph,
        document: _Document,
        article: ArticleIR,
        location: str,
    ) -> list[FigureBlock]:
        figures: list[FigureBlock] = []
        blips = paragraph._p.xpath(".//a:blip")
        for blip in blips:
            rel_id = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if not rel_id or rel_id not in document.part.rels:
                article.warnings.append(f"Не удалось определить связь изображения в {location}.")
                continue
            relationship = document.part.rels[rel_id]
            try:
                part = relationship.target_part
                blob = part.blob
                original_name = Path(str(part.partname)).name
            except Exception:
                article.warnings.append(f"Не удалось извлечь изображение {rel_id} в {location}.")
                continue

            filename = self._unique_asset_name(original_name)
            target = self.assets_dir / filename
            target.write_bytes(blob)
            asset = Asset(
                id=self._next_id("asset"),
                path=f"assets/{filename}",
                media_type=getattr(part, "content_type", None) or mimetypes.guess_type(filename)[0],
                original_name=original_name,
                sha256=sha256_file(target),
            )
            article.assets.append(asset)
            if Path(filename).suffix.lower() not in {".png", ".jpg", ".jpeg", ".pdf"}:
                article.warnings.append(
                    f"Формат изображения {Path(filename).suffix or 'без расширения'} "
                    "может не поддерживаться XeLaTeX без предварительной конвертации."
                )
            figures.append(
                FigureBlock(
                    id=self._next_id("fig"),
                    asset_id=asset.id,
                    source=SourceTrace(format="docx_image", location=location),
                )
            )
        return figures

    def _unique_asset_name(self, original_name: str) -> str:
        base = safe_filename(original_name, fallback="image.bin")
        stem = Path(base).stem
        suffix = Path(base).suffix
        candidate = base
        index = 2
        while candidate.lower() in self._used_asset_names:
            candidate = f"{stem}_{index}{suffix}"
            index += 1
        self._used_asset_names.add(candidate.lower())
        return candidate

    def _heading_level(self, paragraph: Paragraph, style_lower: str) -> int | None:
        text = clean_text(paragraph.text)
        numbered = re.match(r"^(\d+(?:\.\d+)*)\.?\s+\S", text)
        if numbered:
            return max(1, min(6, len(numbered.group(1).split("."))))
        if self._is_references_heading(text):
            return 1

        match = re.search(r"(?:heading|заголовок)\s*(\d+)", style_lower)
        if match:
            return max(1, min(6, int(match.group(1))))
        outline = paragraph._p.pPr
        if outline is not None and outline.outlineLvl is not None:
            value = outline.outlineLvl.val
            try:
                numeric = int(value)
                # Word normally stores Heading 1 as 0. Some imported documents use
                # 1 for top-level headings; numbered text above handles that case.
                return max(1, min(6, numeric + 1))
            except (TypeError, ValueError):
                pass
        return None

    def _is_title(self, style_lower: str, text: str, title_found: bool) -> bool:
        if not text:
            return False
        normalized = re.sub(r"\s+", " ", style_lower).strip()
        title_styles = {
            "title",
            "document title",
            "название",
            "заглавие",
            "название документа",
        }
        return normalized in title_styles

    def _is_author(self, style_lower: str, text: str, title_found: bool, article: ArticleIR) -> bool:
        if not text or not title_found or article.metadata.abstracts:
            return False
        if any(token in style_lower for token in ("author", "автор", "subtitle", "подзаголовок")):
            return True
        return bool(re.match(r"^(автор(?:ы)?|authors?)\s*[:.]", text, flags=re.IGNORECASE))

    def _is_abstract(self, style_lower: str, text: str) -> bool:
        return "abstract" in style_lower or "аннотац" in style_lower or bool(
            re.match(r"^(аннотация|abstract)\s*[:.]", text, flags=re.IGNORECASE)
        )

    def _is_keywords(self, style_lower: str, text: str) -> bool:
        return "keyword" in style_lower or "ключев" in style_lower or bool(
            re.match(r"^(ключевые слова|keywords|key words)\s*[:.]", text, flags=re.IGNORECASE)
        )

    def _is_caption(self, style_lower: str, text: str) -> bool:
        if not text:
            return False
        if "caption" in style_lower or "подпись" in style_lower:
            return True
        return bool(
            re.match(
                r"^(?:рисунок|рис\.?|figure|fig\.?|таблица|table)\s*\d+",
                text,
                flags=re.IGNORECASE,
            )
        )

    def _is_list_item(self, paragraph: Paragraph) -> bool:
        p_pr = paragraph._p.pPr
        if p_pr is not None and p_pr.numPr is not None:
            return True
        style = ((paragraph.style.name if paragraph.style else "") or "").lower()
        return any(token in style for token in ("list", "список", "маркир", "нумер"))

    def _list_level(self, paragraph: Paragraph) -> int:
        p_pr = paragraph._p.pPr
        if p_pr is not None and p_pr.numPr is not None and p_pr.numPr.ilvl is not None:
            try:
                return int(p_pr.numPr.ilvl.val)
            except (TypeError, ValueError):
                pass
        style = ((paragraph.style.name if paragraph.style else "") or "").lower()
        match = re.search(r"(?:list|список|маркир|нумер)[^0-9]*(\d+)", style)
        return max(0, int(match.group(1)) - 1) if match else 0

    def _is_ordered_list(self, paragraph: Paragraph) -> bool:
        # python-docx не раскрывает numbering.xml в удобном API. Для MVP используем
        # имя стиля как признак; неизвестный список считаем маркированным.
        style = ((paragraph.style.name if paragraph.style else "") or "").lower()
        return any(token in style for token in ("number", "нумер"))


    def _clean_caption(self, text: str) -> str:
        value = clean_text(text)
        pattern = (
            r"^(?:рисунок|рис\.?|figure|fig\.?|таблица|table)"
            r"\s*\d+[A-Za-zА-Яа-я]?(?:[.:-]|\s*[–—-])?\s*"
        )
        cleaned = re.sub(pattern, "", value, flags=re.IGNORECASE).strip()
        return cleaned or value

    def _attach_caption(self, article: ArticleIR, caption: str) -> bool:
        for block in reversed(article.body):
            if isinstance(block, (FigureBlock, TableBlock)):
                if block.caption is None:
                    block.caption = caption
                    return True
                return False
            if isinstance(block, (ParagraphBlock, SectionBlock, ListItemBlock)):
                break
        return False

    def _save_authors(self, article: ArticleIR, lines: list[str]) -> None:
        names: list[str] = []
        for line in lines:
            parts = re.split(r"\s*[;,]\s*|\s+и\s+", line)
            names.extend(part.strip() for part in parts if part.strip())
        for name in names:
            article.metadata.authors.append(
                Author(id=f"author-{len(article.metadata.authors) + 1}", name=name)
            )

    def _postprocess_metadata(self, article: ArticleIR) -> None:
        if not article.metadata.titles:
            for index, block in enumerate(article.body[:4]):
                if isinstance(block, ParagraphBlock):
                    text = clean_text(block.text)
                    if 5 <= len(text) <= 300 and not self._looks_like_author_line(text):
                        article.metadata.titles.append(
                            LocalizedText(language=self._guess_language(text), text=text)
                        )
                        article.body.pop(index)
                        article.warnings.append(
                            "Название определено эвристически по первому содержательному абзацу."
                        )
                        break

        if not article.metadata.authors:
            for index, block in enumerate(article.body[:6]):
                if isinstance(block, SectionBlock):
                    break
                if isinstance(block, ParagraphBlock):
                    text = clean_text(block.text)
                    if self._looks_like_author_line(text):
                        names = [
                            part.strip()
                            for part in re.split(r"\s*[;,]\s*|\s+и\s+", text)
                            if part.strip()
                        ]
                        for name in names:
                            article.metadata.authors.append(
                                Author(id=f"author-{len(article.metadata.authors) + 1}", name=name)
                            )
                        article.body.pop(index)
                        article.warnings.append(
                            "Авторы определены эвристически по строке после названия."
                        )
                        break

    def _looks_like_author_line(self, text: str) -> bool:
        value = clean_text(text)
        if not value or len(value) > 300:
            return False
        initials = re.findall(
            r"(?:[А-ЯЁA-Z][а-яёA-Za-z'’\-]+)\s+[А-ЯЁA-Z]\.\s*[А-ЯЁA-Z]\.?",
            value,
        )
        return len(initials) >= 1 and not value.endswith(("?", "!"))

    def _is_secondary_title_paragraph(self, paragraph: Paragraph, text: str) -> bool:
        value = clean_text(text)
        letters = [character for character in value if character.isalpha()]
        if (
            not letters
            or len(value) < 15
            or len(value) > 500
            or any(character.isdigit() for character in value)
            or "doi" in value.lower()
        ):
            return False
        latin = sum("A" <= character.upper() <= "Z" for character in letters)
        uppercase = sum(character.isupper() for character in letters)
        return latin / len(letters) >= 0.75 and uppercase / len(letters) >= 0.85

    @staticmethod
    def _looks_like_secondary_author_line(text: str) -> bool:
        value = clean_text(text)
        if not value or len(value) > 300:
            return False
        initials_first = re.findall(
            r"\b[A-Z]\.\s*(?:[A-Z]\.\s*)?[A-Z][A-Za-z'’\-]+",
            value,
        )
        return bool(initials_first) and not value.endswith(("?", "!"))

    def _clean_heading_title(self, text: str) -> str:
        cleaned = re.sub(r"^\s*\d+(?:\.\d+)*\.?\s*", "", clean_text(text)).strip()
        return cleaned or clean_text(text)

    def _is_references_heading(self, text: str) -> bool:
        normalized = re.sub(r"^\d+(?:\.\d+)*\.?\s*", "", clean_text(text)).strip().lower()
        return normalized in {"литература", "список литературы", "references", "bibliography"}

    def _extract_doi(self, text: str) -> str | None:
        match = re.search(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+", text, flags=re.IGNORECASE)
        return match.group(0).rstrip(".,;)") if match else None

    def _append_ole_summary(self, article: ArticleIR) -> None:
        if not self._ole_formula_count:
            return
        article.warnings.append(
            "MathType/OLE: обнаружено "
            f"{self._ole_formula_count}, сохранено как изображения: {self._ole_formula_converted}, "
            f"не удалось преобразовать: {self._ole_formula_failed}."
        )
        if self._ole_formula_failed:
            article.warnings.append(
                "Для преобразования WMF/EMF установите Pillow на Windows или Inkscape и добавьте его в PATH."
            )
        for error in self._ole_conversion_errors[:5]:
            article.warnings.append(f"Ошибка преобразования формулы: {error}")

    def _strip_prefix(self, text: str, prefixes: tuple[str, ...]) -> str:
        value = text.strip()
        for prefix in prefixes:
            pattern = rf"^{re.escape(prefix)}\s*[:.\-–—]?\s*"
            value, count = re.subn(pattern, "", value, flags=re.IGNORECASE)
            if count:
                break
        return value.strip()

    def _split_keywords(self, text: str) -> list[str]:
        return [part.strip(" .") for part in re.split(r"[;,]", text) if part.strip(" .")]

    def _guess_language(self, text: str) -> str | None:
        cyrillic = len(re.findall(r"[А-Яа-яЁё]", text))
        latin = len(re.findall(r"[A-Za-z]", text))
        if cyrillic > latin:
            return "ru"
        if latin > cyrillic:
            return "en"
        return None
